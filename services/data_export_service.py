from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from exceptions import DataExportError, ValidationError
from models.dto.data_export import (
    ScoreExportContext,
    ScoreExportData,
    ScoreExportRow,
    StudentExportContext,
    StudentExportData,
    StudentExportRow,
)
from models.dto.student_filter import StudentFilter
from services.score_contract import ScoreServiceContract
from services.student_contract import StudentListServiceContract


class _WorkbookExportBase:
    HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")

    @staticmethod
    def _output_path(output_path: str | Path) -> Path:
        try:
            path = Path(output_path)
        except (TypeError, ValueError) as exc:
            raise ValidationError("Đường dẫn file xuất Excel không hợp lệ.") from exc
        if path.suffix.lower() != ".xlsx":
            raise ValidationError("File xuất phải có định dạng .xlsx.")
        if path.exists() and path.is_dir():
            raise ValidationError("Đường dẫn xuất phải là một file.")
        if not path.parent.exists() or not path.parent.is_dir():
            raise ValidationError("Thư mục lưu file xuất không tồn tại.")
        return path

    @staticmethod
    def _safe_text(cell, value: str | None) -> None:
        cell.value = "" if value is None else str(value)
        cell.data_type = "s"

    @classmethod
    def _write_headers(cls, worksheet, row: int, headers: tuple[str, ...]) -> None:
        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row, column, header)
            cell.font = Font(bold=True)
            cell.fill = cls.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def _finish_table(worksheet, header_row: int, columns: int, widths: tuple[int, ...]) -> None:
        worksheet.freeze_panes = f"A{header_row + 1}"
        worksheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(columns)}"
            f"{max(header_row, worksheet.max_row)}"
        )
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width

    @staticmethod
    def _save(workbook: Workbook, path: Path, message: str) -> Path:
        try:
            workbook.save(path)
        except Exception as exc:
            raise DataExportError(message) from exc
        finally:
            workbook.close()
        return path


class StudentExportService(_WorkbookExportBase):
    SHEET_TITLE = "Danh sách học sinh"
    EXPORT_TITLE = "DANH SÁCH HỌC SINH"
    HEADER_ROW = 7
    HEADERS = (
        "STT", "Mã học sinh", "Họ và tên", "Ngày sinh", "Giới tính",
        "Năm học", "Khối", "Lớp",
    )

    def __init__(self, student_service: StudentListServiceContract):
        self.student_service = student_service

    def get_export_data(self, context: StudentExportContext) -> StudentExportData:
        self._validate_context(context)
        items = self.student_service.list_students(context.filters)
        seen: set[str] = set()
        rows: list[StudentExportRow] = []
        for item in items:
            if item.student_id in seen:
                continue
            seen.add(item.student_id)
            rows.append(StudentExportRow(
                student_id=item.student_id,
                student_code=item.student_code,
                full_name=item.full_name,
                date_of_birth=item.date_of_birth,
                gender=item.gender,
                school_year_name=(
                    item.current_school_year_name or context.school_year_name
                ),
                grade_number=item.current_grade_number,
                class_name=item.current_class_name,
            ))
        rows.sort(key=lambda row: (row.full_name.casefold(), row.student_code, row.student_id))
        return StudentExportData(context=context, rows=tuple(rows))

    def export_xlsx(self, context: StudentExportContext, output_path: str | Path) -> Path:
        data = self.get_export_data(context)
        path = self._output_path(output_path)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.SHEET_TITLE
        worksheet.merge_cells("A1:H1")
        worksheet["A1"] = self.EXPORT_TITLE
        worksheet["A1"].font = Font(size=16, bold=True)
        worksheet["A1"].alignment = Alignment(horizontal="center")
        context_values = (
            ("Năm học", context.school_year_name or "Tất cả"),
            ("Khối", context.grade_name or "Tất cả"),
            ("Lớp", context.class_name or "Tất cả"),
        )
        for row_index, (label, value) in enumerate(context_values, start=3):
            worksheet.cell(row_index, 1, label).font = Font(bold=True)
            self._safe_text(worksheet.cell(row_index, 2), value)
        self._write_headers(worksheet, self.HEADER_ROW, self.HEADERS)
        for index, item in enumerate(data.rows, start=1):
            row = self.HEADER_ROW + index
            worksheet.cell(row, 1, index)
            self._safe_text(worksheet.cell(row, 2), item.student_code)
            self._safe_text(worksheet.cell(row, 3), item.full_name)
            worksheet.cell(row, 4, item.date_of_birth)
            if item.date_of_birth is not None:
                worksheet.cell(row, 4).number_format = "dd/mm/yyyy"
            self._safe_text(worksheet.cell(row, 5), item.gender or "")
            self._safe_text(worksheet.cell(row, 6), item.school_year_name or "")
            worksheet.cell(row, 7, item.grade_number)
            self._safe_text(worksheet.cell(row, 8), item.class_name or "")
        self._finish_table(worksheet, self.HEADER_ROW, len(self.HEADERS), (7, 18, 28, 14, 12, 16, 10, 16))
        return self._save(
            workbook, path,
            "Không thể ghi file danh sách học sinh. Vui lòng kiểm tra đường dẫn và quyền truy cập.",
        )

    @staticmethod
    def _validate_context(context: StudentExportContext) -> None:
        if not isinstance(context, StudentExportContext) or not isinstance(context.filters, StudentFilter):
            raise ValidationError("Ngữ cảnh xuất danh sách học sinh không hợp lệ.")


class ScoreExportService(_WorkbookExportBase):
    SHEET_TITLE = "Bảng điểm"
    EXPORT_TITLE = "BẢNG ĐIỂM"
    HEADER_ROW = 8
    HEADERS = (
        "STT", "Mã học sinh", "Họ và tên", "Lớp", "Môn",
        "Bài đánh giá", "Điểm",
    )

    def __init__(self, score_service: ScoreServiceContract):
        self.score_service = score_service

    def get_export_data(self, context: ScoreExportContext) -> ScoreExportData:
        self._validate_context(context)
        source = self.score_service.list_score_roster(
            context.class_id,
            context.school_year_id,
            context.subject_id,
            context.assessment_id,
        )
        seen: set[int] = set()
        rows: list[ScoreExportRow] = []
        for item in source:
            if item.enrollment_id in seen:
                continue
            seen.add(item.enrollment_id)
            rows.append(ScoreExportRow(
                enrollment_id=item.enrollment_id,
                student_id=item.student_id,
                student_code=item.student_code,
                full_name=item.full_name,
                score=item.score,
            ))
        rows.sort(key=lambda row: (row.full_name.casefold(), row.student_code, row.enrollment_id))
        return ScoreExportData(context=context, rows=tuple(rows))

    def export_xlsx(self, context: ScoreExportContext, output_path: str | Path) -> Path:
        data = self.get_export_data(context)
        path = self._output_path(output_path)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.SHEET_TITLE
        worksheet.merge_cells("A1:G1")
        worksheet["A1"] = self.EXPORT_TITLE
        worksheet["A1"].font = Font(size=16, bold=True)
        worksheet["A1"].alignment = Alignment(horizontal="center")
        context_values = (
            ("Năm học", context.school_year_name),
            ("Lớp", context.class_name),
            ("Môn", context.subject_name),
            ("Bài đánh giá", context.assessment_name),
        )
        for row_index, (label, value) in enumerate(context_values, start=3):
            worksheet.cell(row_index, 1, label).font = Font(bold=True)
            self._safe_text(worksheet.cell(row_index, 2), value)
        self._write_headers(worksheet, self.HEADER_ROW, self.HEADERS)
        for index, item in enumerate(data.rows, start=1):
            row = self.HEADER_ROW + index
            worksheet.cell(row, 1, index)
            self._safe_text(worksheet.cell(row, 2), item.student_code)
            self._safe_text(worksheet.cell(row, 3), item.full_name)
            self._safe_text(worksheet.cell(row, 4), context.class_name)
            self._safe_text(worksheet.cell(row, 5), context.subject_name)
            self._safe_text(worksheet.cell(row, 6), context.assessment_name)
            worksheet.cell(row, 7, item.score)
        self._finish_table(worksheet, self.HEADER_ROW, len(self.HEADERS), (7, 18, 28, 16, 22, 24, 12))
        return self._save(
            workbook, path,
            "Không thể ghi file bảng điểm. Vui lòng kiểm tra đường dẫn và quyền truy cập.",
        )

    @staticmethod
    def _validate_context(context: ScoreExportContext) -> None:
        if not isinstance(context, ScoreExportContext):
            raise ValidationError("Ngữ cảnh xuất bảng điểm không hợp lệ.")
        for field_name in ("school_year_id", "class_id", "subject_id", "assessment_id"):
            value = getattr(context, field_name)
            if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
                raise ValidationError(f"{field_name} không hợp lệ.")
        for field_name in ("school_year_name", "class_name", "subject_name", "assessment_name"):
            value = getattr(context, field_name)
            if not isinstance(value, str) or not value.strip():
                raise ValidationError(f"{field_name} không được để trống.")
