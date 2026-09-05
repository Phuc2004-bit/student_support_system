from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from exceptions import ReportExportError, ValidationError
from models.dto.report_dto import SupportReportData
from models.dto.report_export import (
    SupportReportExportContext,
    SupportReportExportData,
)
from models.enums import InterventionStatus
from utils.report_labels import review_result_label, status_label


class ReportExportService:
    SHEET_TITLE = "Báo cáo bổ trợ"
    REPORT_TITLE = "BÁO CÁO BỔ TRỢ HỌC TẬP"
    TABLE_HEADER_ROW = 15

    SUMMARY_FIELDS = (
        ("Tổng số ca", "total_cases"),
        ("Mới phát hiện", "detected_count"),
        ("Đã lập kế hoạch", "planned_count"),
        ("Đang bổ trợ", "in_progress_count"),
        ("Chờ đánh giá", "waiting_review_count"),
        ("Cần tiếp tục", "continue_count"),
        ("Đã đạt ngưỡng", "completed_count"),
    )

    TABLE_HEADERS = (
        "STT",
        "Mã học sinh",
        "Họ và tên",
        "Khối",
        "Lớp",
        "Năm học",
        "Môn",
        "Điểm phát hiện",
        "Ngày phát hiện",
        "Trạng thái",
        "Ngày đánh giá gần nhất",
        "Điểm đánh giá gần nhất",
        "Kết quả đánh giá gần nhất",
    )

    def export_xlsx(
        self,
        export_data: SupportReportExportData,
        output_path: str | Path,
    ) -> Path:
        self._validate_export_data(export_data)
        path = self._validate_output_path(output_path)

        workbook = Workbook()
        try:
            self._build_workbook(workbook, export_data)
            workbook.save(path)
        except Exception as exc:
            raise ReportExportError(
                "Không thể ghi file báo cáo. Vui lòng kiểm tra đường dẫn "
                "và quyền truy cập."
            ) from exc
        finally:
            workbook.close()

        return path

    def _build_workbook(
        self,
        workbook: Workbook,
        export_data: SupportReportExportData,
    ) -> None:
        worksheet = workbook.active
        worksheet.title = self.SHEET_TITLE
        context = export_data.context

        worksheet.merge_cells("A1:M1")
        title_cell = worksheet["A1"]
        title_cell.value = self.REPORT_TITLE
        title_cell.font = Font(size=16, bold=True)
        title_cell.alignment = Alignment(horizontal="center")

        context_values = (
            ("A3", "Năm học", "B3", context.school_year_name),
            ("D3", "Khối", "E3", context.grade_name or "Tất cả"),
            ("G3", "Lớp", "H3", context.class_name or "Tất cả"),
            ("J3", "Môn", "K3", context.subject_name or "Tất cả"),
            (
                "A4",
                "Trạng thái",
                "B4",
                (
                    status_label(self._status_value(context.status))
                    if context.status
                    else "Tất cả"
                ),
            ),
        )
        for label_cell, label, value_cell, value in context_values:
            worksheet[label_cell] = label
            worksheet[label_cell].font = Font(bold=True)
            worksheet[value_cell] = value

        worksheet.merge_cells("A6:B6")
        worksheet["A6"] = "Tổng quan"
        worksheet["A6"].font = Font(bold=True)
        summary = export_data.report.summary
        for row_index, (label, field_name) in enumerate(
            self.SUMMARY_FIELDS,
            start=7,
        ):
            worksheet.cell(row_index, 1, label)
            worksheet.cell(row_index, 2, getattr(summary, field_name))

        header_fill = PatternFill("solid", fgColor="D9EAF7")
        for column, header in enumerate(self.TABLE_HEADERS, start=1):
            cell = worksheet.cell(self.TABLE_HEADER_ROW, column, header)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for row_index, item in enumerate(
            export_data.report.rows,
            start=self.TABLE_HEADER_ROW + 1,
        ):
            values = (
                row_index - self.TABLE_HEADER_ROW,
                item.student_code,
                item.full_name,
                item.grade_number,
                item.class_name,
                context.school_year_name,
                item.subject_name,
                item.trigger_score,
                item.detected_date,
                status_label(item.status),
                item.latest_review_date or "-",
                (
                    item.latest_review_score
                    if item.latest_review_score is not None
                    else "-"
                ),
                review_result_label(item.latest_review_result),
            )
            for column, value in enumerate(values, start=1):
                worksheet.cell(row_index, column, value)

            worksheet.cell(row_index, 9).number_format = "dd/mm/yyyy"
            if item.latest_review_date is not None:
                worksheet.cell(row_index, 11).number_format = "dd/mm/yyyy"

        worksheet.freeze_panes = f"A{self.TABLE_HEADER_ROW + 1}"
        worksheet.auto_filter.ref = (
            f"A{self.TABLE_HEADER_ROW}:"
            f"M{max(self.TABLE_HEADER_ROW, worksheet.max_row)}"
        )
        widths = (7, 16, 26, 9, 14, 14, 20, 16, 17, 20, 23, 24, 28)
        for column, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(column)].width = width

    @classmethod
    def _validate_export_data(
        cls,
        export_data: SupportReportExportData,
    ) -> None:
        if not isinstance(export_data, SupportReportExportData):
            raise ValidationError("Dữ liệu xuất báo cáo không hợp lệ.")

        context = export_data.context
        if not isinstance(context, SupportReportExportContext):
            raise ValidationError("Ngữ cảnh xuất báo cáo không hợp lệ.")
        if not isinstance(export_data.report, SupportReportData):
            raise ValidationError("Snapshot xuất báo cáo không hợp lệ.")
        if (
            isinstance(context.school_year_id, bool)
            or not isinstance(context.school_year_id, int)
            or context.school_year_id <= 0
        ):
            raise ValidationError("Năm học xuất báo cáo không hợp lệ.")
        if (
            not isinstance(context.school_year_name, str)
            or not context.school_year_name.strip()
        ):
            raise ValidationError("Tên năm học xuất báo cáo là bắt buộc.")

        for field_name in ("grade_id", "class_id", "subject_id"):
            value = getattr(context, field_name)
            if value is not None and (
                isinstance(value, bool)
                or not isinstance(value, int)
                or value <= 0
            ):
                raise ValidationError(
                    f"{field_name} của báo cáo không hợp lệ."
                )

        if context.status is not None:
            valid_statuses = {
                status.value for status in InterventionStatus
            }
            if cls._status_value(context.status) not in valid_statuses:
                raise ValidationError(
                    "Trạng thái lọc báo cáo không hợp lệ."
                )

    @staticmethod
    def _status_value(
        status: str | InterventionStatus,
    ) -> str:
        return status.value if isinstance(status, InterventionStatus) else status

    @staticmethod
    def _validate_output_path(output_path: str | Path) -> Path:
        try:
            path = Path(output_path)
        except (TypeError, ValueError) as exc:
            raise ValidationError(
                "Đường dẫn file báo cáo không hợp lệ."
            ) from exc

        if path.suffix.lower() != ".xlsx":
            raise ValidationError("File báo cáo phải có định dạng .xlsx.")
        if path.exists() and path.is_dir():
            raise ValidationError("Đường dẫn xuất báo cáo phải là một file.")
        if not path.parent.exists() or not path.parent.is_dir():
            raise ValidationError(
                "Thư mục lưu báo cáo không tồn tại."
            )
        return path
