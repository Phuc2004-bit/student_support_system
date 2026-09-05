from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from exceptions import ScoreImportError
from models.dto.score_import import (
    ScoreImportIssue,
    ScoreImportRow,
    ScoreImportSourceMetadata,
    ScoreImportWorkbook,
)


class ScoreImportWorkbookParser:
    SHEET_TITLE = "Nhập điểm"
    HEADER_ROW = 6
    HEADERS = ("STT", "Mã học sinh", "Họ và tên", "Điểm")
    HEADER_ALIASES = {
        "stt": "sequence",
        "mã học sinh": "student_id",
        "student_id": "student_id",
        "họ và tên": "student_name",
        "full_name": "student_name",
        "điểm": "score",
        "score": "score",
    }
    REQUIRED_FIELDS = ("sequence", "student_id", "student_name", "score")

    def parse_workbook(self, file_path: str | Path) -> ScoreImportWorkbook:
        path = self._validate_input_path(file_path)
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            raise ScoreImportError(
                "Không thể đọc file Excel. File có thể bị hỏng hoặc không có quyền truy cập."
            ) from exc

        try:
            if self.SHEET_TITLE not in workbook.sheetnames:
                raise ScoreImportError(
                    f"File Excel phải có sheet '{self.SHEET_TITLE}'."
                )
            worksheet = workbook[self.SHEET_TITLE]
            header_map = self._build_header_map(
                [
                    worksheet.cell(self.HEADER_ROW, column).value
                    for column in range(1, worksheet.max_column + 1)
                ]
            )
            return ScoreImportWorkbook(
                file_name=path.name,
                sheet_name=worksheet.title,
                metadata=self._read_metadata(worksheet),
                rows=tuple(self._read_rows(worksheet, header_map)),
            )
        finally:
            workbook.close()

    @staticmethod
    def _validate_input_path(file_path: str | Path) -> Path:
        try:
            path = Path(file_path)
        except (TypeError, ValueError) as exc:
            raise ScoreImportError("Đường dẫn file Excel không hợp lệ.") from exc
        if path.suffix.lower() != ".xlsx":
            raise ScoreImportError("Chỉ hỗ trợ file Excel định dạng .xlsx.")
        if not path.exists() or not path.is_file():
            raise ScoreImportError("Không tìm thấy file Excel.")
        return path

    def _build_header_map(self, headers: list[Any]) -> dict[str, int]:
        mapped: dict[str, int] = {}
        for index, header in enumerate(headers, start=1):
            field_name = self.HEADER_ALIASES.get(self._normalize_header(header))
            if field_name is None:
                continue
            if field_name in mapped:
                raise ScoreImportError(
                    f"Cột '{header}' bị trùng ý nghĩa trong file Excel."
                )
            mapped[field_name] = index
        missing = [field for field in self.REQUIRED_FIELDS if field not in mapped]
        if missing:
            labels = {
                "sequence": "STT",
                "student_id": "Mã học sinh",
                "student_name": "Họ và tên",
                "score": "Điểm",
            }
            raise ScoreImportError(
                "Thiếu cột bắt buộc: "
                + ", ".join(labels[field] for field in missing)
                + "."
            )
        return mapped

    def _read_rows(self, worksheet, header_map: dict[str, int]) -> list[ScoreImportRow]:
        result: list[ScoreImportRow] = []
        for excel_row in range(self.HEADER_ROW + 1, worksheet.max_row + 1):
            values = {
                field: worksheet.cell(excel_row, column).value
                for field, column in header_map.items()
            }
            if self._is_blank_row(values.values()):
                continue
            issues: list[ScoreImportIssue] = []
            raw_score = values.get("score")
            result.append(
                ScoreImportRow(
                    row_number=excel_row,
                    student_id=self._normalize_student_id(
                        values.get("student_id"), issues
                    ),
                    student_name=self._normalize_student_name(
                        values.get("student_name"), issues
                    ),
                    raw_score=raw_score,
                    normalized_score=self._normalize_score(raw_score, issues),
                    issues=tuple(issues),
                )
            )
        return result

    @staticmethod
    def _read_metadata(worksheet) -> ScoreImportSourceMetadata:
        return ScoreImportSourceMetadata(
            school_year_name=ScoreImportWorkbookParser._optional_text(
                worksheet["B1"].value
            ),
            class_name=ScoreImportWorkbookParser._optional_text(
                worksheet["B2"].value
            ),
            subject_name=ScoreImportWorkbookParser._optional_text(
                worksheet["B3"].value
            ),
            assessment_name=ScoreImportWorkbookParser._optional_text(
                worksheet["B4"].value
            ),
        )

    @staticmethod
    def _normalize_student_id(value: Any, issues: list[ScoreImportIssue]) -> str:
        if value is None or (isinstance(value, str) and not value.strip()):
            issues.append(
                ScoreImportIssue(
                    "MISSING_STUDENT_ID",
                    "student_id",
                    "Mã học sinh không được để trống.",
                )
            )
            return ""
        if isinstance(value, bool) or not isinstance(value, (str, int)):
            issues.append(
                ScoreImportIssue(
                    "INVALID_STUDENT_ID_TYPE",
                    "student_id",
                    "Kiểu dữ liệu mã học sinh không được hỗ trợ.",
                )
            )
            return ""
        return str(value).strip()

    @staticmethod
    def _normalize_student_name(
        value: Any,
        issues: list[ScoreImportIssue],
    ) -> str | None:
        if value is None or (isinstance(value, str) and not value.strip()):
            return None
        if not isinstance(value, str):
            issues.append(
                ScoreImportIssue(
                    "INVALID_STUDENT_NAME_TYPE",
                    "student_name",
                    "Kiểu dữ liệu họ tên không được hỗ trợ.",
                )
            )
            return None
        return value.strip()

    @staticmethod
    def _normalize_score(
        value: Any,
        issues: list[ScoreImportIssue],
    ) -> Decimal | None:
        if value is None or (isinstance(value, str) and not value.strip()):
            return None
        if isinstance(value, bool):
            issues.append(
                ScoreImportIssue(
                    "BOOLEAN_SCORE",
                    "score",
                    "Giá trị đúng/sai không phải là điểm số.",
                )
            )
            return None
        if isinstance(value, str) and value.lstrip().startswith("="):
            issues.append(
                ScoreImportIssue(
                    "FORMULA_SCORE",
                    "score",
                    "Ô điểm chứa công thức và không được tự động tính khi import.",
                )
            )
            return None
        if not isinstance(value, (Decimal, int, float, str)):
            issues.append(
                ScoreImportIssue(
                    "UNSUPPORTED_SCORE_TYPE",
                    "score",
                    "Kiểu dữ liệu điểm không được hỗ trợ.",
                )
            )
            return None
        try:
            normalized = Decimal(str(value).strip())
        except (InvalidOperation, ValueError):
            issues.append(
                ScoreImportIssue(
                    "INVALID_SCORE",
                    "score",
                    "Điểm không phải là giá trị số hợp lệ.",
                )
            )
            return None
        if not normalized.is_finite():
            issues.append(
                ScoreImportIssue(
                    "INVALID_SCORE",
                    "score",
                    "Điểm không phải là giá trị số hữu hạn.",
                )
            )
            return None
        return normalized

    @staticmethod
    def _normalize_header(value: Any) -> str:
        if not isinstance(value, str):
            return ""
        return " ".join(value.strip().casefold().split())

    @staticmethod
    def _optional_text(value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _is_blank_row(values) -> bool:
        return all(
            value is None or (isinstance(value, str) and not value.strip())
            for value in values
        )

