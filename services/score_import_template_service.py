from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection

from exceptions import ScoreImportError
from models.dto.score_import import ScoreImportContext, ScoreImportTemplateStudent


class ScoreImportTemplateService:
    SHEET_TITLE = "Nhập điểm"
    HEADER_ROW = 6
    HEADERS = ("STT", "Mã học sinh", "Họ và tên", "Điểm")

    def create_template(
        self,
        context: ScoreImportContext,
        roster: Iterable[ScoreImportTemplateStudent],
        output_path: str | Path,
    ) -> Path:
        self._validate_context(context)
        students = self._normalize_roster(roster)
        path = self._validate_output_path(output_path)
        workbook = Workbook()
        try:
            self._build_workbook(workbook, context, students)
            workbook.save(path)
        except Exception as exc:
            raise ScoreImportError(
                "Không thể tạo file mẫu nhập điểm. "
                "Vui lòng kiểm tra đường dẫn và quyền truy cập."
            ) from exc
        finally:
            workbook.close()
        return path

    def _build_workbook(
        self,
        workbook: Workbook,
        context: ScoreImportContext,
        students: tuple[ScoreImportTemplateStudent, ...],
    ) -> None:
        worksheet = workbook.active
        worksheet.title = self.SHEET_TITLE
        metadata = (
            ("Năm học", context.school_year_name),
            ("Lớp", context.class_name),
            ("Môn", context.subject_name),
            ("Bài đánh giá", context.assessment_name),
        )
        for row, (label, value) in enumerate(metadata, start=1):
            worksheet.cell(row, 1, label).font = Font(bold=True)
            self._set_text(worksheet.cell(row, 2), value)

        fill = PatternFill("solid", fgColor="D9EAF7")
        for column, header in enumerate(self.HEADERS, start=1):
            cell = worksheet.cell(self.HEADER_ROW, column, header)
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

        for index, student in enumerate(students, start=1):
            row = self.HEADER_ROW + index
            worksheet.cell(row, 1, index)
            self._set_text(worksheet.cell(row, 2), student.student_id)
            self._set_text(worksheet.cell(row, 3), student.full_name)
            score_cell = worksheet.cell(row, 4)
            score_cell.value = None
            score_cell.protection = Protection(locked=False)

        worksheet.freeze_panes = f"A{self.HEADER_ROW + 1}"
        worksheet.auto_filter.ref = (
            f"A{self.HEADER_ROW}:D{max(self.HEADER_ROW, worksheet.max_row)}"
        )
        worksheet.column_dimensions["A"].width = 8
        worksheet.column_dimensions["B"].width = 24
        worksheet.column_dimensions["C"].width = 32
        worksheet.column_dimensions["D"].width = 14
        worksheet.protection.sheet = True

    @staticmethod
    def _set_text(cell, value: str) -> None:
        cell.value = value
        cell.data_type = "s"

    @staticmethod
    def _validate_context(context: ScoreImportContext) -> None:
        if not isinstance(context, ScoreImportContext):
            raise ScoreImportError("Ngữ cảnh tạo file mẫu không hợp lệ.")
        for field_name in (
            "school_year_id",
            "class_id",
            "subject_id",
            "assessment_id",
        ):
            value = getattr(context, field_name)
            if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
                raise ScoreImportError(f"{field_name} không hợp lệ.")
        for field_name in (
            "school_year_name",
            "class_name",
            "subject_name",
            "assessment_name",
        ):
            value = getattr(context, field_name)
            if not isinstance(value, str) or not value.strip():
                raise ScoreImportError(f"{field_name} không được để trống.")

    @staticmethod
    def _normalize_roster(
        roster: Iterable[ScoreImportTemplateStudent],
    ) -> tuple[ScoreImportTemplateStudent, ...]:
        try:
            raw_students = tuple(roster)
        except TypeError as exc:
            raise ScoreImportError("Danh sách học sinh không hợp lệ.") from exc
        normalized: list[ScoreImportTemplateStudent] = []
        seen_ids: set[str] = set()
        for item in raw_students:
            if not isinstance(item, ScoreImportTemplateStudent):
                raise ScoreImportError("Dữ liệu học sinh trong roster không hợp lệ.")
            student_id = (
                item.student_id.strip()
                if isinstance(item.student_id, str)
                else ""
            )
            full_name = (
                item.full_name.strip()
                if isinstance(item.full_name, str)
                else ""
            )
            if not student_id or not full_name:
                raise ScoreImportError("Mã học sinh và họ tên là bắt buộc.")
            if student_id in seen_ids:
                raise ScoreImportError("Danh sách học sinh có student_id bị trùng.")
            seen_ids.add(student_id)
            normalized.append(ScoreImportTemplateStudent(student_id, full_name))
        return tuple(
            sorted(
                normalized,
                key=lambda item: (item.full_name.casefold(), item.student_id),
            )
        )

    @staticmethod
    def _validate_output_path(output_path: str | Path) -> Path:
        try:
            path = Path(output_path)
        except (TypeError, ValueError) as exc:
            raise ScoreImportError("Đường dẫn file mẫu không hợp lệ.") from exc
        if path.suffix.lower() != ".xlsx":
            raise ScoreImportError("File mẫu phải có định dạng .xlsx.")
        if not path.parent.exists() or not path.parent.is_dir():
            raise ScoreImportError("Thư mục lưu file mẫu không tồn tại.")
        if path.exists() and path.is_dir():
            raise ScoreImportError("Đường dẫn file mẫu không hợp lệ.")
        return path

