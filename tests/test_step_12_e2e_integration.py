from __future__ import annotations

from datetime import date
from decimal import Decimal
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import bcrypt
from openpyxl import load_workbook
import pyodbc
import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QDialog

from app_context import AppContext
from config.database import db_settings
from database.connection import DatabaseManager
from exceptions import BusinessRuleError, DatabaseError, ValidationError
from models.dto import StudentCreateData
from models.dto.report_export import (
    SupportReportExportContext,
    SupportReportExportData,
)
from models.enums import AssessmentStatus, InterventionStatus, UserRole
from repositories import (
    EnrollmentRepository,
    InterventionRepository,
    ScoreRepository,
    StudentRepository,
    UserRepository,
)
from services.academic_service import AcademicService
from services.auth_service import AuthService
from services.dashboard_service import DashboardService
from services.enrollment_service import EnrollmentService
from services.permission_service import PermissionService
from services.report_export_service import ReportExportService
from services.report_service import ReportService
from services.score_service import ScoreService
from services.student_list_service import StudentListService
from services.student_profile_service import StudentProfileService
from services.student_service import StudentService
from services.support_service import SupportService
from services.user_service import UserService
from ui.main_window import MainWindow
from ui.pages.catalog_page import CatalogPage
from ui.pages.reports_page import ReportsPage
from ui.pages.scores_page import ScoresPage
from ui.pages.students_page import StudentsPage
from ui.pages.support_page import SupportPage
from ui.pages.system_page import SystemPage
from ui.widgets.support_filter_widget import SupportFilterWidget
from utils.report_labels import review_result_label, status_label


PREFIX = "T1210"
YEAR = f"{PREFIX}_2040_2041"
CLASS_NAME = f"{PREFIX}_11A"
CLASS_UPDATED = f"{PREFIX}_11B"
GRADE_NUMBER = 11
GRADE_NAME = f"{PREFIX} Grade 11"
GRADE_UPDATED = f"{PREFIX} Grade Eleven"
SUBJECT_CODE = f"{PREFIX}_SUB"
SUBJECT_NAME = f"{PREFIX} Subject"
SUBJECT_UPDATED = f"{PREFIX} Subject Updated"
ASSESSMENT = f"{PREFIX}_TRIGGER"
ASSESSMENT_UPDATED = f"{PREFIX}_TRIGGER_UPDATED"
REVIEW_ASSESSMENT = f"{PREFIX}_REVIEW"
STUDENT_IDS = (f"{PREFIX}-student-1", f"{PREFIX}-student-2")
STUDENT_CODES = (f"{PREFIX}S1", f"{PREFIX}S2")
ADMIN_USERNAME = f"{PREFIX.lower()}_admin"
TEACHER_USERNAME = f"{PREFIX.lower()}_teacher"
ROLLBACK_USERNAME = f"{PREFIX.lower()}_rollback"
OLD_PASSWORD = "T1210Password@123"
NEW_PASSWORD = "T1210Changed@456"


def qapp() -> QApplication:
    return QApplication.instance() or QApplication([])


def require_test_database(connection_string: str) -> str:
    normalized = connection_string.upper().replace(" ", "")
    if "DATABASE=STUDENT_SUPPORT_DB_TEST;" not in normalized:
        raise RuntimeError("Step 12.10 writes require student_support_db_test.")
    return connection_string


def get_test_db() -> DatabaseManager:
    connection_string = db_settings.connection_string().replace(
        "DATABASE=student_support_db;",
        "DATABASE=student_support_db_test;",
    )
    return DatabaseManager(require_test_database(connection_string))


def assert_connected_test_database(db: DatabaseManager) -> None:
    require_test_database(db.connection_string)
    connection = db.get_connection()
    try:
        actual = connection.cursor().execute("SELECT DB_NAME()").fetchone()[0]
        assert actual == "student_support_db_test"
    finally:
        connection.rollback()
        connection.close()


def cleanup(db: DatabaseManager) -> None:
    assert_connected_test_database(db)
    with db.transaction() as connection:
        cursor = connection.cursor()
        cursor.execute(
            """
            DELETE FROM dbo.INTERVENTION_REVIEWS
            WHERE intervention_id IN
            (
                SELECT i.intervention_id
                FROM dbo.INTERVENTIONS AS i
                INNER JOIN dbo.STUDENT_ENROLLMENTS AS e
                    ON e.enrollment_id = i.enrollment_id
                WHERE e.student_id IN (?, ?)
            )
            """,
            *STUDENT_IDS,
        )
        cursor.execute(
            """
            DELETE FROM dbo.INTERVENTIONS
            WHERE enrollment_id IN
            (
                SELECT enrollment_id FROM dbo.STUDENT_ENROLLMENTS
                WHERE student_id IN (?, ?)
            )
            """,
            *STUDENT_IDS,
        )
        cursor.execute(
            """
            DELETE FROM dbo.SCORES
            WHERE enrollment_id IN
            (
                SELECT enrollment_id FROM dbo.STUDENT_ENROLLMENTS
                WHERE student_id IN (?, ?)
            )
            """,
            *STUDENT_IDS,
        )
        cursor.execute(
            "DELETE FROM dbo.STUDENT_ENROLLMENTS WHERE student_id IN (?, ?)",
            *STUDENT_IDS,
        )
        cursor.execute(
            """
            DELETE FROM dbo.STUDENTS
            WHERE student_id IN (?, ?) OR student_code IN (?, ?)
            """,
            *STUDENT_IDS,
            *STUDENT_CODES,
        )
        cursor.execute(
            """
            DELETE FROM dbo.SUPPORT_RULES
            WHERE subject_id IN
            (
                SELECT subject_id FROM dbo.SUBJECTS WHERE subject_code = ?
            )
            """,
            SUBJECT_CODE,
        )
        cursor.execute(
            """
            DELETE FROM dbo.ASSESSMENTS
            WHERE subject_id IN
            (
                SELECT subject_id FROM dbo.SUBJECTS WHERE subject_code = ?
            )
               OR assessment_name IN (?, ?, ?)
            """,
            SUBJECT_CODE,
            ASSESSMENT,
            ASSESSMENT_UPDATED,
            REVIEW_ASSESSMENT,
        )
        cursor.execute(
            """
            DELETE c FROM dbo.CLASSES AS c
            INNER JOIN dbo.SCHOOL_YEARS AS sy
                ON sy.school_year_id = c.school_year_id
            WHERE sy.year_name = ? OR c.class_name IN (?, ?)
            """,
            YEAR,
            CLASS_NAME,
            CLASS_UPDATED,
        )
        cursor.execute("DELETE FROM dbo.SUBJECTS WHERE subject_code = ?", SUBJECT_CODE)
        cursor.execute("DELETE FROM dbo.SCHOOL_YEARS WHERE year_name = ?", YEAR)
        cursor.execute(
            """
            DELETE FROM dbo.GRADES
            WHERE grade_number = ? AND grade_name IN (?, ?)
            """,
            GRADE_NUMBER,
            GRADE_NAME,
            GRADE_UPDATED,
        )
        cursor.execute(
            "DELETE FROM dbo.USERS WHERE username IN (?, ?, ?)",
            ADMIN_USERNAME,
            TEACHER_USERNAME,
            ROLLBACK_USERNAME,
        )


def build_context(db: DatabaseManager, session) -> AppContext:
    academic = AcademicService(db)
    report = ReportService(db)
    users = UserService(db)
    return AppContext(
        db=db,
        auth_service=AuthService(db),
        permission_service=PermissionService(),
        session=session,
        academic_service=academic,
        dashboard_service=DashboardService(db),
        student_list_service=StudentListService(db),
        student_service=StudentService(db),
        enrollment_service=EnrollmentService(db),
        student_profile_service=StudentProfileService(db),
        score_service=ScoreService(db),
        report_service=report,
        support_service=SupportService(db),
        user_service=users,
    )


def select(combo, value) -> None:
    index = combo.findData(value)
    assert index >= 0
    combo.setCurrentIndex(index)


def select_table_identity(table, identity: int) -> None:
    for row in range(table.rowCount()):
        if table.item(row, 0).data(Qt.ItemDataRole.UserRole) == identity:
            table.setCurrentCell(row, 0)
            return
    raise AssertionError(f"Identity {identity} is not present in the table")


def table_counts(db: DatabaseManager) -> tuple[int, ...]:
    tables = (
        "SCHOOL_YEARS",
        "GRADES",
        "CLASSES",
        "SUBJECTS",
        "ASSESSMENTS",
        "SUPPORT_RULES",
        "USERS",
        "SCORES",
        "INTERVENTIONS",
        "INTERVENTION_REVIEWS",
    )
    with db.transaction() as connection:
        cursor = connection.cursor()
        return tuple(
            int(cursor.execute(f"SELECT COUNT_BIG(*) FROM dbo.{table}").fetchone()[0])
            for table in tables
        )


class AcceptedDialog:
    DialogCode = QDialog.DialogCode

    def __init__(self, values):
        self._values = values

    def exec(self):
        return self.DialogCode.Accepted

    def values(self):
        return self._values


class AcceptedUserDialog:
    DialogCode = QDialog.DialogCode

    def __init__(self, create_values=None, update_values=None):
        self._create_values = create_values
        self._update_values = update_values

    def exec(self):
        return self.DialogCode.Accepted

    def create_values(self):
        return self._create_values

    def update_values(self):
        return self._update_values


def test_database_guard_rejects_production_connection_before_write():
    with pytest.raises(RuntimeError):
        require_test_database(
            "SERVER=.\\SQLEXPRESS;DATABASE=student_support_db;Trusted_Connection=yes;"
        )


def test_admin_catalog_reports_export_and_historical_safety_e2e(tmp_path):
    qapp()
    db = get_test_db()
    assert_connected_test_database(db)
    cleanup(db)
    academic = AcademicService(db)
    users = UserService(db)
    auth = AuthService(db)
    window = None
    output = tmp_path / f"{PREFIX}_support_report.xlsx"

    try:
        users.create_user(
            ADMIN_USERNAME,
            OLD_PASSWORD,
            f"{PREFIX} Administrator",
            UserRole.ADMIN,
        )
        admin = auth.login(ADMIN_USERNAME, OLD_PASSWORD)
        context = build_context(db, admin)
        window = MainWindow(context)

        assert all(window.can_navigate_to(key) for key in MainWindow.PAGE_TITLES)
        assert isinstance(window.pages["students"], StudentsPage)
        assert isinstance(window.pages["scores"], ScoresPage)
        assert isinstance(window.pages["support"], SupportPage)
        assert isinstance(window.pages["reports"], ReportsPage)
        assert isinstance(window.pages["catalogs"], CatalogPage)
        assert isinstance(window.pages["system"], SystemPage)

        before_navigation = table_counts(db)
        for key in ("reports", "catalogs", "system"):
            window.navigate_to(key)
        assert table_counts(db) == before_navigation
        assert not list(tmp_path.glob("*.xlsx"))

        page = window.pages["catalogs"]
        page.school_year_dialog_factory = lambda *_: AcceptedDialog(
            (YEAR, date(2040, 9, 1), date(2041, 5, 31), False)
        )
        assert page.create_school_year() is True
        year_id = int(academic.get_school_year(YEAR)[0])

        page.grade_dialog_factory = lambda *_: AcceptedDialog(
            (GRADE_NUMBER, GRADE_NAME)
        )
        assert page.create_grade() is True
        grade_id = next(
            item.grade_id
            for item in academic.list_catalog_grades()
            if item.grade_number == GRADE_NUMBER and item.grade_name == GRADE_NAME
        )

        page.class_dialog_factory = lambda *_: AcceptedDialog(
            (CLASS_NAME, grade_id, year_id, f"{PREFIX} Homeroom", "ACTIVE")
        )
        assert page.create_class() is True
        class_id = next(
            item.class_id
            for item in academic.list_catalog_classes(year_id, grade_id)
            if item.class_name == CLASS_NAME
        )

        page.subject_dialog_factory = lambda *_: AcceptedDialog(
            (SUBJECT_CODE, SUBJECT_NAME, True)
        )
        assert page.create_subject() is True
        subject_id = int(academic.get_subject(SUBJECT_CODE)[0])

        page.assessment_dialog_factory = lambda *_: AcceptedDialog(
            (
                subject_id,
                year_id,
                ASSESSMENT,
                1,
                "QUIZ",
                date(2040, 10, 1),
                AssessmentStatus.ACTIVE,
            )
        )
        assert page.create_assessment() is True
        trigger_assessment = next(
            item
            for item in academic.list_assessments(year_id, subject_id)
            if item.assessment_name == ASSESSMENT
        )

        page.support_rule_dialog_factory = lambda *_: AcceptedDialog(
            (subject_id, year_id, "5.00", True)
        )
        assert page.create_support_rule() is True
        rule = next(
            item
            for item in academic.list_catalog_support_rules(year_id, subject_id)
            if item.is_active
        )

        academic.update_school_year(
            year_id, YEAR, date(2040, 8, 25), date(2041, 5, 25), False
        )
        academic.update_grade_name(grade_id, GRADE_UPDATED)
        academic.update_class(
            class_id,
            CLASS_UPDATED,
            grade_id,
            year_id,
            f"{PREFIX} Homeroom Updated",
            "ACTIVE",
        )
        academic.update_subject(subject_id, SUBJECT_CODE, SUBJECT_UPDATED, True)
        academic.update_assessment(
            trigger_assessment.assessment_id,
            subject_id,
            year_id,
            ASSESSMENT_UPDATED,
            1,
            "MIDTERM",
            date(2040, 10, 2),
            AssessmentStatus.ACTIVE,
        )
        academic.update_support_rule_threshold(rule.rule_id, Decimal("4.75"))
        review_assessment = academic.create_assessment(
            subject_id,
            year_id,
            REVIEW_ASSESSMENT,
            1,
            "REVIEW",
            date(2040, 11, 1),
        )

        student_repo = StudentRepository()
        enrollment_repo = EnrollmentRepository()
        score_repo = ScoreRepository()
        intervention_repo = InterventionRepository()
        with db.transaction() as connection:
            enrollments = []
            for index, (student_id, student_code) in enumerate(
                zip(STUDENT_IDS, STUDENT_CODES), start=1
            ):
                student_repo.create(
                    connection,
                    student_id,
                    StudentCreateData(student_code, f"{PREFIX} Student {index}"),
                )
                enrollments.append(
                    enrollment_repo.create(
                        connection, student_id, class_id, date(2040, 9, 1)
                    )
                )

            completed_trigger = score_repo.create(
                connection,
                enrollments[0].enrollment_id,
                trigger_assessment.assessment_id,
                Decimal("2.80"),
            )
            completed = intervention_repo.create(
                connection,
                enrollments[0].enrollment_id,
                subject_id,
                completed_trigger.score_id,
                date(2040, 10, 2),
            )
            review_score = score_repo.create(
                connection,
                enrollments[0].enrollment_id,
                review_assessment.assessment_id,
                Decimal("4.20"),
            )
            connection.cursor().execute(
                """
                INSERT INTO dbo.INTERVENTION_REVIEWS
                    (intervention_id, score_id, review_date, result, notes, created_at)
                VALUES (?, ?, ?, ?, ?, GETDATE())
                """,
                completed.intervention_id,
                review_score.score_id,
                date(2040, 11, 1),
                "PASSED",
                f"{PREFIX} historical review",
            )
            connection.cursor().execute(
                """
                UPDATE dbo.INTERVENTIONS
                SET status = 'COMPLETED', updated_at = GETDATE()
                WHERE intervention_id = ?
                """,
                completed.intervention_id,
            )

            detected_trigger = score_repo.create(
                connection,
                enrollments[1].enrollment_id,
                trigger_assessment.assessment_id,
                Decimal("2.50"),
            )
            detected = intervention_repo.create(
                connection,
                enrollments[1].enrollment_id,
                subject_id,
                detected_trigger.score_id,
                date(2040, 10, 2),
            )

        window.navigate_to("catalogs")
        select(page.class_year_combo, year_id)
        select(page.class_grade_combo, grade_id)
        assert any(item.class_name == CLASS_UPDATED for item in page.classes)
        assert any(item.subject_name == SUBJECT_UPDATED for item in page.subjects)
        select(page.assessment_year_combo, year_id)
        select(page.assessment_subject_combo, subject_id)
        assert any(item.assessment_name == ASSESSMENT_UPDATED for item in page.assessments)
        select(page.rule_year_combo, year_id)
        select(page.rule_subject_combo, subject_id)
        assert any(item.threshold == Decimal("4.75") for item in page.support_rules)

        window.navigate_to("reports")
        reports_page = window.pages["reports"]
        select(reports_page.school_year_combo, year_id)
        select(reports_page.grade_combo, grade_id)
        select(reports_page.class_combo, class_id)
        select(reports_page.subject_combo, subject_id)
        select(reports_page.status_combo, InterventionStatus.DETECTED)
        assert reports_page.refresh_report() is True
        assert reports_page.load_state == ReportsPage.STATE_READY
        assert reports_page.report_data.summary.total_cases == 1
        assert reports_page.report_data.summary.detected_count == 1
        assert reports_page.table.rowCount() == 1
        assert reports_page.intervention_id_at_row(0) == detected.intervention_id
        assert sum(bar.get_height() for bar in reports_page.status_chart.axes.patches) == 1
        assert sum(bar.get_height() for bar in reports_page.subject_chart.axes.patches) == 1

        snapshot = context.report_service.get_support_report(
            year_id,
            grade_id=grade_id,
            class_id=class_id,
            subject_id=subject_id,
        )
        assert snapshot.summary.total_cases == len(snapshot.rows) == 2
        export_data = SupportReportExportData(
            context=SupportReportExportContext(
                school_year_id=year_id,
                school_year_name=YEAR,
                grade_id=grade_id,
                grade_name=GRADE_UPDATED,
                class_id=class_id,
                class_name=CLASS_UPDATED,
                subject_id=subject_id,
                subject_name=SUBJECT_UPDATED,
            ),
            report=snapshot,
        )
        ReportExportService().export_xlsx(export_data, output)
        workbook = load_workbook(output, data_only=True)
        try:
            sheet = workbook["Báo cáo bổ trợ"]
            assert sheet["B3"].value == YEAR
            assert sheet["E3"].value == GRADE_UPDATED
            assert sheet["H3"].value == CLASS_UPDATED
            assert sheet["K3"].value == SUBJECT_UPDATED
            assert [sheet.cell(row, 2).value for row in range(7, 14)] == [
                2, 1, 0, 0, 0, 0, 1
            ]
            assert len([sheet.cell(15, column).value for column in range(1, 14)]) == 13
            exported = {
                sheet.cell(row, 2).value: [
                    sheet.cell(row, column).value for column in range(1, 14)
                ]
                for row in range(16, sheet.max_row + 1)
            }
            assert exported[STUDENT_CODES[1]][9] == status_label("DETECTED")
            assert exported[STUDENT_CODES[1]][10:13] == ["-", "-", "-"]
            assert exported[STUDENT_CODES[0]][12] == review_result_label("PASSED")
        finally:
            workbook.close()
        output.unlink()

        academic.set_assessment_active(trigger_assessment.assessment_id, False)
        assert trigger_assessment.assessment_id not in {
            item.assessment_id
            for item in academic.list_active_assessments(year_id, subject_id)
        }
        assert trigger_assessment.assessment_id in {
            item.assessment_id
            for item in academic.list_assessments(year_id, subject_id)
        }
        academic.set_class_active(class_id, False)
        class_filter = SupportFilterWidget(academic)
        class_filter.load_options()
        select(class_filter.school_year_combo, year_id)
        assert class_filter.class_combo.findData(class_id) == -1
        academic.set_subject_active(subject_id, False)
        assert subject_id not in {item[0] for item in academic.list_active_subjects()}
        academic.update_support_rule_threshold(rule.rule_id, Decimal("4.25"))
        academic.set_support_rule_active(rule.rule_id, False)

        with db.transaction() as connection:
            assert enrollment_repo.get_by_id(
                connection, enrollments[0].enrollment_id
            ) is not None
            assert score_repo.get_by_id(connection, completed_trigger.score_id) is not None
            preserved = intervention_repo.get_by_id(
                connection, completed.intervention_id
            )
            assert preserved is not None
            review_count = connection.cursor().execute(
                "SELECT COUNT(*) FROM dbo.INTERVENTION_REVIEWS WHERE intervention_id = ?",
                completed.intervention_id,
            ).fetchone()[0]
            assert review_count == 1
        assert context.report_service.get_support_report(year_id).summary.total_cases == 2

        academic.set_class_active(class_id, True)
        academic.set_subject_active(subject_id, True)
        academic.set_assessment_active(trigger_assessment.assessment_id, True)
        academic.set_support_rule_active(rule.rule_id, True)
        assert class_id in {
            item[0] for item in academic.list_classes_by_school_year(year_id)
            if item[-1]
        }
        assert subject_id in {item[0] for item in academic.list_active_subjects()}
    finally:
        if output.exists():
            output.unlink()
        if window is not None:
            window.close()
        cleanup(db)


def test_admin_teacher_system_profile_password_permissions_and_relogin_e2e():
    qapp()
    db = get_test_db()
    assert_connected_test_database(db)
    cleanup(db)
    users = UserService(db)
    auth = AuthService(db)
    admin_window = None
    teacher_window = None

    try:
        admin_user = users.create_user(
            ADMIN_USERNAME,
            OLD_PASSWORD,
            f"{PREFIX} Admin",
            UserRole.ADMIN,
        )
        admin = auth.login(ADMIN_USERNAME, OLD_PASSWORD)
        admin_context = build_context(db, admin)
        admin_window = MainWindow(admin_context)
        admin_window.navigate_to("system")
        admin_page = admin_window.pages["system"]
        assert admin_page.tabs.count() == 2
        assert admin_page.own_profile.user_id == admin_user.user_id
        assert all(not hasattr(item, "password_hash") for item in admin_page.users)

        admin_page.user_dialog_factory = lambda *_: AcceptedUserDialog(
            create_values=(
                TEACHER_USERNAME,
                OLD_PASSWORD,
                f"{PREFIX} Teacher",
                UserRole.TEACHER,
                f"{PREFIX.lower()}@example.com",
                "0901210123",
                True,
            )
        )
        assert admin_page.create_user() is True
        teacher_user = users.get_by_username(TEACHER_USERNAME)
        assert teacher_user is not None
        select_table_identity(admin_page.user_table, teacher_user.user_id)
        admin_page.user_dialog_factory = lambda *_: AcceptedUserDialog(
            update_values=(
                f"{PREFIX} Teacher Updated",
                UserRole.TEACHER,
                f"{PREFIX.lower()}.updated@example.com",
                "0901210456",
            )
        )
        assert admin_page.edit_user() is True
        updated_teacher = users.get_by_username(TEACHER_USERNAME)
        assert updated_teacher.username == TEACHER_USERNAME
        assert updated_teacher.role == UserRole.TEACHER
        select_table_identity(admin_page.user_table, teacher_user.user_id)
        assert admin_page.toggle_user_active() is True
        with pytest.raises(ValidationError):
            auth.login(TEACHER_USERNAME, OLD_PASSWORD)
        select_table_identity(admin_page.user_table, teacher_user.user_id)
        assert admin_page.toggle_user_active() is True
        assert auth.login(TEACHER_USERNAME, OLD_PASSWORD).user_id == teacher_user.user_id
        assert admin_page.initialize_system() is True
        assert TEACHER_USERNAME in {item.username for item in admin_page.users}
        assert all(not hasattr(item, "password_hash") for item in admin_page.users)

        own = users.update_own_profile(
            admin,
            f"{PREFIX} Admin Updated",
            f"{PREFIX.lower()}.admin@example.com",
            "0901210789",
        )
        assert (own.username, own.role, own.is_active) == (
            ADMIN_USERNAME,
            UserRole.ADMIN,
            True,
        )
        with pytest.raises(BusinessRuleError):
            users.admin_set_user_active(admin, admin.user_id, False)
        with pytest.raises(BusinessRuleError):
            users.admin_update_user(
                admin,
                admin.user_id,
                own.full_name,
                UserRole.TEACHER,
                own.email,
                own.phone,
            )

        users.change_own_password(admin, OLD_PASSWORD, NEW_PASSWORD)
        persisted_admin = users.get_by_username(ADMIN_USERNAME)
        assert persisted_admin.password_hash not in {OLD_PASSWORD, NEW_PASSWORD}
        assert bcrypt.checkpw(
            NEW_PASSWORD.encode(), persisted_admin.password_hash.encode()
        )
        admin_window.request_logout()
        assert admin_context.session is None
        with pytest.raises(ValidationError):
            auth.login(ADMIN_USERNAME, OLD_PASSWORD)
        relogged_admin = auth.login(ADMIN_USERNAME, NEW_PASSWORD)
        assert relogged_admin.user_id == admin_user.user_id

        teacher = auth.login(TEACHER_USERNAME, OLD_PASSWORD)
        teacher_context = build_context(db, teacher)
        teacher_window = MainWindow(teacher_context)
        assert teacher_window.can_navigate_to("reports") is True
        assert teacher_window.can_navigate_to("system") is True
        assert teacher_window.can_navigate_to("catalogs") is False
        with pytest.raises(PermissionError):
            teacher_window.navigate_to("catalogs")
        teacher_window.navigate_to("reports")
        teacher_window.navigate_to("system")
        teacher_page = teacher_window.pages["system"]
        assert teacher_page.tabs.count() == 1
        assert teacher_page.tabs.tabText(0) == "Tài khoản của tôi"
        assert teacher_page.own_profile.user_id == teacher_user.user_id

        teacher_profile = users.update_own_profile(
            teacher,
            f"{PREFIX} Teacher Self",
            f"{PREFIX.lower()}.self@example.com",
            "0901210999",
        )
        assert (teacher_profile.username, teacher_profile.role, teacher_profile.is_active) == (
            TEACHER_USERNAME,
            UserRole.TEACHER,
            True,
        )
        with pytest.raises(ValidationError):
            users.admin_list_users(teacher)
        users.change_own_password(teacher, OLD_PASSWORD, NEW_PASSWORD)
        teacher_window.request_logout()
        assert teacher_context.session is None
        with pytest.raises(ValidationError):
            auth.login(TEACHER_USERNAME, OLD_PASSWORD)
        relogged_teacher = auth.login(TEACHER_USERNAME, NEW_PASSWORD)
        assert relogged_teacher.user_id == teacher_user.user_id

        users.admin_set_user_active(relogged_admin, teacher_user.user_id, False)
        with pytest.raises(ValidationError):
            auth.login(TEACHER_USERNAME, NEW_PASSWORD)
        users.admin_set_user_active(relogged_admin, teacher_user.user_id, True)
        assert auth.login(TEACHER_USERNAME, NEW_PASSWORD).user_id == teacher_user.user_id
    finally:
        if teacher_window is not None:
            teacher_window.close()
        if admin_window is not None:
            admin_window.close()
        cleanup(db)


class FailingAfterProfileUpdateRepository(UserRepository):
    def update_profile(self, connection, user_id, full_name, email, phone):
        super().update_profile(connection, user_id, full_name, email, phone)
        raise pyodbc.Error("forced repository failure")


def test_user_repository_failure_rolls_back_and_exposes_only_app_error():
    qapp()
    db = get_test_db()
    assert_connected_test_database(db)
    cleanup(db)
    users = UserService(db)

    try:
        users.create_user(
            ROLLBACK_USERNAME,
            OLD_PASSWORD,
            f"{PREFIX} Before Rollback",
            UserRole.TEACHER,
        )
        actor = AuthService(db).login(ROLLBACK_USERNAME, OLD_PASSWORD)
        failing_service = UserService(db, FailingAfterProfileUpdateRepository())

        with pytest.raises(DatabaseError) as exc_info:
            failing_service.update_own_profile(
                actor,
                f"{PREFIX} Must Roll Back",
                f"{PREFIX.lower()}.rollback@example.com",
                "0901210000",
            )

        assert "pyodbc" not in str(exc_info.value).lower()
        persisted = users.get_own_profile(actor)
        assert persisted.full_name == f"{PREFIX} Before Rollback"
        assert persisted.email is None
        assert persisted.phone is None

        page = SystemPage(
            user_service=failing_service,
            session=actor,
            permission_service=PermissionService(),
            profile_dialog_factory=lambda *_: AcceptedDialog(
                (
                    f"{PREFIX} UI Must Roll Back",
                    f"{PREFIX.lower()}.ui@example.com",
                    "0901210001",
                )
            ),
        )
        assert page.initialize_system() is True
        assert page.edit_own_profile() is False
        assert "pyodbc" not in page.state_label.text().lower()
        assert users.get_own_profile(actor).full_name == f"{PREFIX} Before Rollback"
    finally:
        cleanup(db)
