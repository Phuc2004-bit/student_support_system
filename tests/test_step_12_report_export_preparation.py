from dataclasses import FrozenInstanceError, replace
from datetime import date, datetime
from decimal import Decimal
import inspect

import pytest
from openpyxl import load_workbook

from exceptions import ReportExportError, ValidationError
from models.dto.report_dto import (
    SupportReportData,
    SupportReportRow,
    SupportReportSummary,
)
from models.dto.report_export import (
    SupportReportExportContext,
    SupportReportExportData,
)
from models.enums import InterventionStatus
from services.report_export_contract import ReportExportServiceContract
from services.report_export_service import ReportExportService
from ui.pages.reports_page import ReportsPage


def context(**changes):
    value = SupportReportExportContext(
        school_year_id=2,
        school_year_name="2026-2027",
        grade_id=6,
        grade_name="Khối 6",
        class_id=61,
        class_name="6A1",
        subject_id=99,
        subject_name="Robotics",
        status="CONTINUE",
    )
    return replace(value, **changes)


def summary():
    return SupportReportSummary(
        total_cases=21,
        detected_count=1,
        planned_count=2,
        in_progress_count=3,
        waiting_review_count=4,
        continue_count=5,
        completed_count=6,
    )


def row(
    *,
    status="CONTINUE",
    latest_review_date=date(2026, 11, 20),
    latest_review_score=Decimal("3.20"),
    latest_review_result="NOT_PASSED",
    subject_name="Robotics",
):
    return SupportReportRow(
        intervention_id=101,
        student_code="HS001",
        full_name="Nguyễn Văn A",
        grade_number=6,
        class_name="6A1",
        subject_code="ROBOT",
        subject_name=subject_name,
        detected_date=date(2026, 10, 10),
        start_date=date(2026, 10, 12),
        status=status,
        trigger_score=Decimal("2.80"),
        latest_review_date=latest_review_date,
        latest_review_score=latest_review_score,
        latest_review_result=latest_review_result,
        subject_id=99,
    )


def export_data(*, report_rows=None, export_context=None):
    rows = (row(),) if report_rows is None else tuple(report_rows)
    return SupportReportExportData(
        context=export_context or context(),
        report=SupportReportData(summary(), rows),
    )


def export_and_load(tmp_path, data=None):
    path = tmp_path / "support-report.xlsx"
    result = ReportExportService().export_xlsx(
        data or export_data(),
        path,
    )
    return result, load_workbook(path)


def as_date(value):
    return value.date() if isinstance(value, datetime) else value


def test_export_service_matches_contract_and_returns_output_path(tmp_path):
    service: ReportExportServiceContract = ReportExportService()
    path = tmp_path / "report.xlsx"

    result = service.export_xlsx(export_data(), path)

    assert result == path
    assert path.is_file()


def test_export_contract_reuses_exact_support_report_data_snapshot():
    report = SupportReportData(summary(), (row(),))
    value = SupportReportExportData(context(), report)

    assert value.report is report
    assert value.report.rows[0].student_code == "HS001"


def test_export_context_is_immutable_and_preserves_filter_ids_and_labels():
    value = context()

    assert value.school_year_id == 2
    assert value.grade_id == 6
    assert value.class_id == 61
    assert value.subject_id == 99
    assert value.school_year_name == "2026-2027"
    assert value.status == "CONTINUE"
    with pytest.raises(FrozenInstanceError):
        value.school_year_id = 3


def test_workbook_has_expected_sheet_title_sections_and_table(tmp_path):
    _, workbook = export_and_load(tmp_path)
    try:
        worksheet = workbook["Báo cáo bổ trợ"]
        assert worksheet["A1"].value == "BÁO CÁO BỔ TRỢ HỌC TẬP"
        assert worksheet["A6"].value == "Tổng quan"
        assert worksheet["A15"].value == "STT"
        assert worksheet["M15"].value == "Kết quả đánh giá gần nhất"
        assert worksheet.freeze_panes == "A16"
        assert worksheet.auto_filter.ref == "A15:M16"
    finally:
        workbook.close()


def test_workbook_exports_all_context_labels(tmp_path):
    _, workbook = export_and_load(tmp_path)
    try:
        worksheet = workbook.active
        assert worksheet["B3"].value == "2026-2027"
        assert worksheet["E3"].value == "Khối 6"
        assert worksheet["H3"].value == "6A1"
        assert worksheet["K3"].value == "Robotics"
        assert worksheet["B4"].value == "Cần tiếp tục"
    finally:
        workbook.close()


def test_optional_context_uses_explicit_all_label(tmp_path):
    all_context = context(
        grade_id=None,
        grade_name=None,
        class_id=None,
        class_name=None,
        subject_id=None,
        subject_name=None,
        status=None,
    )
    _, workbook = export_and_load(
        tmp_path,
        export_data(export_context=all_context),
    )
    try:
        worksheet = workbook.active
        assert worksheet["E3"].value == "Tất cả"
        assert worksheet["H3"].value == "Tất cả"
        assert worksheet["K3"].value == "Tất cả"
        assert worksheet["B4"].value == "Tất cả"
    finally:
        workbook.close()


def test_context_accepts_filter_status_enum_without_ui_conversion(tmp_path):
    enum_context = context(status=InterventionStatus.COMPLETED)

    _, workbook = export_and_load(
        tmp_path,
        export_data(export_context=enum_context),
    )
    try:
        assert workbook.active["B4"].value == "Đã đạt ngưỡng"
    finally:
        workbook.close()


def test_summary_values_are_exported_without_recalculation(tmp_path):
    _, workbook = export_and_load(tmp_path)
    try:
        worksheet = workbook.active
        assert [worksheet.cell(index, 1).value for index in range(7, 14)] == [
            "Tổng số ca",
            "Mới phát hiện",
            "Đã lập kế hoạch",
            "Đang bổ trợ",
            "Chờ đánh giá",
            "Cần tiếp tục",
            "Đã đạt ngưỡng",
        ]
        assert [worksheet.cell(index, 2).value for index in range(7, 14)] == [
            21, 1, 2, 3, 4, 5, 6,
        ]
    finally:
        workbook.close()


def test_report_row_and_school_year_are_exported_from_snapshot_and_context(
    tmp_path,
):
    _, workbook = export_and_load(tmp_path)
    try:
        worksheet = workbook.active
        assert [worksheet.cell(16, column).value for column in range(1, 8)] == [
            1,
            "HS001",
            "Nguyễn Văn A",
            6,
            "6A1",
            "2026-2027",
            "Robotics",
        ]
        assert worksheet["H16"].value == 2.8
        assert as_date(worksheet["I16"].value) == date(2026, 10, 10)
        assert worksheet["J16"].value == "Cần tiếp tục"
        assert as_date(worksheet["K16"].value) == date(2026, 11, 20)
        assert worksheet["L16"].value == 3.2
        assert worksheet["M16"].value == "Chưa đạt"
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("status", "expected"),
    [
        ("DETECTED", "Mới phát hiện"),
        ("PLANNED", "Đã lập kế hoạch"),
        ("IN_PROGRESS", "Đang bổ trợ"),
        ("WAITING_REVIEW", "Chờ đánh giá"),
        ("CONTINUE", "Cần tiếp tục"),
        ("COMPLETED", "Đã đạt ngưỡng"),
    ],
)
def test_export_uses_shared_vietnamese_status_labels(
    tmp_path,
    status,
    expected,
):
    _, workbook = export_and_load(
        tmp_path,
        export_data(report_rows=(row(status=status),)),
    )
    try:
        assert workbook.active["J16"].value == expected
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("result", "expected"),
    [("PASSED", "Đạt"), ("NOT_PASSED", "Chưa đạt")],
)
def test_export_uses_shared_review_result_labels(
    tmp_path,
    result,
    expected,
):
    _, workbook = export_and_load(
        tmp_path,
        export_data(report_rows=(row(latest_review_result=result),)),
    )
    try:
        assert workbook.active["M16"].value == expected
    finally:
        workbook.close()


def test_null_review_values_are_exported_as_hyphens(tmp_path):
    no_review = row(
        latest_review_date=None,
        latest_review_score=None,
        latest_review_result=None,
    )
    _, workbook = export_and_load(
        tmp_path,
        export_data(report_rows=(no_review,)),
    )
    try:
        assert [workbook.active.cell(16, column).value for column in (11, 12, 13)] == [
            "-", "-", "-",
        ]
    finally:
        workbook.close()


def test_empty_report_still_exports_summary_and_headers(tmp_path):
    _, workbook = export_and_load(
        tmp_path,
        export_data(report_rows=()),
    )
    try:
        worksheet = workbook.active
        assert worksheet.max_row == 15
        assert worksheet["A15"].value == "STT"
        assert worksheet["B7"].value == 21
        assert worksheet.auto_filter.ref == "A15:M15"
    finally:
        workbook.close()


@pytest.mark.parametrize(
    "invalid_path",
    ["report.csv", "report", ""],
)
def test_invalid_output_extension_is_rejected(tmp_path, invalid_path):
    with pytest.raises(ValidationError):
        ReportExportService().export_xlsx(
            export_data(),
            tmp_path / invalid_path,
        )


def test_missing_output_directory_is_rejected_safely(tmp_path):
    path = tmp_path / "missing" / "report.xlsx"

    with pytest.raises(ValidationError, match="không tồn tại"):
        ReportExportService().export_xlsx(export_data(), path)

    assert not path.exists()


def test_directory_named_xlsx_is_rejected_safely(tmp_path):
    path = tmp_path / "folder.xlsx"
    path.mkdir()

    with pytest.raises(ValidationError, match="phải là một file"):
        ReportExportService().export_xlsx(export_data(), path)


def test_permission_failure_is_normalized(monkeypatch, tmp_path):
    def deny_write(_workbook, _path):
        raise PermissionError("C:/secret/raw/access denied")

    monkeypatch.setattr(
        "services.report_export_service.Workbook.save",
        deny_write,
    )

    with pytest.raises(ReportExportError) as error:
        ReportExportService().export_xlsx(
            export_data(),
            tmp_path / "report.xlsx",
        )

    assert "secret" not in str(error.value).lower()
    assert "access denied" not in str(error.value).lower()
    assert "quyền truy cập" in str(error.value)


def test_generic_write_failure_is_normalized(monkeypatch, tmp_path):
    def fail_write(_workbook, _path):
        raise RuntimeError("raw writer internals")

    monkeypatch.setattr(
        "services.report_export_service.Workbook.save",
        fail_write,
    )

    with pytest.raises(ReportExportError) as error:
        ReportExportService().export_xlsx(
            export_data(),
            tmp_path / "report.xlsx",
        )

    assert "raw writer" not in str(error.value).lower()


@pytest.mark.parametrize(
    "bad_context",
    [
        context(school_year_id=0),
        context(school_year_id=True),
        context(school_year_name="  "),
        context(grade_id=-1),
        context(class_id=True),
        context(subject_id=0),
        context(status="UNKNOWN"),
    ],
)
def test_invalid_export_context_is_rejected_before_file_write(
    tmp_path,
    bad_context,
):
    path = tmp_path / "report.xlsx"

    with pytest.raises(ValidationError):
        ReportExportService().export_xlsx(
            export_data(export_context=bad_context),
            path,
        )

    assert not path.exists()


@pytest.mark.parametrize("invalid_data", [None, object()])
def test_invalid_export_data_contract_is_rejected(tmp_path, invalid_data):
    with pytest.raises(ValidationError):
        ReportExportService().export_xlsx(
            invalid_data,
            tmp_path / "report.xlsx",
        )


def test_exporter_has_no_database_repository_or_transaction_dependency():
    source = inspect.getsource(ReportExportService).upper()
    for forbidden in (
        "DATABASE",
        "REPOSITORY",
        "TRANSACTION",
        "COMMIT",
        "ROLLBACK",
        "SELECT ",
        "INSERT ",
        "UPDATE ",
        "DELETE ",
    ):
        assert forbidden not in source


def test_exporter_does_not_query_per_row_or_mutate_report_snapshot(tmp_path):
    report_rows = tuple(
        replace(row(), intervention_id=index, student_code=f"HS{index}")
        for index in range(1, 51)
    )
    value = export_data(report_rows=report_rows)

    ReportExportService().export_xlsx(
        value,
        tmp_path / "many.xlsx",
    )

    assert value.report.rows is report_rows
    assert len(value.report.rows) == 50


def test_exporter_has_no_hard_coded_subject_catalog():
    source = inspect.getsource(ReportExportService).upper()

    assert "TOAN" not in source
    assert "VAN" not in source
    assert "ROBOTICS" not in source


def test_report_page_is_not_integrated_with_export_preparation():
    source = inspect.getsource(ReportsPage).upper()

    assert "REPORTEXPORTSERVICE" not in source
    assert "EXPORT_XLSX" not in source
