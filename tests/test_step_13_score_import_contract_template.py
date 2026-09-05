from __future__ import annotations

from dataclasses import FrozenInstanceError
from datetime import date
from decimal import Decimal
import inspect

from openpyxl import Workbook, load_workbook
import pytest

from exceptions import AppError, ScoreImportError
from models.dto.score_import import (
    ScoreImportContext,
    ScoreImportRow,
    ScoreImportTemplateStudent,
    ScoreImportWorkbook,
)
from services.score_import_contract import (
    ScoreImportParserContract,
    ScoreImportTemplateServiceContract,
)
from services.score_import_parser import ScoreImportWorkbookParser
from services.score_import_template_service import ScoreImportTemplateService


def context(**changes) -> ScoreImportContext:
    values = {
        "school_year_id": 2,
        "school_year_name": "2026-2027",
        "class_id": 61,
        "class_name": "6A1",
        "subject_id": 9,
        "subject_name": "Khoa học",
        "assessment_id": 101,
        "assessment_name": "Giữa kỳ I",
    }
    values.update(changes)
    return ScoreImportContext(**values)


def roster():
    return (
        ScoreImportTemplateStudent(" student-b ", " Trần Bình "),
        ScoreImportTemplateStudent("student-a", "An Nguyễn"),
    )


def make_workbook(path, *, sheet="Nhập điểm", headers=None, rows=()):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row, value in enumerate(
        ("2026-2027", "6A1", "Khoa học", "Giữa kỳ I"), start=1
    ):
        worksheet.cell(row, 2, value)
    for column, value in enumerate(
        headers or ("STT", "Mã học sinh", "Họ và tên", "Điểm"), start=1
    ):
        worksheet.cell(6, column, value)
    for row_number, values in enumerate(rows, start=7):
        for column, value in enumerate(values, start=1):
            worksheet.cell(row_number, column, value)
    workbook.save(path)
    workbook.close()
    return path


def test_services_implement_read_only_import_contracts():
    assert isinstance(ScoreImportWorkbookParser(), ScoreImportParserContract)
    assert isinstance(
        ScoreImportTemplateService(), ScoreImportTemplateServiceContract
    )
    assert not hasattr(ScoreImportWorkbookParser(), "commit")
    assert not hasattr(ScoreImportTemplateService(), "save_scores")


def test_import_read_models_are_immutable_and_database_free():
    row = ScoreImportRow(7, "student-a", "An", "8.5", Decimal("8.5"))
    workbook = ScoreImportWorkbook("scores.xlsx", "Nhập điểm", None, (row,))
    with pytest.raises(FrozenInstanceError):
        row.student_id = "changed"
    assert not hasattr(row, "connection")
    assert not hasattr(workbook, "repository")


def test_template_has_metadata_headers_roster_and_blank_score_cells(tmp_path):
    path = tmp_path / "score-template.xlsx"
    result = ScoreImportTemplateService().create_template(context(), roster(), path)
    assert result == path

    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames == ["Nhập điểm"]
        worksheet = workbook["Nhập điểm"]
        assert [worksheet.cell(6, column).value for column in range(1, 5)] == [
            "STT",
            "Mã học sinh",
            "Họ và tên",
            "Điểm",
        ]
        assert [worksheet.cell(row, 1).value for row in range(7, 9)] == [1, 2]
        assert [worksheet.cell(row, 2).value for row in range(7, 9)] == [
            "student-a",
            "student-b",
        ]
        assert [worksheet.cell(row, 3).value for row in range(7, 9)] == [
            "An Nguyễn",
            "Trần Bình",
        ]
        assert [worksheet.cell(row, 4).value for row in range(7, 9)] == [None, None]
        assert [worksheet.cell(row, 2).value for row in range(1, 5)] == [
            "2026-2027",
            "6A1",
            "Khoa học",
            "Giữa kỳ I",
        ]
        assert worksheet.protection.sheet is True
        assert worksheet["D7"].protection.locked is False
    finally:
        workbook.close()


def test_template_order_is_deterministic_and_text_is_not_a_formula(tmp_path):
    path = tmp_path / "formula-safe.xlsx"
    students = (
        ScoreImportTemplateStudent("=2+2", "=HYPERLINK(\"x\")"),
        ScoreImportTemplateStudent("plain", "An"),
    )
    ScoreImportTemplateService().create_template(context(), students, path)
    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook["Nhập điểm"]
        assert worksheet["B7"].value == "=2+2"
        assert worksheet["B7"].data_type == "s"
        assert worksheet["C7"].data_type == "s"
    finally:
        workbook.close()


@pytest.mark.parametrize(
    "students",
    [
        None,
        (object(),),
        (ScoreImportTemplateStudent("", "Name"),),
        (
            ScoreImportTemplateStudent("same", "A"),
            ScoreImportTemplateStudent("same", "B"),
        ),
    ],
)
def test_template_rejects_invalid_or_duplicate_roster(tmp_path, students):
    with pytest.raises(ScoreImportError):
        ScoreImportTemplateService().create_template(
            context(), students, tmp_path / "invalid.xlsx"
        )


@pytest.mark.parametrize(
    "bad_context",
    [
        None,
        context(school_year_id=0),
        context(class_id=True),
        context(subject_name="   "),
        context(assessment_name=None),
    ],
)
def test_template_rejects_invalid_context_before_file_write(tmp_path, bad_context):
    output = tmp_path / "invalid-context.xlsx"
    with pytest.raises(ScoreImportError):
        ScoreImportTemplateService().create_template(bad_context, roster(), output)
    assert not output.exists()


@pytest.mark.parametrize("name", ("scores.csv", "scores", "scores.xls"))
def test_template_rejects_non_xlsx_output(tmp_path, name):
    with pytest.raises(ScoreImportError):
        ScoreImportTemplateService().create_template(
            context(), roster(), tmp_path / name
        )


def test_template_rejects_missing_output_directory(tmp_path):
    with pytest.raises(ScoreImportError):
        ScoreImportTemplateService().create_template(
            context(), roster(), tmp_path / "missing" / "scores.xlsx"
        )


def test_parser_reads_generated_template_and_source_metadata(tmp_path):
    path = tmp_path / "generated.xlsx"
    ScoreImportTemplateService().create_template(context(), roster(), path)
    parsed = ScoreImportWorkbookParser().parse_workbook(path)

    assert parsed.file_name == "generated.xlsx"
    assert parsed.sheet_name == "Nhập điểm"
    assert parsed.metadata.school_year_name == "2026-2027"
    assert parsed.metadata.class_name == "6A1"
    assert parsed.metadata.subject_name == "Khoa học"
    assert parsed.metadata.assessment_name == "Giữa kỳ I"
    assert [row.row_number for row in parsed.rows] == [7, 8]
    assert [row.student_id for row in parsed.rows] == ["student-a", "student-b"]
    assert all(row.normalized_score is None for row in parsed.rows)
    assert all(row.is_structurally_valid for row in parsed.rows)


def test_parser_trims_student_identity_and_informational_name(tmp_path):
    path = make_workbook(
        tmp_path / "trim.xlsx",
        rows=((1, "  student-01  ", "  Nguyễn An  ", " 8.50 "),),
    )
    row = ScoreImportWorkbookParser().parse_workbook(path).rows[0]
    assert row.row_number == 7
    assert row.student_id == "student-01"
    assert row.student_name == "Nguyễn An"
    assert row.raw_score == " 8.50 "
    assert row.normalized_score == Decimal("8.50")


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (8, Decimal("8")),
        (8.25, Decimal("8.25")),
        ("8.50", Decimal("8.50")),
        (" -1 ", Decimal("-1")),
        ("11", Decimal("11")),
    ],
)
def test_parser_normalizes_basic_numeric_cells_without_business_range_check(
    tmp_path, value, expected
):
    path = make_workbook(
        tmp_path / f"numeric-{str(value).strip()}.xlsx",
        rows=((1, "student-01", "Nguyễn An", value),),
    )
    row = ScoreImportWorkbookParser().parse_workbook(path).rows[0]
    assert row.raw_score == value
    assert row.normalized_score == expected
    assert row.issues == ()


def test_blank_score_is_explicitly_preserved_without_business_validation(tmp_path):
    path = make_workbook(
        tmp_path / "blank-score.xlsx",
        rows=((1, "student-01", "Nguyễn An", None),),
    )
    row = ScoreImportWorkbookParser().parse_workbook(path).rows[0]
    assert row.raw_score is None
    assert row.normalized_score is None
    assert row.issues == ()


def test_completely_blank_rows_are_ignored_but_excel_row_numbers_are_kept(tmp_path):
    path = make_workbook(
        tmp_path / "blank-row.xlsx",
        rows=(
            (1, "student-01", "An", "7"),
            (None, None, None, None),
            (3, "student-03", "Bình", "8"),
        ),
    )
    parsed = ScoreImportWorkbookParser().parse_workbook(path)
    assert [row.row_number for row in parsed.rows] == [7, 9]


def test_missing_student_id_is_a_structural_row_issue(tmp_path):
    path = make_workbook(
        tmp_path / "missing-id.xlsx",
        rows=((1, "  ", "Nguyễn An", "7"),),
    )
    row = ScoreImportWorkbookParser().parse_workbook(path).rows[0]
    assert row.student_id == ""
    assert row.is_structurally_valid is False
    assert [issue.code for issue in row.issues] == ["MISSING_STUDENT_ID"]


@pytest.mark.parametrize(
    ("value", "code"),
    [
        (True, "BOOLEAN_SCORE"),
        ("=1+1", "FORMULA_SCORE"),
        (date(2026, 9, 5), "UNSUPPORTED_SCORE_TYPE"),
        ("not-a-number", "INVALID_SCORE"),
        ("NaN", "INVALID_SCORE"),
    ],
)
def test_unsupported_score_cells_become_issues_without_crashing(
    tmp_path, value, code
):
    path = make_workbook(
        tmp_path / f"unsupported-{code}.xlsx",
        rows=((1, "student-01", "Nguyễn An", value),),
    )
    row = ScoreImportWorkbookParser().parse_workbook(path).rows[0]
    if isinstance(value, date):
        assert isinstance(row.raw_score, date)
    else:
        assert row.raw_score == value
    assert row.normalized_score is None
    assert code in {issue.code for issue in row.issues}


@pytest.mark.parametrize(
    ("student_id", "student_name", "code"),
    [
        (True, "Nguyễn An", "INVALID_STUDENT_ID_TYPE"),
        ("student-01", 123, "INVALID_STUDENT_NAME_TYPE"),
    ],
)
def test_unsupported_identity_cells_become_row_issues(
    tmp_path, student_id, student_name, code
):
    path = make_workbook(
        tmp_path / f"identity-{code}.xlsx",
        rows=((1, student_id, student_name, "7"),),
    )
    row = ScoreImportWorkbookParser().parse_workbook(path).rows[0]
    assert code in {issue.code for issue in row.issues}


@pytest.mark.parametrize(
    "headers",
    [
        (None, "Mã học sinh", "Họ và tên", "Điểm"),
        ("STT", None, "Họ và tên", "Điểm"),
        ("STT", "Mã học sinh", None, "Điểm"),
        ("STT", "Mã học sinh", "Họ và tên", None),
    ],
)
def test_missing_required_header_blocks_workbook(tmp_path, headers):
    path = make_workbook(tmp_path / "missing-header.xlsx", headers=headers)
    with pytest.raises(ScoreImportError, match="Thiếu cột bắt buộc"):
        ScoreImportWorkbookParser().parse_workbook(path)


def test_duplicate_semantic_header_blocks_workbook(tmp_path):
    path = make_workbook(
        tmp_path / "duplicate-header.xlsx",
        headers=("STT", "Mã học sinh", "student_id", "Họ và tên", "Điểm"),
    )
    with pytest.raises(ScoreImportError, match="trùng ý nghĩa"):
        ScoreImportWorkbookParser().parse_workbook(path)


def test_wrong_sheet_is_normalized_to_project_error(tmp_path):
    path = make_workbook(tmp_path / "wrong-sheet.xlsx", sheet="Other")
    with pytest.raises(ScoreImportError, match="Nhập điểm"):
        ScoreImportWorkbookParser().parse_workbook(path)


@pytest.mark.parametrize("name", ("scores.csv", "scores.xls", "scores"))
def test_parser_rejects_invalid_extension(tmp_path, name):
    path = tmp_path / name
    path.write_text("not excel", encoding="utf-8")
    with pytest.raises(ScoreImportError, match=".xlsx"):
        ScoreImportWorkbookParser().parse_workbook(path)


def test_parser_rejects_missing_file(tmp_path):
    with pytest.raises(ScoreImportError, match="Không tìm thấy"):
        ScoreImportWorkbookParser().parse_workbook(tmp_path / "missing.xlsx")


def test_corrupt_workbook_is_normalized_without_openpyxl_details(tmp_path):
    path = tmp_path / "corrupt.xlsx"
    path.write_bytes(b"not a zip workbook")
    with pytest.raises(ScoreImportError) as exc_info:
        ScoreImportWorkbookParser().parse_workbook(path)
    assert "openpyxl" not in str(exc_info.value).lower()
    assert isinstance(exc_info.value, AppError)


def test_unreadable_workbook_error_is_normalized(monkeypatch, tmp_path):
    path = tmp_path / "unreadable.xlsx"
    path.touch()

    def fail(*_args, **_kwargs):
        raise PermissionError("raw permission detail")

    monkeypatch.setattr("services.score_import_parser.load_workbook", fail)
    with pytest.raises(ScoreImportError) as exc_info:
        ScoreImportWorkbookParser().parse_workbook(path)
    assert "raw permission detail" not in str(exc_info.value)


def test_template_write_error_is_normalized(monkeypatch, tmp_path):
    def fail(*_args, **_kwargs):
        raise PermissionError("raw permission detail")

    monkeypatch.setattr(
        "services.score_import_template_service.Workbook.save", fail
    )
    with pytest.raises(ScoreImportError) as exc_info:
        ScoreImportTemplateService().create_template(
            context(), roster(), tmp_path / "blocked.xlsx"
        )
    assert "raw permission detail" not in str(exc_info.value)


def test_parser_and_template_have_no_database_or_repository_dependency():
    source = (
        inspect.getsource(inspect.getmodule(ScoreImportWorkbookParser))
        + inspect.getsource(inspect.getmodule(ScoreImportTemplateService))
    ).upper()
    for forbidden in (
        "DATABASE",
        "REPOSITORY",
        "SCORESERVICE",
        "SUPPORTSERVICE",
        "INSERT ",
        "UPDATE ",
        ".COMMIT(",
    ):
        assert forbidden not in source
