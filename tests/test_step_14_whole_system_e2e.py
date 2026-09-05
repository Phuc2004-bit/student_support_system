from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
import inspect
import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import bcrypt
from openpyxl import load_workbook
import pyodbc
import pytest
from PySide6.QtWidgets import QApplication, QDialog

from app_context import AppContext
from config.database import db_settings
from database.connection import DatabaseManager
from exceptions import (
    BusinessRuleError,
    DatabaseError,
    DuplicateError,
    InvalidStateTransitionError,
    MissingSupportRuleError,
    ValidationError,
)
from models.dto import StudentCreateData
from models.dto.data_export import ScoreExportContext, StudentExportContext
from models.dto.report_export import SupportReportExportContext, SupportReportExportData
from models.dto.score_import import ScoreImportContext, ScoreImportTemplateStudent
from models.dto.student_filter import StudentFilter
from models.enums import InterventionStatus, ReviewResult, UserRole
from repositories import UserRepository
from services.academic_service import AcademicService
from services.auth_service import AuthService
from services.data_export_service import ScoreExportService, StudentExportService
from services.enrollment_service import EnrollmentService
from services.permission_service import PermissionService
from services.report_export_service import ReportExportService
from services.report_service import ReportService
from services.score_import_commit_service import ScoreImportCommitService
from services.score_import_parser import ScoreImportWorkbookParser
from services.score_import_preview_service import ScoreImportPreviewService
from services.score_import_template_service import ScoreImportTemplateService
from services.score_service import ScoreService
from services.student_list_service import StudentListService
from services.student_profile_service import StudentProfileService
from services.student_service import StudentService
from services.support_service import SupportService
from services.user_service import UserService
from ui.main_window import MainWindow


PREFIX = "T14"
YEAR_NAME = "T14_2627"
CLASS_NAME = "-T14_7A"
SUBJECT_CODE = "T14_M1"
SUBJECT_NAME = "@T14 Môn tích hợp"
NO_RULE_CODE = "T14_NR"
ADMIN_USERNAME = "t14_admin"
TEACHER_USERNAME = "t14_teacher"
ROLLBACK_USERNAME = "t14_rollback"
OLD_PASSWORD = "T14OldPassword!"
NEW_PASSWORD = "T14NewPassword!"
THRESHOLD = Decimal("6.25")
STUDENT_CODES = ("T1401", "T1402", "T1499")
ASSESSMENT_NAMES = (
    "T14 Trigger",
    "T14 Review 1",
    "T14 Review 2",
    "T14 New episode",
    "T14 Ordinary",
    "T14 Excel",
    "T14 No rule",
)


def qapp() -> QApplication:
    return QApplication.instance() or QApplication([])


def step14_db() -> DatabaseManager:
    configured = db_settings.connection_string()
    target = configured.replace(
        "DATABASE=student_support_db;", "DATABASE=student_support_db_test;"
    )
    if "DATABASE=student_support_db_test;" not in target:
        raise RuntimeError("Bước 14 chỉ được phép dùng student_support_db_test.")
    db = DatabaseManager(target)
    assert_test_db(db)
    return db


def assert_test_db(db: DatabaseManager) -> None:
    with db.transaction() as connection:
        actual = connection.cursor().execute("SELECT DB_NAME()").fetchone()[0]
    if actual != "student_support_db_test":
        raise RuntimeError(f"Đã hủy write vì database hiện tại là {actual!r}.")


def cleanup(db: DatabaseManager) -> None:
    assert_test_db(db)
    with db.transaction() as connection:
        cursor = connection.cursor()
        cursor.execute(
            f"DELETE FROM dbo.INTERVENTION_REVIEWS WHERE intervention_id IN "
            f"(SELECT intervention_id FROM dbo.INTERVENTIONS WHERE enrollment_id IN "
            f"(SELECT e.enrollment_id FROM dbo.STUDENT_ENROLLMENTS e "
            f"INNER JOIN dbo.STUDENTS s ON s.student_id=e.student_id "
            f"WHERE s.student_code LIKE ?))",
            "T14%",
        )
        cursor.execute(
            f"DELETE FROM dbo.INTERVENTIONS WHERE enrollment_id IN "
            f"(SELECT e.enrollment_id FROM dbo.STUDENT_ENROLLMENTS e "
            f"INNER JOIN dbo.STUDENTS s ON s.student_id=e.student_id "
            f"WHERE s.student_code LIKE ?)",
            "T14%",
        )
        cursor.execute(
            f"DELETE FROM dbo.SCORES WHERE enrollment_id IN "
            f"(SELECT e.enrollment_id FROM dbo.STUDENT_ENROLLMENTS e "
            f"INNER JOIN dbo.STUDENTS s ON s.student_id=e.student_id "
            f"WHERE s.student_code LIKE ?)",
            "T14%",
        )
        cursor.execute(
            "DELETE FROM dbo.SUPPORT_RULES WHERE subject_id IN "
            "(SELECT subject_id FROM dbo.SUBJECTS WHERE subject_code IN (?, ?))",
            SUBJECT_CODE,
            NO_RULE_CODE,
        )
        cursor.execute(
            "DELETE FROM dbo.ASSESSMENTS WHERE subject_id IN "
            "(SELECT subject_id FROM dbo.SUBJECTS WHERE subject_code IN (?, ?))",
            SUBJECT_CODE,
            NO_RULE_CODE,
        )
        cursor.execute(
            "DELETE FROM dbo.STUDENT_ENROLLMENTS WHERE student_id IN "
            "(SELECT student_id FROM dbo.STUDENTS WHERE student_code LIKE ?)",
            "T14%",
        )
        cursor.execute(
            "DELETE FROM dbo.STUDENTS WHERE student_code LIKE ?",
            "T14%",
        )
        cursor.execute("DELETE FROM dbo.CLASSES WHERE class_name = ?", CLASS_NAME)
        cursor.execute(
            "DELETE FROM dbo.SUBJECTS WHERE subject_code IN (?, ?)",
            SUBJECT_CODE,
            NO_RULE_CODE,
        )
        cursor.execute("DELETE FROM dbo.SCHOOL_YEARS WHERE year_name = ?", YEAR_NAME)
        cursor.execute(
            "DELETE FROM dbo.USERS WHERE username LIKE ?", "t14%"
        )


def residual_counts(db: DatabaseManager) -> tuple[int, ...]:
    assert_test_db(db)
    with db.transaction() as connection:
        cursor = connection.cursor()
        checks = (
            ("dbo.USERS", "username LIKE ?", "t14%"),
            ("dbo.STUDENTS", "student_code LIKE ?", "T14%"),
            ("dbo.SCHOOL_YEARS", "year_name LIKE ?", "T14%"),
            ("dbo.CLASSES", "class_name LIKE ?", "%T14%"),
            ("dbo.SUBJECTS", "subject_code LIKE ?", "T14%"),
            ("dbo.ASSESSMENTS", "assessment_name LIKE ?", "T14%"),
        )
        return tuple(
            int(cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE {where}", value).fetchone()[0])
            for table, where, value in checks
        )


def build_context(db: DatabaseManager, session) -> AppContext:
    academic = AcademicService(db)
    students = StudentListService(db)
    enrollments = EnrollmentService(db)
    scores = ScoreService(db)
    preview = ScoreImportPreviewService(db)
    return AppContext(
        db=db,
        auth_service=AuthService(db),
        permission_service=PermissionService(),
        session=session,
        academic_service=academic,
        student_list_service=students,
        student_service=StudentService(db),
        enrollment_service=enrollments,
        student_profile_service=StudentProfileService(db),
        score_service=scores,
        score_import_parser=ScoreImportWorkbookParser(),
        score_import_template_service=ScoreImportTemplateService(),
        score_import_preview_service=preview,
        score_import_commit_service=ScoreImportCommitService(
            db, preview_service=preview, score_service=scores
        ),
        student_export_service=StudentExportService(students),
        score_export_service=ScoreExportService(scores),
        report_service=ReportService(db),
        report_export_service=ReportExportService(),
        support_service=SupportService(db),
        user_service=UserService(db),
    )


@dataclass(frozen=True)
class AcademicFixture:
    grade_id: int
    year_id: int
    class_id: int
    subject_id: int
    assessment_ids: tuple[int, ...]
    rule_id: int


def create_catalog(academic: AcademicService) -> AcademicFixture:
    grades = academic.list_grades()
    if not grades:
        grade_id = academic.create_grade(7, "Khối 7")
    else:
        chosen = next((row for row in grades if int(row[1]) == 7), grades[0])
        grade_id = int(chosen[0])
    year_id = academic.create_school_year(
        YEAR_NAME, date(2026, 9, 1), date(2027, 5, 31), False
    )
    class_id = academic.create_class(CLASS_NAME, grade_id, year_id)
    subject_id = academic.create_subject(SUBJECT_CODE, SUBJECT_NAME)
    assessments = tuple(
        academic.create_assessment(
            subject_id,
            year_id,
            name,
            1,
            "TEST" if index not in (1, 2) else "REVIEW",
            date(2026, 10, 1 + index),
        )
        for index, name in enumerate(ASSESSMENT_NAMES[:6])
    )
    rule = academic.create_support_rule(subject_id, year_id, THRESHOLD)
    assert academic.get_school_year(YEAR_NAME)[0] == year_id
    assert academic.get_subject(SUBJECT_CODE)[0] == subject_id
    assert academic.get_class(class_id)[0] == class_id
    assert {item.assessment_id for item in academic.list_assessments(year_id, subject_id)} == {
        item.assessment_id for item in assessments
    }
    assert academic.list_catalog_support_rules(year_id, subject_id)[0].threshold == THRESHOLD
    return AcademicFixture(
        grade_id,
        year_id,
        class_id,
        subject_id,
        tuple(item.assessment_id for item in assessments),
        rule.rule_id,
    )


def select(combo, value) -> None:
    index = combo.findData(value)
    assert index >= 0
    combo.setCurrentIndex(index)


def fill_score_template(path: Path, values: dict[str, Decimal]) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook["Nhập điểm"]
        for row in range(7, sheet.max_row + 1):
            sheet.cell(row, 4, values[sheet.cell(row, 2).value])
        workbook.save(path)
    finally:
        workbook.close()


def count_for_students(db: DatabaseManager, table: str) -> int:
    assert table in {"SCORES", "INTERVENTIONS", "INTERVENTION_REVIEWS"}
    with db.transaction() as connection:
        if table == "INTERVENTION_REVIEWS":
            sql = (
                f"SELECT COUNT(*) FROM dbo.INTERVENTION_REVIEWS WHERE intervention_id IN "
                f"(SELECT intervention_id FROM dbo.INTERVENTIONS WHERE enrollment_id IN "
                f"(SELECT e.enrollment_id FROM dbo.STUDENT_ENROLLMENTS e "
                f"INNER JOIN dbo.STUDENTS s ON s.student_id=e.student_id "
                f"WHERE s.student_code LIKE ?))"
            )
        else:
            sql = (
                f"SELECT COUNT(*) FROM dbo.{table} WHERE enrollment_id IN "
                f"(SELECT e.enrollment_id FROM dbo.STUDENT_ENROLLMENTS e "
                f"INNER JOIN dbo.STUDENTS s ON s.student_id=e.student_id "
                f"WHERE s.student_code LIKE ?)"
            )
        return int(connection.cursor().execute(sql, "T14%").fetchone()[0])


def test_full_academic_support_reports_and_excel_lifecycle(monkeypatch, tmp_path):
    qapp()
    db = step14_db()
    cleanup(db)
    window = None
    try:
        users = UserService(db)
        admin_user = users.create_user(
            ADMIN_USERNAME, OLD_PASSWORD, "T14 Admin", UserRole.ADMIN
        )
        teacher_user = users.create_user(
            TEACHER_USERNAME, OLD_PASSWORD, "T14 Teacher", UserRole.TEACHER
        )
        session = AuthService(db).login(ADMIN_USERNAME, OLD_PASSWORD)
        context = build_context(db, session)
        window = MainWindow(context)
        assert session.user_id == admin_user.user_id
        assert all(window.can_navigate_to(key) for key in MainWindow.PAGE_TITLES)

        catalog = create_catalog(context.academic_service)
        window.navigate_to("catalogs")
        assert window.pages["catalogs"].initialize_catalogs() is True

        student_a = context.student_service.create_student(
            StudentCreateData(
                STUDENT_CODES[0],
                "=T14 Alpha",
                date(2014, 1, 2),
                "Nữ",
            ),
        )
        student_b = context.student_service.create_student(
            StudentCreateData(STUDENT_CODES[1], "T14 Beta", date(2014, 2, 3), "Nam"),
        )
        enrollment_a = context.enrollment_service.enroll_student(
            student_a.student_id, catalog.class_id, date(2026, 9, 1)
        )
        enrollment_b = context.enrollment_service.enroll_student(
            student_b.student_id, catalog.class_id, date(2026, 9, 1)
        )
        listed = context.student_list_service.list_students(
            StudentFilter(
                school_year_id=catalog.year_id,
                grade_id=catalog.grade_id,
                class_id=catalog.class_id,
            )
        )
        assert {item.student_id for item in listed} == {student_a.student_id, student_b.student_id}
        assert context.enrollment_service.get_active_enrollment(student_a.student_id) == enrollment_a

        trigger_score, intervention = context.score_service.create_score_and_detect(
            enrollment_a.enrollment_id,
            catalog.assessment_ids[0],
            Decimal("5.00"),
        )
        assert intervention.status is InterventionStatus.DETECTED
        assert intervention.trigger_score_id == trigger_score.score_id
        with pytest.raises(InvalidStateTransitionError):
            context.support_service.start_intervention(intervention.intervention_id)

        planned = context.support_service.plan_intervention(
            intervention.intervention_id,
            teacher_user.user_id,
            date(2026, 10, 10),
            "T14 phụ đạo",
            "T14 theo dõi",
        )
        assert planned.status is InterventionStatus.PLANNED
        assert context.support_service.start_intervention(intervention.intervention_id).status is InterventionStatus.IN_PROGRESS
        assert context.support_service.mark_waiting_review(intervention.intervention_id).status is InterventionStatus.WAITING_REVIEW
        first_review = context.support_service.review_intervention(
            intervention.intervention_id,
            review_date=date(2026, 11, 1),
            assessment_id=catalog.assessment_ids[1],
            score_value=THRESHOLD - Decimal("0.01"),
            notes="T14 chưa đạt",
        )
        assert first_review.status is InterventionStatus.CONTINUE
        assert context.support_service.continue_intervention(intervention.intervention_id).status is InterventionStatus.IN_PROGRESS
        context.support_service.mark_waiting_review(intervention.intervention_id)
        completed = context.support_service.review_intervention(
            intervention.intervention_id,
            review_date=date(2026, 12, 1),
            assessment_id=catalog.assessment_ids[2],
            score_value=THRESHOLD,
            notes="T14 đạt ngưỡng",
        )
        assert completed.status is InterventionStatus.COMPLETED
        with pytest.raises(InvalidStateTransitionError):
            context.support_service.continue_intervention(intervention.intervention_id)

        detail = context.support_service.get_intervention_detail(intervention.intervention_id)
        assert [review.result for review in detail.reviews] == [
            ReviewResult.NOT_PASSED,
            ReviewResult.PASSED,
        ]
        assert [review.review_date for review in detail.reviews] == [
            date(2026, 11, 1),
            date(2026, 12, 1),
        ]
        assert context.score_service.can_edit_score(trigger_score.score_id) is False
        assert all(context.score_service.can_edit_score(review.score_id) is False for review in detail.reviews)

        ordinary, ordinary_intervention = context.score_service.create_score_and_detect(
            enrollment_b.enrollment_id,
            catalog.assessment_ids[4],
            Decimal("8.00"),
        )
        assert ordinary_intervention is None
        assert context.score_service.can_edit_score(ordinary.score_id) is True
        with pytest.raises(DuplicateError):
            context.score_service.create_score_and_detect(
                enrollment_b.enrollment_id, catalog.assessment_ids[4], Decimal("9")
            )

        new_score, new_episode = context.score_service.create_score_and_detect(
            enrollment_a.enrollment_id,
            catalog.assessment_ids[3],
            Decimal("4.00"),
        )
        assert new_episode.intervention_id != intervention.intervention_id
        assert new_episode.status is InterventionStatus.DETECTED
        assert context.support_service.get_intervention(intervention.intervention_id).status is InterventionStatus.COMPLETED

        import_context = ScoreImportContext(
            catalog.year_id,
            YEAR_NAME,
            catalog.class_id,
            CLASS_NAME,
            catalog.subject_id,
            SUBJECT_NAME,
            catalog.assessment_ids[5],
            ASSESSMENT_NAMES[5],
        )
        import_path = tmp_path / "T14-import.xlsx"
        roster = context.enrollment_service.list_class_enrollments(
            catalog.class_id, catalog.year_id
        )
        context.score_import_template_service.create_template(
            import_context,
            tuple(ScoreImportTemplateStudent(item.student_id, item.full_name) for item in roster),
            import_path,
        )
        fill_score_template(
            import_path,
            {student_a.student_id: Decimal("9.00"), student_b.student_id: Decimal("5.50")},
        )
        parsed = context.score_import_parser.parse_workbook(import_path)
        preview = context.score_import_preview_service.preview_import(import_context, parsed)
        assert preview.can_commit and preview.invalid_count == 0
        imported = context.score_import_commit_service.commit_import(preview)
        assert imported.imported_count == 2
        assert imported.intervention_created_count == 1

        profile_a = context.student_profile_service.get_profile(student_a.student_id)
        profile_b = context.student_profile_service.get_profile(student_b.student_id)
        assert len(profile_a.enrollment_history) == 1
        assert {item.status for item in profile_a.intervention_history} == {
            InterventionStatus.COMPLETED,
            InterventionStatus.DETECTED,
        }
        assert len(profile_b.intervention_history) == 1

        report = context.report_service.get_support_report(
            catalog.year_id,
            catalog.grade_id,
            catalog.class_id,
            catalog.subject_id,
        )
        assert report.summary.total_cases == 3
        assert report.summary.completed_count == 1
        assert report.summary.detected_count == 2
        completed_report = context.report_service.get_support_report(
            catalog.year_id,
            catalog.grade_id,
            catalog.class_id,
            catalog.subject_id,
            InterventionStatus.COMPLETED,
        )
        assert completed_report.summary.total_cases == 1
        assert completed_report.rows[0].latest_review_score == THRESHOLD

        window.navigate_to("support")
        support_page = window.pages["support"]
        select(support_page.school_year_combo, catalog.year_id)
        select(support_page.grade_combo, catalog.grade_id)
        select(support_page.class_combo, catalog.class_id)
        select(support_page.subject_combo, catalog.subject_id)
        select(support_page.status_combo, None)
        assert {item.intervention_id for item in support_page.items} == {
            row.intervention_id for row in report.rows
        }
        window.navigate_to("reports")
        reports_page = window.pages["reports"]
        select(reports_page.school_year_combo, catalog.year_id)
        select(reports_page.grade_combo, catalog.grade_id)
        select(reports_page.class_combo, catalog.class_id)
        select(reports_page.subject_combo, catalog.subject_id)
        select(reports_page.status_combo, None)
        assert reports_page.report_data == report
        assert reports_page.table.rowCount() == 3
        assert reports_page.kpi_cards["total_cases"].value_label.text() == "3"

        student_path = tmp_path / "T14-students.xlsx"
        score_path = tmp_path / "T14-scores.xlsx"
        support_path = tmp_path / "T14-support.xlsx"
        report_path = tmp_path / "T14-report.xlsx"
        context.student_export_service.export_xlsx(
            StudentExportContext(
                StudentFilter(
                    school_year_id=catalog.year_id,
                    grade_id=catalog.grade_id,
                    class_id=catalog.class_id,
                ),
                YEAR_NAME,
                f"Khối {context.academic_service.get_class(catalog.class_id)[3]}",
                CLASS_NAME,
            ),
            student_path,
        )
        context.score_export_service.export_xlsx(
            ScoreExportContext(
                catalog.year_id,
                YEAR_NAME,
                catalog.class_id,
                CLASS_NAME,
                catalog.subject_id,
                SUBJECT_NAME,
                catalog.assessment_ids[5],
                ASSESSMENT_NAMES[5],
            ),
            score_path,
        )
        export_context = SupportReportExportContext(
            catalog.year_id,
            YEAR_NAME,
            catalog.grade_id,
            "Khối",
            catalog.class_id,
            CLASS_NAME,
            catalog.subject_id,
            SUBJECT_NAME,
        )
        export_data = SupportReportExportData(export_context, report)
        context.report_export_service.export_xlsx(export_data, support_path)
        context.report_export_service.export_xlsx(export_data, report_path)

        workbook = load_workbook(student_path, data_only=False)
        try:
            sheet = workbook["Danh sách học sinh"]
            assert sheet.max_row == 9
            assert {sheet["B8"].value, sheet["B9"].value} == set(STUDENT_CODES[:2])
            formula_row = 8 if sheet["C8"].value == "=T14 Alpha" else 9
            assert sheet.cell(formula_row, 3).data_type == "s"
        finally:
            workbook.close()
        workbook = load_workbook(score_path, data_only=False)
        try:
            sheet = workbook["Bảng điểm"]
            assert {sheet["G9"].value, sheet["G10"].value} == {9, 5.5}
            assert sheet["G9"].data_type == sheet["G10"].data_type == "n"
        finally:
            workbook.close()
        for path in (support_path, report_path):
            workbook = load_workbook(path, data_only=False)
            try:
                sheet = workbook["Báo cáo bổ trợ"]
                assert sheet.max_row == 18
                assert sheet.max_column == 13
                assert sheet["B7"].value == 3
                assert sheet["B13"].value == 1
                assert all(sheet.cell(row, 3).data_type == "s" for row in range(16, 19))
            finally:
                workbook.close()
        assert not list(Path.cwd().glob("T14*.xlsx"))
    finally:
        if window is not None:
            window.close()
        cleanup(db)
        assert residual_counts(db) == (0, 0, 0, 0, 0, 0)


class AcceptedUserDialog:
    DialogCode = QDialog.DialogCode

    def __init__(self, values):
        self._values = values

    def exec(self):
        return self.DialogCode.Accepted

    def create_values(self):
        return self._values


class AcceptedPasswordDialog:
    DialogCode = QDialog.DialogCode

    def __init__(self, old_password, new_password):
        self._values = (old_password, new_password)
        self.cleared = False

    def exec(self):
        return self.DialogCode.Accepted

    def values(self):
        return self._values

    def clear_passwords(self):
        self.cleared = True


def test_account_activation_permission_and_session_lifecycle(monkeypatch):
    qapp()
    db = step14_db()
    cleanup(db)
    old_window = new_window = None
    try:
        users = UserService(db)
        admin_user = users.create_user(
            ADMIN_USERNAME, OLD_PASSWORD, "T14 Admin", UserRole.ADMIN
        )
        auth = AuthService(db)
        admin = auth.login(ADMIN_USERNAME, OLD_PASSWORD)
        context = build_context(db, admin)
        old_window = MainWindow(context)
        read_only_before = residual_counts(db) + tuple(
            count_for_students(db, table)
            for table in ("SCORES", "INTERVENTIONS", "INTERVENTION_REVIEWS")
        )
        for page_key in MainWindow.PAGE_TITLES:
            old_window.navigate_to(page_key)
        read_only_after = residual_counts(db) + tuple(
            count_for_students(db, table)
            for table in ("SCORES", "INTERVENTIONS", "INTERVENTION_REVIEWS")
        )
        assert read_only_after == read_only_before
        assert not list(Path.cwd().glob("T14*.xlsx"))

        old_window.navigate_to("system")
        system_page = old_window.pages["system"]
        assert system_page.tabs.count() == 2
        system_page.user_dialog_factory = lambda *_: AcceptedUserDialog(
            (
                TEACHER_USERNAME,
                OLD_PASSWORD,
                "T14 Teacher",
                UserRole.TEACHER,
                "t14.teacher@example.com",
                "0901400000",
                True,
            )
        )
        assert system_page.create_user() is True
        teacher_user = users.get_by_username(TEACHER_USERNAME)
        assert teacher_user is not None
        assert all(not hasattr(item, "password_hash") for item in system_page.users)

        password_dialog = AcceptedPasswordDialog(OLD_PASSWORD, NEW_PASSWORD)
        system_page.change_password_dialog_factory = lambda *_: password_dialog
        assert system_page.change_own_password() is True
        assert password_dialog.cleared is True
        stored = users.get_by_username(ADMIN_USERNAME)
        assert stored.password_hash not in {OLD_PASSWORD, NEW_PASSWORD}
        assert bcrypt.checkpw(NEW_PASSWORD.encode(), stored.password_hash.encode())
        with pytest.raises(ValidationError):
            auth.login(ADMIN_USERNAME, OLD_PASSWORD)
        relogged_admin = auth.login(ADMIN_USERNAME, NEW_PASSWORD)

        with pytest.raises(BusinessRuleError):
            users.admin_set_user_active(relogged_admin, admin_user.user_id, False)
        with pytest.raises(BusinessRuleError):
            users.admin_update_user(
                relogged_admin,
                admin_user.user_id,
                "T14 Demoted",
                UserRole.TEACHER,
            )

        users.admin_set_user_active(relogged_admin, teacher_user.user_id, False)
        with pytest.raises(ValidationError):
            auth.login(TEACHER_USERNAME, OLD_PASSWORD)
        users.admin_set_user_active(relogged_admin, teacher_user.user_id, True)
        teacher = auth.login(TEACHER_USERNAME, OLD_PASSWORD)

        old_window.request_logout()
        assert context.session is None
        assert all(not old_window.can_navigate_to(key) for key in MainWindow.PAGE_TITLES)
        context.set_session(teacher)
        new_window = MainWindow(context)
        expected = {
            "dashboard": True,
            "students": True,
            "scores": True,
            "support": True,
            "reports": True,
            "catalogs": False,
            "system": True,
        }
        assert {key: new_window.can_navigate_to(key) for key in expected} == expected
        with pytest.raises(PermissionError):
            new_window.navigate_to("catalogs")
        new_window.navigate_to("system")
        teacher_page = new_window.pages["system"]
        assert teacher_page.tabs.count() == 1
        assert teacher_page.own_profile.user_id == teacher_user.user_id
        with pytest.raises(ValidationError):
            users.admin_list_users(teacher)
        assert old_window.can_navigate_to("students") is False
        warnings = []
        monkeypatch.setattr(
            "ui.pages.students_page.QMessageBox.warning",
            lambda _parent, title, message: warnings.append((title, message)),
        )
        assert old_window.pages["students"].export_students() is False
        assert len(warnings) == 1
    finally:
        if new_window is not None:
            new_window.close()
        if old_window is not None:
            old_window.close()
        cleanup(db)
        assert residual_counts(db) == (0, 0, 0, 0, 0, 0)


class FailingUserRepository(UserRepository):
    def update_profile(self, connection, user_id, full_name, email, phone):
        super().update_profile(connection, user_id, full_name, email, phone)
        raise pyodbc.Error("raw SQL driver failure")


def test_failure_rollbacks_integrity_and_architecture_guards():
    db = step14_db()
    cleanup(db)
    try:
        users = UserService(db)
        users.create_user(
            ROLLBACK_USERNAME,
            OLD_PASSWORD,
            "T14 Before rollback",
            UserRole.TEACHER,
        )
        actor = AuthService(db).login(ROLLBACK_USERNAME, OLD_PASSWORD)
        failing_users = UserService(db, FailingUserRepository())
        with pytest.raises(DatabaseError) as user_error:
            failing_users.update_own_profile(
                actor, "T14 Must rollback", "t14.rollback@example.com", "0901400001"
            )
        assert "pyodbc" not in str(user_error.value).lower()
        assert users.get_own_profile(actor).full_name == "T14 Before rollback"

        academic = AcademicService(db)
        catalog = create_catalog(academic)
        no_rule_subject = academic.create_subject(NO_RULE_CODE, "T14 Không rule")
        no_rule_assessment = academic.create_assessment(
            no_rule_subject,
            catalog.year_id,
            ASSESSMENT_NAMES[6],
            1,
            "TEST",
            date(2026, 11, 20),
        )
        student = StudentService(db).create_student(
            StudentCreateData(STUDENT_CODES[2], "T14 Rollback Student"),
        )
        enrollment = EnrollmentService(db).enroll_student(
            student.student_id, catalog.class_id, date(2026, 9, 1)
        )
        with pytest.raises(MissingSupportRuleError) as score_error:
            ScoreService(db).create_score_and_detect(
                enrollment.enrollment_id,
                no_rule_assessment.assessment_id,
                Decimal("1.00"),
            )
        assert "pyodbc" not in str(score_error.value).lower()
        assert count_for_students(db, "SCORES") == 0
        assert count_for_students(db, "INTERVENTIONS") == 0
        assert count_for_students(db, "INTERVENTION_REVIEWS") == 0

        with db.transaction() as connection:
            cursor = connection.cursor()
            orphan_scores = cursor.execute(
                "SELECT COUNT(*) FROM dbo.SCORES s LEFT JOIN dbo.STUDENT_ENROLLMENTS e "
                "ON e.enrollment_id=s.enrollment_id WHERE e.enrollment_id IS NULL"
            ).fetchone()[0]
            orphan_interventions = cursor.execute(
                "SELECT COUNT(*) FROM dbo.INTERVENTIONS i "
                "LEFT JOIN dbo.SCORES s ON s.score_id=i.trigger_score_id "
                "LEFT JOIN dbo.STUDENT_ENROLLMENTS e ON e.enrollment_id=i.enrollment_id "
                "WHERE s.score_id IS NULL OR e.enrollment_id IS NULL"
            ).fetchone()[0]
            orphan_reviews = cursor.execute(
                "SELECT COUNT(*) FROM dbo.INTERVENTION_REVIEWS r "
                "LEFT JOIN dbo.INTERVENTIONS i ON i.intervention_id=r.intervention_id "
                "LEFT JOIN dbo.SCORES s ON s.score_id=r.score_id "
                "WHERE i.intervention_id IS NULL OR s.score_id IS NULL"
            ).fetchone()[0]
            duplicate_scores = cursor.execute(
                "SELECT COUNT(*) FROM (SELECT enrollment_id, assessment_id FROM dbo.SCORES "
                "GROUP BY enrollment_id, assessment_id HAVING COUNT(*) > 1) d"
            ).fetchone()[0]
            duplicate_open_interventions = cursor.execute(
                "SELECT COUNT(*) FROM (SELECT enrollment_id, subject_id "
                "FROM dbo.INTERVENTIONS WHERE status <> 'COMPLETED' "
                "GROUP BY enrollment_id, subject_id HAVING COUNT(*) > 1) d"
            ).fetchone()[0]
        assert (
            orphan_scores
            == orphan_interventions
            == orphan_reviews
            == duplicate_scores
            == duplicate_open_interventions
            == 0
        )

        from repositories.intervention_repository import InterventionRepository
        from repositories.report_repository import ReportRepository
        from repositories.score_repository import ScoreRepository
        from ui.pages.catalog_page import CatalogPage
        from ui.pages.reports_page import ReportsPage
        from ui.pages.scores_page import ScoresPage
        from ui.pages.students_page import StudentsPage
        from ui.pages.support_page import SupportPage

        ui_source = "\n".join(
            inspect.getsource(item)
            for item in (CatalogPage, StudentsPage, ScoresPage, SupportPage, ReportsPage)
        ).upper()
        for forbidden in (
            "SELECT ",
            "INSERT INTO",
            "DELETE FROM",
            "UPDATE DBO",
            ".TRANSACTION(",
            "OPENPYXL",
            "REPOSITORY",
        ):
            assert forbidden not in ui_source
        repository_source = "\n".join(
            inspect.getsource(item)
            for item in (ScoreRepository, InterventionRepository, ReportRepository)
        ).upper()
        assert ".COMMIT(" not in repository_source
        report_source = inspect.getsource(ReportRepository.list_support_cases).upper()
        assert "OUTER APPLY" in report_source and "GET_DETAIL" not in report_source
        preview_source = inspect.getsource(ScoreImportPreviewService._build_preview)
        assert all(
            name in preview_source
            for name in ("list_by_ids", "list_by_student_ids", "list_by_class_assessment")
        )
        from bootstrap import build_app_context
        assert "ImportService" not in inspect.getsource(build_app_context)
    finally:
        cleanup(db)
        assert residual_counts(db) == (0, 0, 0, 0, 0, 0)
