from __future__ import annotations

import inspect
import os
from dataclasses import FrozenInstanceError
from datetime import date
from decimal import Decimal

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from bootstrap import build_app_context
from exceptions import DataExportError, ValidationError
from models.dto.data_export import ScoreExportContext, StudentExportContext
from models.dto.score import ScoreRosterItem
from models.dto.student_filter import StudentFilter
from models.dto.student_list import StudentListItem
from models.enums import StudentStatus
from services.data_export_contract import (
    ScoreExportServiceContract,
    StudentExportServiceContract,
)
from services.data_export_service import ScoreExportService, StudentExportService
from ui.pages.scores_page import ScoresPage
from ui.pages.students_page import StudentsPage


def app():
    return QApplication.instance() or QApplication([])


def student_item(
    student_id="s1",
    student_code="HS01",
    full_name="An",
    school_year="2026-2027",
):
    return StudentListItem(
        student_id,
        student_code,
        full_name,
        date(2012, 1, 2),
        "Nam",
        "6A1",
        6,
        StudentStatus.ACTIVE,
        school_year,
    )


class StudentReader:
    def __init__(self, rows=()):
        self.rows = list(rows)
        self.calls = []

    def list_students(self, filters):
        self.calls.append(filters)
        return list(self.rows)


def student_context(**changes):
    values = dict(
        filters=StudentFilter(
            search_text="An",
            school_year_id=2,
            grade_id=6,
            class_id=61,
            status=StudentStatus.ACTIVE,
        ),
        school_year_name="2026-2027",
        grade_name="Khối 6",
        class_name="6A1",
    )
    values.update(changes)
    return StudentExportContext(**values)


def score_item(
    enrollment_id=1,
    student_id="s1",
    student_code="HS01",
    full_name="An",
    score=Decimal("7.25"),
):
    return ScoreRosterItem(
        enrollment_id=enrollment_id,
        student_id=student_id,
        student_code=student_code,
        full_name=full_name,
        assessment_id=101,
        score_id=1000 + enrollment_id if score is not None else None,
        score=score,
    )


class ScoreReader:
    def __init__(self, rows=()):
        self.rows = list(rows)
        self.calls = []

    def list_score_roster(self, *args):
        self.calls.append(args)
        return list(self.rows)


def score_context(**changes):
    values = dict(
        school_year_id=2,
        school_year_name="2026-2027",
        class_id=61,
        class_name="6A1",
        subject_id=11,
        subject_name="Vật lý",
        assessment_id=101,
        assessment_name="Giữa kỳ",
    )
    values.update(changes)
    return ScoreExportContext(**values)


def test_export_contracts_and_contexts_are_immutable():
    students: StudentExportServiceContract = StudentExportService(StudentReader())
    scores: ScoreExportServiceContract = ScoreExportService(ScoreReader())
    assert students is not None and scores is not None
    with pytest.raises(FrozenInstanceError):
        score_context().class_id = 62


def test_student_export_reads_filter_once_sorts_and_deduplicates_identity():
    reader = StudentReader((
        student_item("s2", "HS02", "Bình"),
        student_item("s1", "HS01", "An"),
        student_item("s1", "DUP", "Trùng"),
    ))
    service = StudentExportService(reader)
    context = student_context()
    data = service.get_export_data(context)
    assert reader.calls == [context.filters]
    assert [(row.student_id, row.student_code) for row in data.rows] == [
        ("s1", "HS01"), ("s2", "HS02")
    ]
    assert data.rows[0].school_year_name == "2026-2027"


def test_student_workbook_structure_context_and_fields(tmp_path):
    service = StudentExportService(StudentReader((student_item(),)))
    path = service.export_xlsx(student_context(), tmp_path / "students.xlsx")
    workbook = load_workbook(path)
    try:
        sheet = workbook["Danh sách học sinh"]
        assert sheet["A1"].value == "DANH SÁCH HỌC SINH"
        assert (sheet["B3"].value, sheet["B4"].value, sheet["B5"].value) == (
            "2026-2027", "Khối 6", "6A1"
        )
        assert [sheet.cell(7, column).value for column in range(1, 9)] == list(
            StudentExportService.HEADERS
        )
        values = [sheet.cell(8, column).value for column in range(1, 9)]
        assert values[:3] == [1, "HS01", "An"]
        assert values[3].date() == date(2012, 1, 2)
        assert values[4:] == ["Nam", "2026-2027", 6, "6A1"]
        assert sheet.freeze_panes == "A8"
    finally:
        workbook.close()


def test_student_export_empty_dataset_creates_headers_only(tmp_path):
    path = StudentExportService(StudentReader()).export_xlsx(
        student_context(), tmp_path / "empty.xlsx"
    )
    workbook = load_workbook(path)
    try:
        assert workbook.active.max_row == StudentExportService.HEADER_ROW
        assert workbook.active.auto_filter.ref == "A7:H7"
    finally:
        workbook.close()


def test_student_formula_like_text_is_forced_to_string(tmp_path):
    item = student_item(student_code="+CMD", full_name="=SUM(1,1)", school_year="@YEAR")
    path = StudentExportService(StudentReader((item,))).export_xlsx(
        student_context(class_name="-6A1"), tmp_path / "safe.xlsx"
    )
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        for coordinate in ("B8", "C8", "F8", "H8"):
            assert sheet[coordinate].data_type == "s"
    finally:
        workbook.close()


def test_student_export_rejects_bad_path_and_normalizes_write_error(monkeypatch, tmp_path):
    service = StudentExportService(StudentReader())
    with pytest.raises(ValidationError):
        service.export_xlsx(student_context(), tmp_path / "students.csv")
    monkeypatch.setattr("services.data_export_service.Workbook.save", lambda *_: (_ for _ in ()).throw(PermissionError("denied")))
    with pytest.raises(DataExportError, match="Không thể ghi file danh sách"):
        service.export_xlsx(student_context(), tmp_path / "students.xlsx")


def test_score_export_reads_authoritative_context_once_and_keeps_missing_score():
    reader = ScoreReader((
        score_item(2, "s2", "HS02", "Bình", None),
        score_item(1, "s1", "HS01", "An", Decimal("7.25")),
        score_item(1, "s1", "DUP", "Trùng", Decimal("8")),
    ))
    service = ScoreExportService(reader)
    context = score_context()
    data = service.get_export_data(context)
    assert reader.calls == [(61, 2, 11, 101)]
    assert [(row.student_id, row.score) for row in data.rows] == [
        ("s1", Decimal("7.25")), ("s2", None)
    ]


@pytest.mark.parametrize("field,value", [
    ("school_year_id", 0), ("class_id", True), ("subject_id", -1),
    ("assessment_id", None), ("subject_name", ""),
])
def test_score_export_validates_context(field, value):
    service = ScoreExportService(ScoreReader())
    with pytest.raises(ValidationError):
        service.get_export_data(score_context(**{field: value}))


def test_score_workbook_context_numeric_score_and_blank_missing_value(tmp_path):
    reader = ScoreReader((score_item(), score_item(2, "s2", "HS02", "Bình", None)))
    path = ScoreExportService(reader).export_xlsx(
        score_context(), tmp_path / "scores.xlsx"
    )
    workbook = load_workbook(path)
    try:
        sheet = workbook["Bảng điểm"]
        assert sheet["A1"].value == "BẢNG ĐIỂM"
        assert [sheet[coordinate].value for coordinate in ("B3", "B4", "B5", "B6")] == [
            "2026-2027", "6A1", "Vật lý", "Giữa kỳ"
        ]
        assert sheet["B9"].value == "HS01"
        assert sheet["G9"].value == pytest.approx(7.25)
        assert sheet["B10"].value == "HS02"
        assert sheet["G10"].value is None
        assert sheet.freeze_panes == "A9"
    finally:
        workbook.close()


def test_score_formula_like_context_and_student_text_is_safe_but_score_numeric(tmp_path):
    reader = ScoreReader((score_item(student_code="@HS", full_name="=1+1"),))
    context = score_context(class_name="+Class", subject_name="-Subject", assessment_name="@Exam")
    path = ScoreExportService(reader).export_xlsx(context, tmp_path / "safe-score.xlsx")
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        for coordinate in ("B4", "B5", "B6", "B9", "C9", "D9", "E9", "F9"):
            assert sheet[coordinate].data_type == "s"
        assert sheet["G9"].data_type == "n"
    finally:
        workbook.close()


def test_score_export_empty_and_write_error(tmp_path, monkeypatch):
    service = ScoreExportService(ScoreReader())
    path = service.export_xlsx(score_context(), tmp_path / "empty-score.xlsx")
    workbook = load_workbook(path)
    try:
        assert workbook.active.max_row == ScoreExportService.HEADER_ROW
    finally:
        workbook.close()
    monkeypatch.setattr("services.data_export_service.Workbook.save", lambda *_: (_ for _ in ()).throw(OSError("locked")))
    with pytest.raises(DataExportError, match="Không thể ghi file bảng điểm"):
        service.export_xlsx(score_context(), tmp_path / "locked.xlsx")


class ExportRecorder:
    def __init__(self, error=None):
        self.calls = []
        self.error = error

    def export_xlsx(self, context, path):
        self.calls.append((context, path))
        if self.error:
            raise self.error
        return path


def set_student_context(page):
    for combo, label, value in (
        (page.filter_widget.school_year_combo, "2026-2027", 2),
        (page.filter_widget.grade_combo, "Khối 6", 6),
        (page.filter_widget.class_combo, "6A1", 61),
    ):
        combo.addItem(label, value)
        combo.setCurrentIndex(combo.count() - 1)


def set_score_context(page):
    for combo, label, value in (
        (page.school_year_combo, "2026-2027", 2),
        (page.grade_combo, "Khối 6", 6),
        (page.class_combo, "6A1", 61),
        (page.subject_combo, "Vật lý", 11),
        (page.assessment_combo, "Giữa kỳ", 101),
    ):
        combo.addItem(label, value)
        combo.setCurrentIndex(combo.count() - 1)


def test_students_page_export_action_cancel_and_data_api_context(monkeypatch):
    app(); exporter = ExportRecorder(); page = StudentsPage(student_export_service=exporter)
    assert page.export_button.text() == "Xuất Excel"
    monkeypatch.setattr("ui.pages.students_page.QFileDialog.getSaveFileName", lambda *_: ("", ""))
    assert page.export_students() is False and exporter.calls == []
    set_student_context(page)
    monkeypatch.setattr("ui.pages.students_page.QFileDialog.getSaveFileName", lambda *_: ("students", ""))
    monkeypatch.setattr("ui.pages.students_page.QMessageBox.information", lambda *_: None)
    assert page.export_students() is True
    context, path = exporter.calls[0]
    assert (context.filters.school_year_id, context.filters.grade_id, context.filters.class_id) == (2, 6, 61)
    assert (context.school_year_name, context.grade_name, context.class_name) == ("2026-2027", "Khối 6", "6A1")
    assert path.endswith("students.xlsx")


def test_students_page_export_error_is_handled(monkeypatch):
    app(); exporter = ExportRecorder(DataExportError("Không thể ghi file.")); page = StudentsPage(student_export_service=exporter)
    monkeypatch.setattr("ui.pages.students_page.QFileDialog.getSaveFileName", lambda *_: ("students.xlsx", ""))
    warnings = []
    monkeypatch.setattr("ui.pages.students_page.QMessageBox.warning", lambda *args: warnings.append(args))
    assert page.export_students() is False and warnings


def test_scores_page_export_requires_context_and_passes_ids_labels(monkeypatch):
    app(); exporter = ExportRecorder(); page = ScoresPage(score_export_service=exporter)
    assert page.export_button.text() == "Xuất Excel"
    assert page.export_scores() is False and exporter.calls == []
    set_score_context(page)
    monkeypatch.setattr("ui.pages.scores_page.QFileDialog.getSaveFileName", lambda *_: ("scores", ""))
    monkeypatch.setattr("ui.pages.scores_page.QMessageBox.information", lambda *_: None)
    assert page.export_scores() is True
    context, path = exporter.calls[0]
    assert (context.school_year_id, context.class_id, context.subject_id, context.assessment_id) == (2, 61, 11, 101)
    assert (context.school_year_name, context.class_name, context.subject_name, context.assessment_name) == ("2026-2027", "6A1", "Vật lý", "Giữa kỳ")
    assert path.endswith("scores.xlsx")


def test_scores_page_cancel_and_export_error_are_handled(monkeypatch):
    app(); exporter = ExportRecorder(DataExportError("File đang mở.")); page = ScoresPage(score_export_service=exporter)
    set_score_context(page)
    monkeypatch.setattr("ui.pages.scores_page.QFileDialog.getSaveFileName", lambda *_: ("", ""))
    assert page.export_scores() is False and exporter.calls == []
    monkeypatch.setattr("ui.pages.scores_page.QFileDialog.getSaveFileName", lambda *_: ("scores.xlsx", ""))
    warnings = []
    monkeypatch.setattr("ui.pages.scores_page.QMessageBox.warning", lambda *args: warnings.append(args))
    assert page.export_scores() is False and warnings


def test_bootstrap_injects_exporters_using_existing_read_services():
    context = build_app_context()
    assert isinstance(context.student_export_service, StudentExportService)
    assert isinstance(context.score_export_service, ScoreExportService)
    assert context.student_export_service.student_service is context.student_list_service
    assert context.score_export_service.score_service is context.score_service


def test_export_architecture_is_read_only_and_ui_has_no_workbook_or_repository_source():
    service_source = inspect.getsource(StudentExportService) + inspect.getsource(ScoreExportService)
    ui_source = inspect.getsource(StudentsPage) + inspect.getsource(ScoresPage)
    assert "QTableWidget" not in service_source
    for forbidden in ("INSERT ", "UPDATE DBO", "DELETE ", ".COMMIT(", ".TRANSACTION("):
        assert forbidden not in service_source.upper()
    for forbidden in ("FROM OPENPYXL", "IMPORT OPENPYXL", "= WORKBOOK(", "REPOSITORY", "SELECT ", "INSERT ", "UPDATE DBO", "DELETE "):
        assert forbidden not in ui_source.upper()
