from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
import inspect
import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
import pytest
from PySide6.QtWidgets import QApplication, QDialog, QMessageBox

from app_context import AppContext
from config.database import db_settings
from database.connection import DatabaseManager
from exceptions import MissingSupportRuleError
from models.dto import StudentCreateData
from models.dto.data_export import ScoreExportContext, StudentExportContext
from models.dto.score_import import (
    ScoreImportContext,
    ScoreImportTemplateStudent,
)
from models.dto.student_filter import StudentFilter
from models.enums import InterventionStatus, UserRole
from repositories import (
    AcademicRepository,
    EnrollmentRepository,
    StudentRepository,
    SupportRuleRepository,
)
from services.academic_service import AcademicService
from services.auth_service import AuthService
from services.data_export_service import ScoreExportService, StudentExportService
from services.enrollment_service import EnrollmentService
from services.permission_service import PermissionService
from services.report_export_service import ReportExportService
from services.report_service import ReportService
from services.score_import_commit_service import ScoreImportCommitService
from services.score_import_parser import ScoreImportWorkbookParser
from services.score_import_preview_service import ScoreImportPreviewService
from services.score_import_template_service import ScoreImportTemplateService
from services.score_service import ScoreService
from services.student_list_service import StudentListService
from services.student_profile_service import StudentProfileService
from services.student_service import StudentService
from services.support_service import SupportService
from services.user_service import UserService
from ui.main_window import MainWindow


PREFIX = "T138"
YEAR_NAME = "T138_2627"
CLASS_NAME = "-T138_6A"
SUBJECT_CODE = "T138_M1"
SUBJECT_NAME = "@T138 Môn"
NO_RULE_SUBJECT_CODE = "T138_NR"
NO_RULE_SUBJECT_NAME = "+T138 Không rule"
ASSESSMENT_NAMES = (
    "T138 Nhập 1",
    "T138 Nhập 2",
    "T138 Đánh giá lại",
    "T138 Rollback",
)
STUDENT_IDS = tuple(f"t138-student-{suffix}" for suffix in "abc")
STUDENT_CODES = ("T13801", "T13802", "T13803")
STUDENT_NAMES = ("=T138 An", "+T138 Bình", "@T138 Chi")
UNSCORED_STUDENT_ID = "t138-student-d"
UNSCORED_STUDENT_CODE = "T13804"
UNSCORED_STUDENT_NAME = "T138 Dũng"
ALL_STUDENT_IDS = STUDENT_IDS + (UNSCORED_STUDENT_ID,)
ADMIN_USERNAME = "t138_admin"
TEACHER_USERNAME = "t138_teacher"
PASSWORD = "T138Password!"
THRESHOLD = Decimal("6.25")


def qapp() -> QApplication:
    return QApplication.instance() or QApplication([])


def get_test_db() -> DatabaseManager:
    configured = db_settings.connection_string()
    target = configured.replace(
        "DATABASE=student_support_db;",
        "DATABASE=student_support_db_test;",
    )
    if "DATABASE=student_support_db_test;" not in target:
        raise RuntimeError("Step 13.8 chỉ được phép dùng student_support_db_test.")
    db = DatabaseManager(target)
    assert_test_database(db)
    return db


def assert_test_database(db: DatabaseManager) -> None:
    with db.transaction() as connection:
        actual = connection.cursor().execute("SELECT DB_NAME()").fetchone()[0]
    if actual != "student_support_db_test":
        raise RuntimeError(
            f"Đã hủy Step 13.8 vì database hiện tại là {actual!r}."
        )


def cleanup(db: DatabaseManager) -> None:
    assert_test_database(db)
    placeholders = ", ".join("?" for _ in ALL_STUDENT_IDS)
    with db.transaction() as connection:
        cursor = connection.cursor()
        cursor.execute(
            f"DELETE FROM dbo.INTERVENTION_REVIEWS WHERE intervention_id IN "
            f"(SELECT intervention_id FROM dbo.INTERVENTIONS WHERE enrollment_id IN "
            f"(SELECT enrollment_id FROM dbo.STUDENT_ENROLLMENTS "
            f"WHERE student_id IN ({placeholders})))",
            *ALL_STUDENT_IDS,
        )
        cursor.execute(
            f"DELETE FROM dbo.INTERVENTIONS WHERE enrollment_id IN "
            f"(SELECT enrollment_id FROM dbo.STUDENT_ENROLLMENTS "
            f"WHERE student_id IN ({placeholders}))",
            *ALL_STUDENT_IDS,
        )
        cursor.execute(
            f"DELETE FROM dbo.SCORES WHERE enrollment_id IN "
            f"(SELECT enrollment_id FROM dbo.STUDENT_ENROLLMENTS "
            f"WHERE student_id IN ({placeholders}))",
            *ALL_STUDENT_IDS,
        )
        cursor.execute(
            "DELETE FROM dbo.SUPPORT_RULES WHERE subject_id IN "
            "(SELECT subject_id FROM dbo.SUBJECTS WHERE subject_code IN (?, ?))",
            SUBJECT_CODE,
            NO_RULE_SUBJECT_CODE,
        )
        cursor.execute(
            "DELETE FROM dbo.ASSESSMENTS WHERE subject_id IN "
            "(SELECT subject_id FROM dbo.SUBJECTS WHERE subject_code IN (?, ?))",
            SUBJECT_CODE,
            NO_RULE_SUBJECT_CODE,
        )
        cursor.execute(
            f"DELETE FROM dbo.STUDENT_ENROLLMENTS WHERE student_id IN ({placeholders})",
            *ALL_STUDENT_IDS,
        )
        cursor.execute(
            f"DELETE FROM dbo.STUDENTS WHERE student_id IN ({placeholders})",
            *ALL_STUDENT_IDS,
        )
        cursor.execute("DELETE FROM dbo.CLASSES WHERE class_name = ?", CLASS_NAME)
        cursor.execute(
            "DELETE FROM dbo.SUBJECTS WHERE subject_code IN (?, ?)",
            SUBJECT_CODE,
            NO_RULE_SUBJECT_CODE,
        )
        cursor.execute("DELETE FROM dbo.SCHOOL_YEARS WHERE year_name = ?", YEAR_NAME)
        cursor.execute(
            "DELETE FROM dbo.USERS WHERE username IN (?, ?)",
            ADMIN_USERNAME,
            TEACHER_USERNAME,
        )


def fixture_counts(db: DatabaseManager) -> tuple[int, ...]:
    assert_test_database(db)
    with db.transaction() as connection:
        cursor = connection.cursor()
        queries = (
            ("SELECT COUNT(*) FROM dbo.USERS WHERE username LIKE ?", "t138%"),
            ("SELECT COUNT(*) FROM dbo.STUDENTS WHERE student_code LIKE ?", "T138%"),
            ("SELECT COUNT(*) FROM dbo.SCHOOL_YEARS WHERE year_name LIKE ?", "T138%"),
            ("SELECT COUNT(*) FROM dbo.CLASSES WHERE class_name LIKE ?", "%T138%"),
            ("SELECT COUNT(*) FROM dbo.SUBJECTS WHERE subject_code LIKE ?", "T138%"),
            ("SELECT COUNT(*) FROM dbo.ASSESSMENTS WHERE assessment_name LIKE ?", "T138%"),
        )
        return tuple(
            int(cursor.execute(query, value).fetchone()[0])
            for query, value in queries
        )


@dataclass(frozen=True)
class SeedData:
    grade_id: int
    year_id: int
    class_id: int
    subject_id: int
    no_rule_subject_id: int
    assessment_ids: tuple[int, ...]
    enrollment_ids: tuple[int, ...]
    admin_id: int
    teacher_id: int


def seed(db: DatabaseManager) -> SeedData:
    assert_test_database(db)
    user_service = UserService(db)
    admin = user_service.create_user(
        ADMIN_USERNAME, PASSWORD, "T138 Admin", UserRole.ADMIN
    )
    teacher = user_service.create_user(
        TEACHER_USERNAME, PASSWORD, "T138 Teacher", UserRole.TEACHER
    )

    academic = AcademicRepository()
    students = StudentRepository()
    enrollments = EnrollmentRepository()
    rules = SupportRuleRepository()
    with db.transaction() as connection:
        cursor = connection.cursor()
        grade = cursor.execute(
            "SELECT grade_id FROM dbo.GRADES WHERE grade_number = ?", 6
        ).fetchone()
        if grade is None:
            grade_id = int(
                cursor.execute(
                    "INSERT INTO dbo.GRADES (grade_number, grade_name) "
                    "OUTPUT INSERTED.grade_id VALUES (?, ?)",
                    6,
                    "Khối 6",
                ).fetchone()[0]
            )
        else:
            grade_id = int(grade[0])

        year_id = academic.create_school_year(
            connection,
            YEAR_NAME,
            date(2026, 9, 1),
            date(2027, 5, 31),
            False,
        )
        class_id = academic.create_class(
            connection, CLASS_NAME, grade_id, year_id
        )
        subject_id = academic.create_subject(
            connection, SUBJECT_CODE, SUBJECT_NAME
        )
        no_rule_subject_id = academic.create_subject(
            connection, NO_RULE_SUBJECT_CODE, NO_RULE_SUBJECT_NAME
        )
        assessments = (
            academic.create_assessment(
                connection,
                subject_id,
                year_id,
                ASSESSMENT_NAMES[0],
                1,
                "TEST",
                date(2026, 10, 1),
            ),
            academic.create_assessment(
                connection,
                subject_id,
                year_id,
                ASSESSMENT_NAMES[1],
                1,
                "TEST",
                date(2026, 10, 15),
            ),
            academic.create_assessment(
                connection,
                subject_id,
                year_id,
                ASSESSMENT_NAMES[2],
                1,
                "REVIEW",
                date(2026, 12, 1),
            ),
            academic.create_assessment(
                connection,
                no_rule_subject_id,
                year_id,
                ASSESSMENT_NAMES[3],
                1,
                "TEST",
                date(2026, 11, 1),
            ),
        )
        rules.create(connection, subject_id, year_id, THRESHOLD)

        enrollment_ids = []
        for index, (student_id, code, name) in enumerate(
            zip(STUDENT_IDS, STUDENT_CODES, STUDENT_NAMES)
        ):
            students.create(
                connection,
                student_id,
                StudentCreateData(
                    student_code=code,
                    full_name=name,
                    date_of_birth=date(2014, index + 1, index + 2),
                    gender="Nam" if index == 0 else "Nữ",
                ),
            )
            enrollment_ids.append(
                enrollments.create(
                    connection, student_id, class_id, date(2026, 9, 1)
                ).enrollment_id
            )

    return SeedData(
        grade_id=grade_id,
        year_id=year_id,
        class_id=class_id,
        subject_id=subject_id,
        no_rule_subject_id=no_rule_subject_id,
        assessment_ids=tuple(item.assessment_id for item in assessments),
        enrollment_ids=tuple(enrollment_ids),
        admin_id=admin.user_id,
        teacher_id=teacher.user_id,
    )


def build_context(db: DatabaseManager, session) -> AppContext:
    academic = AcademicService(db)
    student_list = StudentListService(db)
    enrollment = EnrollmentService(db)
    score = ScoreService(db)
    preview = ScoreImportPreviewService(db)
    return AppContext(
        db=db,
        auth_service=AuthService(db),
        permission_service=PermissionService(),
        session=session,
        academic_service=academic,
        student_list_service=student_list,
        student_service=StudentService(db),
        enrollment_service=enrollment,
        student_profile_service=StudentProfileService(db),
        score_service=score,
        score_import_parser=ScoreImportWorkbookParser(),
        score_import_template_service=ScoreImportTemplateService(),
        score_import_preview_service=preview,
        score_import_commit_service=ScoreImportCommitService(
            db, preview_service=preview, score_service=score
        ),
        student_export_service=StudentExportService(student_list),
        score_export_service=ScoreExportService(score),
        report_service=ReportService(db),
        report_export_service=ReportExportService(),
        support_service=SupportService(db),
        user_service=UserService(db),
    )


def login_context(db: DatabaseManager, username: str) -> AppContext:
    auth = AuthService(db)
    session = auth.login(username, PASSWORD)
    context = build_context(db, session)
    context.auth_service = auth
    return context


def select(combo, value) -> None:
    index = combo.findData(value)
    assert index >= 0, f"Không tìm thấy {value!r} trong {combo.objectName()}"
    combo.setCurrentIndex(index)


def select_score_context(page, data: SeedData, assessment_id: int, subject_id=None) -> None:
    select(page.school_year_combo, data.year_id)
    select(page.grade_combo, data.grade_id)
    select(page.class_combo, data.class_id)
    select(page.subject_combo, subject_id or data.subject_id)
    select(page.assessment_combo, assessment_id)


def select_report_context(page, data: SeedData, status=None) -> None:
    select(page.school_year_combo, data.year_id)
    select(page.grade_combo, data.grade_id)
    select(page.class_combo, data.class_id)
    select(page.subject_combo, data.subject_id)
    select(page.status_combo, status)


def fill_template(path: Path, values: tuple[Decimal, ...]) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook[ScoreImportTemplateService.SHEET_TITLE]
        by_id = {student_id: value for student_id, value in zip(STUDENT_IDS, values)}
        for row in range(7, sheet.max_row + 1):
            sheet.cell(row, 4, by_id[sheet.cell(row, 2).value])
        # Metadata is informational. The selected application context remains authoritative.
        sheet["B1"] = "WRONG YEAR"
        sheet["B2"] = "WRONG CLASS"
        sheet["B3"] = "WRONG SUBJECT"
        sheet["B4"] = "WRONG ASSESSMENT"
        workbook.save(path)
    finally:
        workbook.close()


class AcceptedPreviewDialog:
    DialogCode = QDialog.DialogCode

    def __init__(self, *_args, **_kwargs):
        pass

    def exec(self):
        return self.DialogCode.Accepted


def count_student_rows(db: DatabaseManager, table: str) -> int:
    assert table in {"SCORES", "INTERVENTIONS"}
    placeholders = ", ".join("?" for _ in ALL_STUDENT_IDS)
    with db.transaction() as connection:
        return int(
            connection.cursor().execute(
                f"SELECT COUNT(*) FROM dbo.{table} WHERE enrollment_id IN "
                f"(SELECT enrollment_id FROM dbo.STUDENT_ENROLLMENTS "
                f"WHERE student_id IN ({placeholders}))",
                *ALL_STUDENT_IDS,
            ).fetchone()[0]
        )


def workbook_rows(path: Path, sheet_name: str, first_row: int, columns: int):
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook[sheet_name]
        return sheet, [
            tuple(sheet.cell(row, column).value for column in range(1, columns + 1))
            for row in range(first_row, sheet.max_row + 1)
        ]
    finally:
        # Callers needing cell data types open the workbook themselves.
        workbook.close()


def test_admin_login_mainwindow_import_detection_and_all_exports_e2e(
    monkeypatch, tmp_path
):
    qapp()
    db = get_test_db()
    cleanup(db)
    try:
        data = seed(db)
        context = login_context(db, ADMIN_USERNAME)
        assert context.session.user_id == data.admin_id
        assert context.session.role is UserRole.ADMIN
        window = MainWindow(context)
        for key in ("students", "scores", "support", "reports"):
            assert window.can_navigate_to(key) is True

        monkeypatch.setattr("ui.pages.scores_page.QMessageBox.information", lambda *_: None)
        monkeypatch.setattr("ui.pages.scores_page.QMessageBox.warning", lambda *_: None)
        monkeypatch.setattr("ui.pages.students_page.QMessageBox.information", lambda *_: None)
        monkeypatch.setattr("ui.pages.students_page.QMessageBox.critical", lambda *_: None)
        monkeypatch.setattr("ui.pages.support_page.QMessageBox.information", lambda *_: None)
        monkeypatch.setattr("ui.pages.support_page.QMessageBox.warning", lambda *_: None)
        monkeypatch.setattr("ui.report_excel_actions.QMessageBox.information", lambda *_: None)
        monkeypatch.setattr("ui.report_excel_actions.QMessageBox.warning", lambda *_: None)

        window.navigate_to("scores")
        scores_page = window.pages["scores"]
        select_score_context(scores_page, data, data.assessment_ids[0])

        template_path = tmp_path / "T138-template.xlsx"
        monkeypatch.setattr(
            "ui.pages.scores_page.QFileDialog.getSaveFileName",
            lambda *_: (str(template_path), ""),
        )
        assert scores_page.download_import_template() is True
        workbook = load_workbook(template_path, data_only=False)
        try:
            sheet = workbook["Nhập điểm"]
            assert workbook.sheetnames == ["Nhập điểm"]
            assert [sheet.cell(row, 2).value for row in range(1, 5)] == [
                YEAR_NAME,
                CLASS_NAME,
                SUBJECT_NAME,
                ASSESSMENT_NAMES[0],
            ]
            roster_ids = [sheet.cell(row, 2).value for row in range(7, 10)]
            assert len(roster_ids) == len(set(roster_ids)) == 3
            assert set(roster_ids) == set(STUDENT_IDS)
            assert all(sheet.cell(row, 4).value is None for row in range(7, 10))
        finally:
            workbook.close()

        imported_values = (Decimal("6.24"), THRESHOLD, Decimal("9.00"))
        fill_template(template_path, imported_values)
        parsed = context.score_import_parser.parse_workbook(template_path)
        preview = context.score_import_preview_service.preview_import(
            scores_page.score_import_context(), parsed
        )
        assert preview.can_commit is True
        assert preview.invalid_count == 0
        assert {row.student_id: row.enrollment_id for row in preview.rows} == dict(
            zip(STUDENT_IDS, data.enrollment_ids)
        )
        assert preview.context.school_year_id == data.year_id
        assert preview.context.class_id == data.class_id
        assert preview.context.subject_id == data.subject_id
        assert preview.context.assessment_id == data.assessment_ids[0]
        assert parsed.metadata.school_year_name == "WRONG YEAR"

        monkeypatch.setattr(
            "ui.pages.scores_page.QFileDialog.getOpenFileName",
            lambda *_: (str(template_path), ""),
        )
        monkeypatch.setattr(
            "ui.pages.scores_page.ScoreImportPreviewDialog", AcceptedPreviewDialog
        )
        monkeypatch.setattr(
            "ui.pages.scores_page.QMessageBox.question",
            lambda *_: QMessageBox.StandardButton.Yes,
        )
        assert scores_page.import_scores_from_excel() is True

        with db.transaction() as connection:
            rows = connection.cursor().execute(
                "SELECT e.student_id, s.score, s.score_id "
                "FROM dbo.SCORES s INNER JOIN dbo.STUDENT_ENROLLMENTS e "
                "ON e.enrollment_id = s.enrollment_id "
                "WHERE s.assessment_id = ? ORDER BY e.student_id",
                data.assessment_ids[0],
            ).fetchall()
            rule_threshold = connection.cursor().execute(
                "SELECT threshold FROM dbo.SUPPORT_RULES "
                "WHERE subject_id = ? AND school_year_id = ? AND is_active = 1",
                data.subject_id,
                data.year_id,
            ).fetchone()[0]
            interventions = connection.cursor().execute(
                "SELECT i.intervention_id, e.student_id, i.trigger_score_id, i.status "
                "FROM dbo.INTERVENTIONS i INNER JOIN dbo.STUDENT_ENROLLMENTS e "
                "ON e.enrollment_id = i.enrollment_id "
                "WHERE i.subject_id = ? ORDER BY e.student_id",
                data.subject_id,
            ).fetchall()
        assert Decimal(str(rule_threshold)) == THRESHOLD
        assert [(row[0], Decimal(str(row[1]))) for row in rows] == list(
            zip(STUDENT_IDS, imported_values)
        )
        assert len(interventions) == 1
        assert interventions[0][1] == STUDENT_IDS[0]
        assert interventions[0][2] == rows[0][2]
        assert interventions[0][3] == InterventionStatus.DETECTED.value

        # A second low score must reuse the still-open intervention.
        second_template = tmp_path / "T138-second.xlsx"
        second_context = ScoreImportContext(
            data.year_id,
            YEAR_NAME,
            data.class_id,
            CLASS_NAME,
            data.subject_id,
            SUBJECT_NAME,
            data.assessment_ids[1],
            ASSESSMENT_NAMES[1],
        )
        context.score_import_template_service.create_template(
            second_context,
            (
                ScoreImportTemplateStudent(STUDENT_IDS[0], STUDENT_NAMES[0]),
                ScoreImportTemplateStudent(STUDENT_IDS[1], STUDENT_NAMES[1]),
            ),
            second_template,
        )
        fill_template(second_template, (Decimal("5.00"), Decimal("5.50")))
        second_preview = context.score_import_preview_service.preview_import(
            second_context,
            context.score_import_parser.parse_workbook(second_template),
        )
        second_result = context.score_import_commit_service.commit_import(second_preview)
        assert second_result.imported_count == 2
        assert second_result.intervention_created_count == 1
        with db.transaction() as connection:
            open_counts = dict(
                connection.cursor().execute(
                    "SELECT e.student_id, COUNT(*) FROM dbo.INTERVENTIONS i "
                    "INNER JOIN dbo.STUDENT_ENROLLMENTS e "
                    "ON e.enrollment_id = i.enrollment_id "
                    "WHERE e.student_id IN (?, ?) GROUP BY e.student_id",
                    STUDENT_IDS[0],
                    STUDENT_IDS[1],
                ).fetchall()
            )
        assert open_counts == {STUDENT_IDS[0]: 1, STUDENT_IDS[1]: 1}

        intervention_id = int(interventions[0][0])
        support = context.support_service
        support.plan_intervention(
            intervention_id,
            data.teacher_id,
            date(2026, 10, 20),
            "T138 phụ đạo",
            "T138 ghi chú",
        )
        support.start_intervention(intervention_id)
        support.mark_waiting_review(intervention_id)
        completed = support.review_intervention(
            intervention_id,
            review_date=date(2026, 12, 5),
            notes="T138 review",
            assessment_id=data.assessment_ids[2],
            score_value=Decimal("7.00"),
        )
        assert completed.status is InterventionStatus.COMPLETED

        # Add one active roster member after import: score export must retain a blank row.
        with db.transaction() as connection:
            StudentRepository().create(
                connection,
                UNSCORED_STUDENT_ID,
                StudentCreateData(
                    student_code=UNSCORED_STUDENT_CODE,
                    full_name=UNSCORED_STUDENT_NAME,
                    date_of_birth=date(2014, 4, 5),
                    gender="Nam",
                ),
            )
            EnrollmentRepository().create(
                connection,
                UNSCORED_STUDENT_ID,
                data.class_id,
                date(2026, 9, 1),
            )

        score_export = tmp_path / "T138-scores.xlsx"
        monkeypatch.setattr(
            "ui.pages.scores_page.QFileDialog.getSaveFileName",
            lambda *_: (str(score_export), ""),
        )
        assert scores_page.export_scores() is True
        workbook = load_workbook(score_export, data_only=False)
        try:
            sheet = workbook["Bảng điểm"]
            exported = [
                (sheet.cell(row, 2).value, sheet.cell(row, 7).value)
                for row in range(9, 13)
            ]
            expected_by_code = dict(zip(STUDENT_CODES, imported_values))
            assert {code for code, _ in exported} == set(STUDENT_CODES) | {
                UNSCORED_STUDENT_CODE
            }
            assert all(
                Decimal(str(score)) == expected_by_code[code]
                for code, score in exported
                if code in expected_by_code
            )
            assert dict(exported)[UNSCORED_STUDENT_CODE] is None
            assert all(
                sheet.cell(row, 7).data_type == "n"
                for row in range(9, 13)
                if sheet.cell(row, 7).value is not None
            )
            assert sheet["B4"].data_type == "s"
            assert sheet["B5"].data_type == "s"
            assert [sheet.cell(row, 2).value for row in range(3, 7)] == [
                YEAR_NAME,
                CLASS_NAME,
                SUBJECT_NAME,
                ASSESSMENT_NAMES[0],
            ]
            assert all(sheet.cell(row, 3).data_type == "s" for row in range(9, 13))
        finally:
            workbook.close()

        window.navigate_to("students")
        students_page = window.pages["students"]
        select(students_page.filter_widget.school_year_combo, data.year_id)
        select(students_page.filter_widget.grade_combo, data.grade_id)
        select(students_page.filter_widget.class_combo, data.class_id)
        student_export = tmp_path / "T138-students.xlsx"
        monkeypatch.setattr(
            "ui.pages.students_page.QFileDialog.getSaveFileName",
            lambda *_: (str(student_export), ""),
        )
        assert students_page.export_students() is True
        workbook = load_workbook(student_export, data_only=False)
        try:
            sheet = workbook["Danh sách học sinh"]
            rows = [
                tuple(sheet.cell(row, column).value for column in range(2, 9))
                for row in range(8, 12)
            ]
            assert [row[1] for row in rows] == sorted(
                STUDENT_NAMES + (UNSCORED_STUDENT_NAME,),
                key=lambda value: value.casefold(),
            )
            assert len({row[0] for row in rows}) == 4
            assert {row[0] for row in rows} == set(STUDENT_CODES) | {
                UNSCORED_STUDENT_CODE
            }
            by_code = {row[0]: row for row in rows}
            assert by_code[STUDENT_CODES[0]][2].date() == date(2014, 1, 2)
            assert by_code[STUDENT_CODES[0]][3] == "Nam"
            assert by_code[STUDENT_CODES[1]][2].date() == date(2014, 2, 3)
            assert by_code[STUDENT_CODES[1]][3] == "Nữ"
            assert [sheet.cell(row, 2).value for row in range(3, 6)] == [
                YEAR_NAME,
                "Khối 6",
                CLASS_NAME,
            ]
            assert all(row[5] == 6 and row[6] == CLASS_NAME for row in rows)
            assert all(sheet.cell(row, 3).data_type == "s" for row in range(8, 12))
            assert all(sheet.cell(row, 8).data_type == "s" for row in range(8, 12))
        finally:
            workbook.close()

        window.navigate_to("support")
        support_page = window.pages["support"]
        select_report_context(support_page, data, InterventionStatus.COMPLETED)
        support_export = tmp_path / "T138-support.xlsx"
        monkeypatch.setattr(
            "ui.pages.support_page.QFileDialog.getSaveFileName",
            lambda *_: (str(support_export), ""),
        )
        assert support_page.export_support_cases() is True

        window.navigate_to("reports")
        reports_page = window.pages["reports"]
        select_report_context(reports_page, data, InterventionStatus.COMPLETED)
        report_export = tmp_path / "T138-report.xlsx"
        monkeypatch.setattr(
            "ui.report_excel_actions.QFileDialog.getSaveFileName",
            lambda *_: (str(report_export), ""),
        )
        assert reports_page.export_report() is True

        for export_path in (support_export, report_export):
            workbook = load_workbook(export_path, data_only=False)
            try:
                sheet = workbook["Báo cáo bổ trợ"]
                assert sheet.max_column == 13
                assert [sheet.cell(row, 2).value for row in range(7, 14)] == [
                    1,
                    0,
                    0,
                    0,
                    0,
                    0,
                    1,
                ]
                assert sheet.max_row == 16
                assert sheet["B16"].value == STUDENT_CODES[0]
                assert sheet["J16"].value == "Đã đạt ngưỡng"
                assert sheet["K16"].value.date() == date(2026, 12, 5)
                assert Decimal(str(sheet["L16"].value)) == Decimal("7.00")
                assert sheet["M16"].value == "Đạt"
                assert sheet["C16"].data_type == "s"
                assert sheet["E16"].data_type == "s"
                assert sheet["G16"].data_type == "s"
                assert sheet["H16"].data_type == "n"
                assert sheet["L16"].data_type == "n"
            finally:
                workbook.close()

        # The still-detected case has no review and must export explicit hyphens.
        select(reports_page.status_combo, InterventionStatus.DETECTED)
        no_review_export = tmp_path / "T138-no-review.xlsx"
        monkeypatch.setattr(
            "ui.report_excel_actions.QFileDialog.getSaveFileName",
            lambda *_: (str(no_review_export), ""),
        )
        assert reports_page.export_report() is True
        workbook = load_workbook(no_review_export)
        try:
            sheet = workbook["Báo cáo bổ trợ"]
            assert sheet.max_row == 16
            assert sheet["B16"].value == STUDENT_CODES[1]
            assert (sheet["K16"].value, sheet["L16"].value, sheet["M16"].value) == (
                "-",
                "-",
                "-",
            )
        finally:
            workbook.close()

        # Empty filtered snapshot stays a valid workbook with headers and no dummy row.
        select(reports_page.status_combo, InterventionStatus.PLANNED)
        empty_export = tmp_path / "T138-empty.xlsx"
        monkeypatch.setattr(
            "ui.report_excel_actions.QFileDialog.getSaveFileName",
            lambda *_: (str(empty_export), ""),
        )
        assert reports_page.export_report() is True
        workbook = load_workbook(empty_export)
        try:
            sheet = workbook["Báo cáo bổ trợ"]
            assert sheet.max_row == ReportExportService.TABLE_HEADER_ROW
            assert sheet["B7"].value == 0
            assert sheet["B13"].value == 0
        finally:
            workbook.close()

        assert not list(Path.cwd().glob("T138*.xlsx"))
    finally:
        cleanup(db)
        assert fixture_counts(db) == (0, 0, 0, 0, 0, 0)


def test_import_rollback_invalid_file_and_export_write_error_are_safe(
    monkeypatch, tmp_path
):
    qapp()
    db = get_test_db()
    cleanup(db)
    try:
        data = seed(db)
        context = login_context(db, ADMIN_USERNAME)
        rollback_context = ScoreImportContext(
            data.year_id,
            YEAR_NAME,
            data.class_id,
            CLASS_NAME,
            data.no_rule_subject_id,
            NO_RULE_SUBJECT_NAME,
            data.assessment_ids[3],
            ASSESSMENT_NAMES[3],
        )
        rollback_file = tmp_path / "T138-rollback.xlsx"
        context.score_import_template_service.create_template(
            rollback_context,
            tuple(
                ScoreImportTemplateStudent(student_id, name)
                for student_id, name in zip(STUDENT_IDS, STUDENT_NAMES)
            ),
            rollback_file,
        )
        fill_template(
            rollback_file,
            (Decimal("1.00"), Decimal("2.00"), Decimal("3.00")),
        )
        preview = context.score_import_preview_service.preview_import(
            rollback_context,
            context.score_import_parser.parse_workbook(rollback_file),
        )
        assert preview.can_commit is True
        with pytest.raises(MissingSupportRuleError) as exc_info:
            context.score_import_commit_service.commit_import(preview)
        assert "pyodbc" not in str(exc_info.value).lower()
        assert count_student_rows(db, "SCORES") == 0
        assert count_student_rows(db, "INTERVENTIONS") == 0

        window = MainWindow(context)
        window.navigate_to("scores")
        scores_page = window.pages["scores"]
        select_score_context(
            scores_page,
            data,
            data.assessment_ids[3],
            subject_id=data.no_rule_subject_id,
        )
        invalid_file = tmp_path / "T138-invalid.xlsx"
        invalid_file.write_bytes(b"not an xlsx zip archive")

        class CommitMustNotRun:
            called = False

            def commit_import(self, _preview):
                self.called = True
                raise AssertionError("commit must not run")

        blocked_commit = CommitMustNotRun()
        scores_page.score_import_commit_service = blocked_commit
        monkeypatch.setattr(
            "ui.pages.scores_page.QFileDialog.getOpenFileName",
            lambda *_: (str(invalid_file), ""),
        )
        monkeypatch.setattr("ui.pages.scores_page.QMessageBox.warning", lambda *_: None)
        assert scores_page.import_scores_from_excel() is False
        assert blocked_commit.called is False
        assert count_student_rows(db, "SCORES") == 0

        window.navigate_to("students")
        students_page = window.pages["students"]
        before = (count_student_rows(db, "SCORES"), count_student_rows(db, "INTERVENTIONS"))
        monkeypatch.setattr(
            "ui.pages.students_page.QFileDialog.getSaveFileName",
            lambda *_: (str(tmp_path / "missing" / "T138-students.xlsx"), ""),
        )
        warnings = []
        monkeypatch.setattr(
            "ui.pages.students_page.QMessageBox.warning",
            lambda *args: warnings.append(args),
        )
        assert students_page.export_students() is False
        assert warnings
        assert "xuất" in str(warnings[-1][1]).lower()
        assert "pyodbc" not in str(warnings[-1][2]).lower()
        assert before == (
            count_student_rows(db, "SCORES"),
            count_student_rows(db, "INTERVENTIONS"),
        )
    finally:
        cleanup(db)
        assert fixture_counts(db) == (0, 0, 0, 0, 0, 0)


def test_teacher_permissions_navigation_and_stale_window_session_e2e(monkeypatch):
    qapp()
    db = get_test_db()
    cleanup(db)
    try:
        seed(db)
        context = login_context(db, ADMIN_USERNAME)
        old_window = MainWindow(context)
        old_scores = old_window.pages["scores"]
        old_students = old_window.pages["students"]
        assert all(
            old_window.can_navigate_to(key)
            for key in ("students", "scores", "support", "reports")
        )

        # Navigation itself is read-only and produces no workbook side effects.
        before = (count_student_rows(db, "SCORES"), count_student_rows(db, "INTERVENTIONS"))
        for key in ("students", "scores", "support", "reports"):
            old_window.navigate_to(key)
        assert before == (
            count_student_rows(db, "SCORES"),
            count_student_rows(db, "INTERVENTIONS"),
        )
        assert not list(Path.cwd().glob("T138*.xlsx"))

        monkeypatch.setattr("ui.pages.students_page.QMessageBox.critical", lambda *_: None)
        monkeypatch.setattr("ui.pages.scores_page.QMessageBox.warning", lambda *_: None)
        old_window.request_logout()
        assert context.session is None
        assert old_window.can_navigate_to("scores") is False
        assert old_students.export_students() is False
        assert old_scores.import_scores_from_excel() is False

        teacher_session = context.auth_service.login(TEACHER_USERNAME, PASSWORD)
        context.set_session(teacher_session)
        new_window = MainWindow(context)
        assert teacher_session.role is UserRole.TEACHER
        for key in ("students", "scores", "support", "reports"):
            assert new_window.can_navigate_to(key) is True
        assert all(
            button.isEnabled() and not button.isHidden()
            for button in (
                new_window.pages["students"].export_button,
                new_window.pages["scores"].download_template_button,
                new_window.pages["scores"].import_excel_button,
                new_window.pages["scores"].export_button,
                new_window.pages["support"].export_button,
                new_window.pages["reports"].export_button,
            )
        )
        assert old_window.can_navigate_to("students") is False
        assert old_students.export_students() is False
        assert old_scores.import_scores_from_excel() is False
        assert count_student_rows(db, "SCORES") == 0
        assert count_student_rows(db, "INTERVENTIONS") == 0
    finally:
        cleanup(db)
        assert fixture_counts(db) == (0, 0, 0, 0, 0, 0)


def test_step_13_e2e_uses_production_layers_without_legacy_import_or_ui_sql():
    from bootstrap import build_app_context
    from services.import_service import ImportService
    from ui.pages.scores_page import ScoresPage
    from ui.pages.students_page import StudentsPage
    from ui.pages.support_page import SupportPage
    from ui.pages.reports_page import ReportsPage

    wiring = inspect.getsource(build_app_context) + inspect.getsource(MainWindow)
    assert "ImportService" not in wiring
    assert "import_service" not in wiring
    assert not hasattr(AppContext, "import_service")
    assert ImportService is not None  # Legacy API remains importable for compatibility.

    ui_source = "\n".join(
        inspect.getsource(item)
        for item in (ScoresPage, StudentsPage, SupportPage, ReportsPage)
    ).upper()
    for forbidden in (
        "SELECT ",
        "INSERT INTO",
        "UPDATE DBO",
        "DELETE FROM",
        "OPENPYXL",
        "REPOSITORY",
        ".TRANSACTION(",
    ):
        assert forbidden not in ui_source

    assert ".transaction(" in inspect.getsource(
        ScoreImportCommitService.commit_import
    )
    assert "create_scores_and_detect" in inspect.getsource(
        ScoreImportCommitService.commit_import
    )
    assert "list_students" in inspect.getsource(StudentExportService.get_export_data)
    assert "list_score_roster" in inspect.getsource(ScoreExportService.get_export_data)
    assert "get_support_report" in inspect.getsource(ReportService)
