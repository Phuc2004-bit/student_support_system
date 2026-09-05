from __future__ import annotations

import inspect
import os
from datetime import date
from decimal import Decimal

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QWidget

from app_context import AppContext
from bootstrap import build_app_context
from exceptions import ReportExportError
from models.dto.report_dto import (
    SupportReportData,
    SupportReportRow,
    SupportReportSummary,
)
from models.dto.report_export import (
    SupportReportExportContext,
    SupportReportExportData,
)
from models.dto import UserSession
from models.enums import InterventionStatus, UserRole
from services.permission_service import PermissionService
from services.report_export_service import ReportExportService
from ui.pages.reports_page import ReportsPage
from ui.pages.support_page import SupportPage
from ui.main_window import MainWindow
from ui.report_excel_actions import ReportExcelActions


def app():
    return QApplication.instance() or QApplication([])


def summary(total=1, completed=0):
    return SupportReportSummary(
        total_cases=total,
        detected_count=total - completed,
        planned_count=0,
        in_progress_count=0,
        waiting_review_count=0,
        continue_count=0,
        completed_count=completed,
    )


def report_row(
    intervention_id=1,
    status="DETECTED",
    full_name="Học sinh A",
    class_name="6A1",
    subject_name="Vật lý",
    review_date=date(2026, 11, 1),
    review_score=Decimal("4.00"),
    review_result="PASSED",
):
    return SupportReportRow(
        intervention_id=intervention_id,
        student_code="HS01",
        full_name=full_name,
        grade_number=6,
        class_name=class_name,
        subject_code="VL",
        subject_name=subject_name,
        detected_date=date(2026, 10, 1),
        start_date=None,
        status=status,
        trigger_score=Decimal("2.75"),
        latest_review_date=review_date,
        latest_review_score=review_score,
        latest_review_result=review_result,
        subject_id=11,
    )


def snapshot(rows=None, report_summary=None):
    values = (report_row(),) if rows is None else tuple(rows)
    return SupportReportData(report_summary or summary(len(values)), values)


class Writer:
    def __init__(self, error=None):
        self.calls = []
        self.error = error

    def export_xlsx(self, data, path):
        self.calls.append((data, path))
        if self.error:
            raise self.error
        return path


class ReportReader:
    def __init__(self, data=None, error=None):
        self.data = data or snapshot()
        self.error = error
        self.report_calls = []
        self.case_calls = []

    def get_support_report(self, **filters):
        self.report_calls.append(filters)
        if self.error:
            raise self.error
        return self.data

    def get_support_cases(self, **filters):
        self.case_calls.append(filters)
        return list(self.data.rows)


def select_context(page, *, status=InterventionStatus.COMPLETED):
    controls = (
        (page.school_year_combo, "2026-2027", 2),
        (page.grade_combo, "Khối 6", 6),
        (page.class_combo, "6A1", 61),
        (page.subject_combo, "Vật lý", 11),
    )
    for combo, label, value in controls:
        combo.addItem(label, value)
        combo.setCurrentIndex(combo.count() - 1)
    index = page.status_combo.findData(status)
    assert index >= 0
    page.status_combo.setCurrentIndex(index)


def expected_filters():
    return {
        "school_year_id": 2,
        "grade_id": 6,
        "class_id": 61,
        "subject_id": 11,
        "status": InterventionStatus.COMPLETED,
    }


def test_reports_page_export_action_cancel_does_not_call_writer(monkeypatch):
    app(); writer = Writer(); page = ReportsPage(excel_writer=writer)
    assert page.export_button.text() == "Xuất Excel"
    select_context(page)
    page.set_report_data(snapshot())
    monkeypatch.setattr("ui.report_excel_actions.QFileDialog.getSaveFileName", lambda *_: ("", ""))
    assert page.export_report() is False
    assert writer.calls == []


def test_reports_page_exports_exact_visible_snapshot_and_current_filter(monkeypatch):
    app(); writer = Writer(); page = ReportsPage(excel_writer=writer)
    select_context(page)
    visible = snapshot((report_row(status="COMPLETED"),), summary(1, completed=1))
    page.set_report_data(visible)
    monkeypatch.setattr("ui.report_excel_actions.QFileDialog.getSaveFileName", lambda *_: ("report", ""))
    monkeypatch.setattr("ui.report_excel_actions.QMessageBox.information", lambda *_: None)
    assert page.export_report() is True
    assert len(writer.calls) == 1
    export_data, path = writer.calls[0]
    assert export_data.report is visible
    assert export_data.report.rows[0].status == "COMPLETED"
    assert export_data.context == SupportReportExportContext(
        school_year_id=2,
        school_year_name="2026-2027",
        grade_id=6,
        grade_name="Khối 6",
        class_id=61,
        class_name="6A1",
        subject_id=11,
        subject_name="Vật lý",
        status=InterventionStatus.COMPLETED,
    )
    assert path.endswith("report.xlsx")


def test_reports_page_refreshes_once_if_snapshot_filter_is_stale(monkeypatch):
    app(); writer = Writer(); reader = ReportReader(snapshot((report_row(status="COMPLETED"),), summary(1, 1)))
    page = ReportsPage(report_service=reader, excel_writer=writer)
    select_context(page)
    reader.report_calls.clear()
    page.report_data = snapshot((report_row(status="DETECTED"),))
    page._report_filters = page.current_filters().__class__(school_year_id=99)
    monkeypatch.setattr("ui.report_excel_actions.QFileDialog.getSaveFileName", lambda *_: ("report.xlsx", ""))
    monkeypatch.setattr("ui.report_excel_actions.QMessageBox.information", lambda *_: None)
    assert page.export_report() is True
    assert reader.report_calls == [expected_filters()]
    assert writer.calls[0][0].report is reader.data


def test_reports_page_empty_snapshot_exports_without_dummy_rows(monkeypatch, tmp_path):
    app(); exporter = ReportExportService(); page = ReportsPage(excel_writer=exporter)
    select_context(page, status=InterventionStatus.DETECTED)
    page.set_report_data(snapshot((), summary(0)))
    output = tmp_path / "empty.xlsx"
    monkeypatch.setattr("ui.report_excel_actions.QFileDialog.getSaveFileName", lambda *_: (str(output), ""))
    monkeypatch.setattr("ui.report_excel_actions.QMessageBox.information", lambda *_: None)
    assert page.export_report() is True
    workbook = load_workbook(output)
    try:
        assert workbook.active.max_row == ReportExportService.TABLE_HEADER_ROW
        assert workbook.active["B7"].value == 0
    finally:
        workbook.close()


def test_reports_page_writer_error_is_safe(monkeypatch):
    app(); writer = Writer(ReportExportError("Không thể ghi file báo cáo.")); page = ReportsPage(excel_writer=writer)
    select_context(page); page.set_report_data(snapshot())
    monkeypatch.setattr("ui.report_excel_actions.QFileDialog.getSaveFileName", lambda *_: ("report.xlsx", ""))
    warnings = []
    monkeypatch.setattr("ui.report_excel_actions.QMessageBox.warning", lambda *args: warnings.append(args))
    assert page.export_report() is False
    assert warnings and "Không thể ghi" in page.state_label.text()


def test_support_page_export_requires_year_and_has_action():
    app(); page = SupportPage()
    assert page.export_button.text() == "Xuất Excel"
    assert page.export_support_cases() is False


def test_support_page_cancel_does_not_read_or_write(monkeypatch):
    app(); reader = ReportReader(); writer = Writer()
    page = SupportPage(support_report_service=reader, report_export_service=writer)
    select_context(page)
    monkeypatch.setattr("ui.pages.support_page.QFileDialog.getSaveFileName", lambda *_: ("", ""))
    assert page.export_support_cases() is False
    assert reader.report_calls == [] and writer.calls == []


def test_support_page_passes_all_filters_to_one_report_snapshot_and_writer(monkeypatch):
    app(); reader = ReportReader(snapshot((report_row(status="COMPLETED"),), summary(1, 1))); writer = Writer()
    page = SupportPage(support_report_service=reader, report_export_service=writer)
    select_context(page)
    monkeypatch.setattr("ui.pages.support_page.QFileDialog.getSaveFileName", lambda *_: ("support", ""))
    monkeypatch.setattr("ui.pages.support_page.QMessageBox.information", lambda *_: None)
    assert page.export_support_cases() is True
    assert reader.report_calls == [expected_filters()]
    assert reader.case_calls == []
    assert len(writer.calls) == 1
    data, path = writer.calls[0]
    assert data.report is reader.data
    assert data.context.status == InterventionStatus.COMPLETED
    assert path.endswith("support.xlsx")


def test_support_page_read_and_write_errors_do_not_expose_raw_details(monkeypatch):
    app(); reader = ReportReader(error=RuntimeError("raw pyodbc SELECT secret")); writer = Writer()
    page = SupportPage(support_report_service=reader, report_export_service=writer)
    select_context(page)
    monkeypatch.setattr("ui.pages.support_page.QFileDialog.getSaveFileName", lambda *_: ("support.xlsx", ""))
    monkeypatch.setattr("ui.pages.support_page.QMessageBox.warning", lambda *_: None)
    assert page.export_support_cases() is False
    assert "pyodbc" not in page.state_label.text().lower()
    reader.error = None
    writer.error = RuntimeError("raw openpyxl internals")
    assert page.export_support_cases() is False
    assert "openpyxl" not in page.state_label.text().lower()


def test_status_filtered_workbook_has_matching_summary_rows_and_latest_review(tmp_path):
    completed = report_row(status="COMPLETED")
    data = SupportReportExportData(
        SupportReportExportContext(2, "2026-2027", status=InterventionStatus.COMPLETED),
        snapshot((completed,), summary(1, completed=1)),
    )
    path = ReportExportService().export_xlsx(data, tmp_path / "completed.xlsx")
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        assert sheet["B7"].value == 1
        assert sheet["B13"].value == 1
        assert sheet["J16"].value == "Đã đạt ngưỡng"
        assert sheet["K16"].value.date() == date(2026, 11, 1)
        assert sheet["L16"].value == pytest.approx(4.0)
        assert sheet["M16"].value == "Đạt"
    finally:
        workbook.close()


def test_report_export_formula_injection_is_blocked_and_scores_stay_numeric(tmp_path):
    dangerous = report_row(
        full_name="=1+1",
        class_name="+Class",
        subject_name="@Subject",
        review_result="-Review",
    )
    data = SupportReportExportData(
        SupportReportExportContext(
            2, "=YEAR", grade_name="+Grade", class_name="-Class",
            subject_name="@Subject",
        ),
        snapshot((dangerous,)),
    )
    path = ReportExportService().export_xlsx(data, tmp_path / "safe.xlsx")
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        for coordinate in ("B3", "E3", "H3", "K3", "C16", "E16", "G16", "M16"):
            assert sheet[coordinate].data_type == "s"
        assert sheet["H16"].data_type == "n"
        assert sheet["L16"].data_type == "n"
    finally:
        workbook.close()


def test_bootstrap_creates_report_exporter_and_main_window_shares_it():
    context = build_app_context()
    assert isinstance(context.report_export_service, ReportExportService)
    shared = ReportExportService()
    window_context = AppContext(
        db=object(),
        auth_service=object(),
        permission_service=PermissionService(),
        session=UserSession(1, "admin", "Admin", UserRole.ADMIN),
        report_export_service=shared,
    )
    window = MainWindow(window_context)
    assert window.pages["reports"].report_export_service is shared
    assert window.pages["support"].report_export_service is shared


def test_constructor_positional_parent_backward_compatibility():
    app(); parent = QWidget()
    reports = ReportsPage(None, None, parent)
    support = SupportPage(None, None, None, None, None, None, None, None, None, None, None, parent)
    assert reports.parent() is parent
    assert support.parent() is parent


def test_export_entry_points_do_not_read_tables_items_or_details():
    reports_source = inspect.getsource(ReportsPage.export_report).upper()
    support_source = inspect.getsource(SupportPage.export_support_cases).upper()
    assert ".TABLE" not in reports_source
    for forbidden in (".TABLE", "._ITEMS", "GET_DETAIL", "GET_INTERVENTION"):
        assert forbidden not in support_source


def test_pages_do_not_import_openpyxl_repository_or_sql():
    source = (
        inspect.getsource(inspect.getmodule(ReportsPage))
        + inspect.getsource(inspect.getmodule(SupportPage))
        + inspect.getsource(inspect.getmodule(ReportExcelActions))
    ).upper()
    for forbidden in ("OPENPYXL", "REPOSITORY", "SELECT ", "INSERT ", "UPDATE DBO", "DELETE "):
        assert forbidden not in source


def test_report_exporter_remains_database_read_write_independent():
    source = inspect.getsource(ReportExportService).upper()
    for forbidden in ("DATABASE", "REPOSITORY", "TRANSACTION", ".COMMIT(", "INSERT ", "UPDATE DBO", "DELETE "):
        assert forbidden not in source
