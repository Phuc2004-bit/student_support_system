from __future__ import annotations

from datetime import date
from decimal import Decimal
import json
import logging
import os
from pathlib import Path
from secrets import token_urlsafe
import sys

from openpyxl import load_workbook
from PySide6.QtWidgets import QFileDialog

from config.database import db_settings
from config.paths import environment_file_path
from config.settings import settings
from exceptions import ValidationError
from models.dto import StudentCreateData
from models.dto.data_export import ScoreExportContext, StudentExportContext
from models.dto.report_export import SupportReportExportContext, SupportReportExportData
from models.dto.score_import import ScoreImportContext, ScoreImportTemplateStudent
from models.dto.student_filter import StudentFilter
from models.enums import AssessmentStatus, InterventionStatus, UserRole
from ui.dialogs.intervention_detail_dialog import InterventionDetailDialog
from ui.dialogs.login_dialog import LoginDialog
from ui.dialogs.score_import_preview_dialog import ScoreImportPreviewDialog
from ui.dialogs.student_profile_dialog import StudentProfileDialog
from ui.main_window import MainWindow
from ui.theme import APP_BACKGROUND


logger = logging.getLogger(__name__)
PREFIX = os.getenv("PACKAGED_SMOKE_PREFIX", "T153").strip().upper()
if PREFIX not in {"T153", "T154", "T155", "T172"}:
    raise RuntimeError("Unsupported packaged smoke fixture prefix.")
YEAR_NAME = f"{PREFIX}_2627"
CLASS_NAME = f"{PREFIX}_7A"
SUBJECT_CODE = f"{PREFIX}_M1"
SUBJECT_NAME = f"{PREFIX} Môn đóng gói"
ADMIN_USERNAME = f"{PREFIX.lower()}_admin"
TEACHER_USERNAME = f"{PREFIX.lower()}_teacher"
OLD_PASSWORD = f"Aa1!{token_urlsafe(18)}"
NEW_PASSWORD = f"Bb2!{token_urlsafe(18)}"
MANUAL_ASSESSMENT_NAME = f"{PREFIX} Manual"
IMPORT_ASSESSMENT_NAME = f"{PREFIX} Import"
REVIEW_FAIL_ASSESSMENT_NAME = f"{PREFIX} Review 1"
REVIEW_PASS_ASSESSMENT_NAME = f"{PREFIX} Review 2"
STUDENT_CODE_A = f"{PREFIX}01"
STUDENT_CODE_B = f"{PREFIX}02"
FORMULA_NAME = f"={PREFIX} Formula"
THRESHOLD = Decimal("6.25")


def _assert_test_database(context) -> str:
    connection_string = context.db.connection_string
    normalized = connection_string.upper().replace(" ", "")
    if "DATABASE=STUDENT_SUPPORT_DB_TEST;" not in normalized:
        raise RuntimeError("Packaged smoke aborted before write: test database is required.")
    with context.db.transaction() as connection:
        actual = str(connection.cursor().execute("SELECT DB_NAME()").fetchone()[0])
    if actual != "student_support_db_test":
        raise RuntimeError("Packaged smoke aborted before write: DB_NAME guard failed.")
    return actual


def _cleanup(context) -> None:
    _assert_test_database(context)
    with context.db.transaction() as connection:
        cursor = connection.cursor()
        cursor.execute(
            "DELETE FROM dbo.INTERVENTION_REVIEWS WHERE intervention_id IN "
            "(SELECT intervention_id FROM dbo.INTERVENTIONS WHERE enrollment_id IN "
            "(SELECT e.enrollment_id FROM dbo.STUDENT_ENROLLMENTS e "
            "INNER JOIN dbo.STUDENTS s ON s.student_id=e.student_id "
            "WHERE s.student_code LIKE ?))",
            f"{PREFIX}%",
        )
        cursor.execute(
            "DELETE FROM dbo.INTERVENTIONS WHERE enrollment_id IN "
            "(SELECT e.enrollment_id FROM dbo.STUDENT_ENROLLMENTS e "
            "INNER JOIN dbo.STUDENTS s ON s.student_id=e.student_id "
            "WHERE s.student_code LIKE ?)",
            f"{PREFIX}%",
        )
        cursor.execute(
            "DELETE FROM dbo.SCORES WHERE enrollment_id IN "
            "(SELECT e.enrollment_id FROM dbo.STUDENT_ENROLLMENTS e "
            "INNER JOIN dbo.STUDENTS s ON s.student_id=e.student_id "
            "WHERE s.student_code LIKE ?)",
            f"{PREFIX}%",
        )
        cursor.execute(
            "DELETE FROM dbo.SUPPORT_RULES WHERE subject_id IN "
            "(SELECT subject_id FROM dbo.SUBJECTS WHERE subject_code LIKE ?)",
            f"{PREFIX}%",
        )
        cursor.execute(
            "DELETE FROM dbo.ASSESSMENTS WHERE subject_id IN "
            "(SELECT subject_id FROM dbo.SUBJECTS WHERE subject_code LIKE ?)",
            f"{PREFIX}%",
        )
        cursor.execute(
            "DELETE FROM dbo.STUDENT_ENROLLMENTS WHERE student_id IN "
            "(SELECT student_id FROM dbo.STUDENTS WHERE student_code LIKE ?)",
            f"{PREFIX}%",
        )
        cursor.execute(
            "DELETE FROM dbo.STUDENTS WHERE student_code LIKE ?", f"{PREFIX}%"
        )
        cursor.execute("DELETE FROM dbo.CLASSES WHERE class_name LIKE ?", f"{PREFIX}%")
        cursor.execute("DELETE FROM dbo.SUBJECTS WHERE subject_code LIKE ?", f"{PREFIX}%")
        cursor.execute("DELETE FROM dbo.SCHOOL_YEARS WHERE year_name LIKE ?", f"{PREFIX}%")
        cursor.execute("DELETE FROM dbo.USERS WHERE username LIKE ?", f"{PREFIX.lower()}%")


def _residual_counts(context) -> tuple[int, ...]:
    _assert_test_database(context)
    checks = (
        ("dbo.USERS", "username LIKE ?", f"{PREFIX.lower()}%"),
        ("dbo.STUDENTS", "student_code LIKE ?", f"{PREFIX}%"),
        ("dbo.SCHOOL_YEARS", "year_name LIKE ?", f"{PREFIX}%"),
        ("dbo.CLASSES", "class_name LIKE ?", f"{PREFIX}%"),
        ("dbo.SUBJECTS", "subject_code LIKE ?", f"{PREFIX}%"),
        ("dbo.ASSESSMENTS", "assessment_name LIKE ?", f"{PREFIX}%"),
    )
    with context.db.transaction() as connection:
        cursor = connection.cursor()
        return tuple(
            int(cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE {where}", value).fetchone()[0])
            for table, where, value in checks
        )


def _select(combo, value) -> None:
    index = combo.findData(value)
    if index < 0:
        raise AssertionError(f"Packaged UI selector is missing database id {value!r}.")
    combo.setCurrentIndex(index)


def _write_import_score(path: Path, student_id: str, value: Decimal) -> None:
    workbook = load_workbook(path)
    try:
        worksheet = workbook.active
        matched = False
        for row in range(7, worksheet.max_row + 1):
            if worksheet.cell(row, 2).value == student_id:
                worksheet.cell(row, 4, value)
                matched = True
        if not matched:
            raise AssertionError("Generated import template did not contain the fixture student.")
        workbook.save(path)
    finally:
        workbook.close()


def _run_checks(app, context, output_dir: Path) -> dict[str, object]:
    result: dict[str, object] = {}
    result["frozen"] = bool(getattr(sys, "frozen", False))
    result["env_file"] = str(environment_file_path())
    result["executable"] = str(Path(sys.executable).resolve())
    result["working_directory"] = str(Path.cwd().resolve())
    result["runtime_paths"] = tuple(sys.path)
    result["fixture_prefix"] = PREFIX
    result["db_name"] = _assert_test_database(context)
    result["configured_db"] = db_settings.DATABASE
    result["app_version"] = settings.APP_VERSION
    result["pyside6_runtime"] = app is not None
    result["pyodbc_runtime"] = result["db_name"] == "student_support_db_test"
    result["dotenv_runtime"] = (
        Path(result["env_file"]).name == ".env"
        and result["configured_db"] == "student_support_db_test"
    )
    _cleanup(context)

    grades = context.academic_service.list_grades()
    if not grades:
        raise AssertionError("The test database must contain at least one grade.")
    grade = next((item for item in grades if int(item[1]) == 7), grades[0])
    grade_id = int(grade[0])
    year_id = context.academic_service.create_school_year(
        YEAR_NAME, date(2026, 9, 1), date(2027, 5, 31), False
    )
    class_id = context.academic_service.create_class(CLASS_NAME, grade_id, year_id)
    subject_id = context.academic_service.create_subject(SUBJECT_CODE, SUBJECT_NAME)
    manual_assessment = context.academic_service.create_assessment(
        subject_id, year_id, MANUAL_ASSESSMENT_NAME, 1, "TEST", date(2026, 10, 1)
    )
    import_assessment = context.academic_service.create_assessment(
        subject_id, year_id, IMPORT_ASSESSMENT_NAME, 1, "TEST", date(2026, 10, 2)
    )
    review_fail_assessment = context.academic_service.create_assessment(
        subject_id, year_id, REVIEW_FAIL_ASSESSMENT_NAME, 1, "TEST", date(2026, 10, 3)
    )
    review_pass_assessment = context.academic_service.create_assessment(
        subject_id, year_id, REVIEW_PASS_ASSESSMENT_NAME, 1, "TEST", date(2026, 10, 4)
    )
    context.academic_service.create_support_rule(subject_id, year_id, THRESHOLD)

    admin_user = context.user_service.create_user(
        ADMIN_USERNAME, OLD_PASSWORD, f"{PREFIX} Admin", UserRole.ADMIN
    )
    teacher_user = context.user_service.create_user(
        TEACHER_USERNAME, OLD_PASSWORD, f"{PREFIX} Teacher", UserRole.TEACHER
    )
    try:
        context.auth_service.login(ADMIN_USERNAME, "wrong-password")
    except ValidationError as exc:
        result["wrong_password_safe"] = "pyodbc" not in str(exc).lower()
    else:
        raise AssertionError("Wrong packaged password was unexpectedly accepted.")
    admin_session = context.auth_service.login(ADMIN_USERNAME, OLD_PASSWORD)
    result["admin_login"] = admin_session.user_id == admin_user.user_id
    context.set_session(admin_session)

    student_a = context.student_service.create_student(
        StudentCreateData(STUDENT_CODE_A, FORMULA_NAME, date(2014, 1, 2), "Nữ")
    )
    student_b = context.student_service.create_student(
        StudentCreateData(STUDENT_CODE_B, f"{PREFIX} Student", date(2014, 2, 3), "Nam")
    )
    enrollment_a = context.enrollment_service.enroll_student(
        student_a.student_id, class_id, date(2026, 9, 1)
    )
    context.enrollment_service.enroll_student(
        student_b.student_id, class_id, date(2026, 9, 1)
    )
    window = MainWindow(context)
    replacement_window = None
    teacher_window = None
    try:
        login_dialog = LoginDialog(context.auth_service)
        result["v11_login_theme"] = (
            settings.APP_VERSION == "1.2.0"
            and APP_BACKGROUND.upper() in login_dialog.styleSheet().upper()
            and "HỆ THỐNG" in login_dialog.title_label.text()
        )
        login_dialog.close()

        for page_key in MainWindow.PAGE_TITLES:
            window.navigate_to(page_key)
            app.processEvents()
        result["mainwindow_navigation"] = all(
            window.can_navigate_to(page_key) for page_key in MainWindow.PAGE_TITLES
        )
        result["v11_dark_theme"] = (
            APP_BACKGROUND.upper() in window.styleSheet().upper()
            and all(window.pages[key] is not None for key in MainWindow.PAGE_TITLES)
        )
        dashboard = window.pages["dashboard"]
        dashboard.bar_chart.canvas.draw()
        dashboard.donut_chart.canvas.draw()
        result["dashboard_charts"] = bool(
            dashboard.bar_chart.figure and dashboard.donut_chart.figure
        )
        result["matplotlib_runtime"] = result["dashboard_charts"]

        students_page = window.pages["students"]
        students_page.initialize_students()
        result["students_page"] = any(
            item.student_id == student_a.student_id for item in students_page.items
        )

        scores_page = window.pages["scores"]
        scores_page.initialize_scores()
        _select(scores_page.school_year_combo, year_id)
        _select(scores_page.grade_combo, grade_id)
        _select(scores_page.class_combo, class_id)
        _select(scores_page.subject_combo, subject_id)
        _select(scores_page.assessment_combo, manual_assessment.assessment_id)
        app.processEvents()
        result["scores_page"] = scores_page.score_table.rowCount() == 2

        manual_row = next(
            row
            for row in range(scores_page.score_table.rowCount())
            if scores_page.enrollment_id_at_row(row) == enrollment_a.enrollment_id
        )
        scores_page.score_table.item(manual_row, 3).setText("5.00")
        result["manual_score"] = scores_page.save_scores()
        manual_roster = context.score_service.list_score_roster(
            class_id,
            year_id,
            subject_id,
            manual_assessment.assessment_id,
        )
        result["manual_score"] = result["manual_score"] and any(
            item.enrollment_id == enrollment_a.enrollment_id
            and item.score == Decimal("5.00")
            for item in manual_roster
        )
        detected_rows = context.report_service.get_support_cases(
            year_id,
            grade_id,
            class_id,
            subject_id,
        )
        detected_row = next(
            item for item in detected_rows if item.student_id == student_a.student_id
        )
        intervention = context.support_service.get_intervention(
            detected_row.intervention_id
        )
        result["support_detection"] = (
            intervention.status is InterventionStatus.DETECTED
        )

        context.academic_service.update_assessment(
            manual_assessment.assessment_id,
            subject_id,
            year_id,
            MANUAL_ASSESSMENT_NAME,
            1,
            "TEST",
            date(2026, 10, 1),
            AssessmentStatus.LOCKED,
        )
        scores_page.context_filter._reload_assessments()
        result["locked_assessment"] = (
            scores_page.assessment_combo.findData(
                manual_assessment.assessment_id
            ) == -1
        )

        profile_dialog = StudentProfileDialog(
            student_a.student_id, context.student_profile_service
        )
        profile = profile_dialog.load_profile()
        result["student_profile"] = (
            profile.student.student_id == student_a.student_id
            and profile_dialog.tabs.count() == 4
        )
        profile_dialog.close()

        detail_dialog = InterventionDetailDialog(
            intervention.intervention_id,
            context.support_service,
            planning_service=context.support_service,
            user_service=context.user_service,
        )
        detail = detail_dialog.load_detail()
        result["support_detail"] = (
            detail is not None
            and detail.intervention_id == intervention.intervention_id
            and detail_dialog.plan_button.isEnabled()
        )
        detail_dialog.close()

        support_page = window.pages["support"]
        _select(support_page.school_year_combo, year_id)
        _select(support_page.grade_combo, grade_id)
        _select(support_page.class_combo, class_id)
        _select(support_page.subject_combo, subject_id)
        app.processEvents()
        result["support_page"] = any(
            item.intervention_id == intervention.intervention_id
            for item in support_page.items
        )

        workflow_statuses = [intervention.status]
        workflow_statuses.append(
            context.support_service.plan_intervention(
                intervention.intervention_id,
                teacher_user.user_id,
                date(2026, 10, 5),
                "Kèm cặp",
                "Packaged workflow",
            ).status
        )
        workflow_statuses.append(
            context.support_service.start_intervention(
                intervention.intervention_id
            ).status
        )
        workflow_statuses.append(
            context.support_service.mark_waiting_review(
                intervention.intervention_id
            ).status
        )
        workflow_statuses.append(
            context.support_service.review_intervention(
                intervention.intervention_id,
                review_date=date(2026, 10, 6),
                assessment_id=review_fail_assessment.assessment_id,
                score_value=Decimal("6.00"),
            ).status
        )
        workflow_statuses.append(
            context.support_service.continue_intervention(
                intervention.intervention_id
            ).status
        )
        context.support_service.mark_waiting_review(
            intervention.intervention_id
        )
        workflow_statuses.append(InterventionStatus.WAITING_REVIEW)
        workflow_statuses.append(
            context.support_service.review_intervention(
                intervention.intervention_id,
                review_date=date(2026, 10, 7),
                assessment_id=review_pass_assessment.assessment_id,
                score_value=THRESHOLD,
            ).status
        )
        result["support_workflow"] = workflow_statuses == [
            InterventionStatus.DETECTED,
            InterventionStatus.PLANNED,
            InterventionStatus.IN_PROGRESS,
            InterventionStatus.WAITING_REVIEW,
            InterventionStatus.CONTINUE,
            InterventionStatus.IN_PROGRESS,
            InterventionStatus.WAITING_REVIEW,
            InterventionStatus.COMPLETED,
        ]
        completed_dialog = InterventionDetailDialog(
            intervention.intervention_id,
            context.support_service,
            planning_service=context.support_service,
            user_service=context.user_service,
        )
        completed_detail = completed_dialog.load_detail()
        result["no_manual_complete"] = (
            completed_detail.status is InterventionStatus.COMPLETED
            and not hasattr(completed_dialog, "complete_button")
            and not any(
                button.isEnabled()
                for button in (
                    completed_dialog.plan_button,
                    completed_dialog.start_button,
                    completed_dialog.waiting_review_button,
                    completed_dialog.review_button,
                    completed_dialog.continue_button,
                )
            )
        )
        completed_dialog.close()

        reports_page = window.pages["reports"]
        _select(reports_page.school_year_combo, year_id)
        _select(reports_page.grade_combo, grade_id)
        _select(reports_page.class_combo, class_id)
        _select(reports_page.subject_combo, subject_id)
        app.processEvents()
        reports_page.status_chart.canvas.draw()
        reports_page.subject_chart.canvas.draw()
        result["reports_charts"] = reports_page.report_data is not None

        system_page = window.pages["system"]
        result["system_page"] = (
            system_page.tabs.count() == 2
            and system_page.own_profile.user_id == admin_user.user_id
            and all(not hasattr(item, "password_hash") for item in system_page.users)
        )

        open_dialog = QFileDialog(window)
        open_dialog.setAcceptMode(QFileDialog.AcceptMode.AcceptOpen)
        save_dialog = QFileDialog(window)
        save_dialog.setAcceptMode(QFileDialog.AcceptMode.AcceptSave)
        open_dialog.show()
        app.processEvents()
        open_visible = open_dialog.isVisible()
        open_dialog.close()
        save_dialog.show()
        app.processEvents()
        save_visible = save_dialog.isVisible()
        save_dialog.close()
        result["file_dialog"] = (
            open_visible
            and save_visible
            and open_dialog.acceptMode() == QFileDialog.AcceptMode.AcceptOpen
            and save_dialog.acceptMode() == QFileDialog.AcceptMode.AcceptSave
        )

        context.user_service.change_own_password(
            admin_session, OLD_PASSWORD, NEW_PASSWORD
        )
        try:
            context.auth_service.login(ADMIN_USERNAME, OLD_PASSWORD)
        except ValidationError:
            old_password_rejected = True
        else:
            old_password_rejected = False
        new_session = context.auth_service.login(ADMIN_USERNAME, NEW_PASSWORD)
        result["bcrypt"] = old_password_rejected and new_session.user_id == admin_user.user_id

        window.request_logout()
        stale_window_blocked = not window.can_navigate_to("students")
        context.set_session(new_session)
        replacement_window = MainWindow(context)
        result["logout_relogin"] = (
            stale_window_blocked and replacement_window.can_navigate_to("students")
        )

        teacher_session = context.auth_service.login(
            TEACHER_USERNAME, OLD_PASSWORD
        )
        context.set_session(teacher_session)
        teacher_window = MainWindow(context)
        teacher_pages = (
            "dashboard", "students", "scores", "support", "reports", "system"
        )
        for page_key in teacher_pages:
            teacher_window.navigate_to(page_key)
            app.processEvents()
        teacher_system = teacher_window.pages["system"]
        result["teacher_permissions"] = (
            all(teacher_window.can_navigate_to(key) for key in teacher_pages)
            and not teacher_window.can_navigate_to("catalogs")
            and teacher_window.sidebar._buttons["catalogs"].isHidden()
            and teacher_system.tabs.count() == 1
        )
        teacher_window.request_logout()
        result["teacher_logout_guard"] = not teacher_window.can_navigate_to(
            "students"
        )
    finally:
        if teacher_window is not None:
            teacher_window.close()
        if replacement_window is not None:
            replacement_window.close()
        window.close()

    roster = context.enrollment_service.list_class_enrollments(class_id, year_id)
    import_context = ScoreImportContext(
        year_id,
        YEAR_NAME,
        class_id,
        CLASS_NAME,
        subject_id,
        SUBJECT_NAME,
        import_assessment.assessment_id,
        IMPORT_ASSESSMENT_NAME,
    )
    template_path = output_dir / f"{PREFIX}-template.xlsx"
    context.score_import_template_service.create_template(
        import_context,
        tuple(ScoreImportTemplateStudent(item.student_id, item.full_name) for item in roster),
        template_path,
    )
    _write_import_score(template_path, student_a.student_id, Decimal("7.00"))
    _write_import_score(template_path, student_b.student_id, Decimal("5.50"))
    parsed = context.score_import_parser.parse_workbook(template_path)
    preview = context.score_import_preview_service.preview_import(import_context, parsed)
    preview_dialog = ScoreImportPreviewDialog(preview)
    result["excel_preview_dialog"] = (
        preview_dialog.table.rowCount() == 2
        and preview_dialog.import_button.isEnabled()
    )
    preview_dialog.close()
    imported = context.score_import_commit_service.commit_import(preview)
    result["excel_import"] = imported.imported_count == 2
    result["openpyxl_runtime"] = True

    report = context.report_service.get_support_report(
        year_id, grade_id, class_id, subject_id
    )
    result["report_snapshot"] = report.summary.total_cases == 2

    student_path = output_dir / f"{PREFIX}-students.xlsx"
    score_path = output_dir / f"{PREFIX}-scores.xlsx"
    support_path = output_dir / f"{PREFIX}-support.xlsx"
    report_path = output_dir / f"{PREFIX}-report.xlsx"
    context.student_export_service.export_xlsx(
        StudentExportContext(
            StudentFilter(
                school_year_id=year_id, grade_id=grade_id, class_id=class_id
            ),
            YEAR_NAME,
            f"Khối {int(grade[1])}",
            CLASS_NAME,
        ),
        student_path,
    )
    context.score_export_service.export_xlsx(
        ScoreExportContext(
            year_id,
            YEAR_NAME,
            class_id,
            CLASS_NAME,
            subject_id,
            SUBJECT_NAME,
            manual_assessment.assessment_id,
            MANUAL_ASSESSMENT_NAME,
        ),
        score_path,
    )
    export_context = SupportReportExportContext(
        year_id,
        YEAR_NAME,
        grade_id,
        f"Khối {int(grade[1])}",
        class_id,
        CLASS_NAME,
        subject_id,
        SUBJECT_NAME,
    )
    export_data = SupportReportExportData(export_context, report)
    context.report_export_service.export_xlsx(export_data, support_path)
    context.report_export_service.export_xlsx(export_data, report_path)

    student_book = load_workbook(student_path, data_only=False)
    try:
        sheet = student_book.active
        formula_cells = [sheet.cell(row, 3) for row in range(8, sheet.max_row + 1)]
        result["student_export"] = len(formula_cells) == 2
        result["formula_safety"] = any(
            cell.value == FORMULA_NAME and cell.data_type == "s"
            for cell in formula_cells
        )
    finally:
        student_book.close()

    score_book = load_workbook(score_path, data_only=False)
    try:
        sheet = score_book.active
        values = [sheet.cell(row, 7).value for row in range(9, sheet.max_row + 1)]
        result["score_export"] = Decimal("5.0") in {
            Decimal(str(value)) for value in values if value is not None
        } and any(value is None for value in values)
    finally:
        score_book.close()

    for key, path in (("support_export", support_path), ("report_export", report_path)):
        workbook = load_workbook(path, data_only=False)
        try:
            sheet = workbook.active
            result[key] = (
                sheet.max_column == 13
                and sheet["B7"].value == 2
                and sheet.max_row == 17
            )
        finally:
            workbook.close()

    result["excel_files"] = [str(path) for path in (
        template_path, student_path, score_path, support_path, report_path
    )]
    return result


def run_packaged_runtime_smoke(app, context, result_path: str) -> int:
    output = Path(result_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    payload: dict[str, object] = {}
    guarded = False
    exit_code = 1
    try:
        logger.info("Packaged functional smoke started")
        _assert_test_database(context)
        guarded = True
        payload = _run_checks(app, context, output.parent)
        excluded_result_keys = {
            "env_file", "executable", "working_directory", "runtime_paths",
            "fixture_prefix", "db_name", "configured_db", "app_version",
            "excel_files"
        }
        failed_checks = [
            key
            for key, value in payload.items()
            if key not in excluded_result_keys and value is not True
        ]
        if failed_checks:
            raise AssertionError(
                "Packaged smoke checks failed: " + ", ".join(failed_checks)
            )
        exit_code = 0
        logger.info("Packaged functional smoke completed")
    except RuntimeError as exc:
        logger.error("Packaged functional smoke aborted safely: %s", type(exc).__name__)
        payload["error_type"] = type(exc).__name__
    except Exception as exc:
        logger.exception("Packaged functional smoke failed")
        payload["error_type"] = type(exc).__name__
    finally:
        if guarded:
            try:
                _cleanup(context)
                payload["residual_counts"] = _residual_counts(context)
            except Exception as cleanup_error:
                logger.exception("Packaged smoke cleanup failed")
                payload["cleanup_error_type"] = type(cleanup_error).__name__
                exit_code = 1
        output.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    return exit_code
